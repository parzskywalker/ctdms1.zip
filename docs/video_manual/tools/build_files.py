"""Creates the demonstration input files (run with the project's .venv python)."""
import csv
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).parents[3]))
from openpyxl import Workbook
from openpyxl.styles import Font

import demo_data as D

OUT = Path(__file__).parents[1] / "demo_files"
OUT.mkdir(exist_ok=True)
cust = {c[0]: c for c in D.CUSTOMERS}

# ---------------- SAP, item-wise Excel (one row per medicine)
wb = Workbook()
ws = wb.active
ws.title = "Billing documents"
head = ["Billing Document", "Billing Date", "Sold-To Party", "Sold-To Name", "Party GSTIN", "Ship-To City",
        "Net Value", "No. of Packages", "Material", "Material Description", "Qty", "Line Value"]
ws.append(head)
names = dict(D.SKUS)
for no, d, cc, val, cases, lines in D.SAP_INV:
    tot_q = sum(q for _, q in lines)
    used = 0
    for k, (sku, q) in enumerate(lines):
        lv = round(val * q / tot_q, 2) if k < len(lines) - 1 else round(val - used, 2)
        used += lv
        ws.append([no, date(int(d[6:]), int(d[3:5]), int(d[:2])), cc, cust[cc][1], cust[cc][4], cust[cc][3],
                   float(val), cases, sku, names[sku], q, lv])
for c in ws[1]:
    c.font = Font(bold=True)
for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
    for c in row:
        c.number_format = "dd-mm-yyyy"
wb.save(OUT / "DEMO_SAP_billing_documents_itemwise.xlsx")

# ---------------- MARG, CSV sales register
with open(OUT / "DEMO_MARG_sales_register.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["Bill No", "Bill Date", "Branch", "Party Code", "Party Name", "Station", "Net Amount", "Cases", "Narration"])
    for no, d, pc, pn, st, val, cases in D.MARG_INV:
        w.writerow([no, d, "RAIPUR", pc, pn, st, f"{val:.2f}", cases, ""])

# ---------------- CBO, CSV invoice export (customer found through GSTIN)
with open(OUT / "DEMO_CBO_invoice_export.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["Invoice No", "Invoice Date", "Customer Code", "Customer Name", "City", "GSTIN", "Net Amount", "No of Cases"])
    for no, d, cc, nm, city, val, cases, g in D.CBO_INV:
        w.writerow([no, d, cc, nm, city, g, f"{val:.2f}", cases])

# ---------------- transporter rate cards
for t, rates in D.RATES.items():
    with open(OUT / f"DEMO_rate_card_{t}.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Area", "Rate", "Effective from"])
        for area, r in rates.items():
            w.writerow([area, f"{r:.2f}", "01-10-2026"])

# ---------------- transporter bill with deliberate irregularities (Beta Cargo Movers)
with open(OUT / "DEMO_transporter_bill_Beta_Cargo.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["LR No", "Invoice No", "Boxes", "Amount", "Remarks"])
    w.writerow(["BCM-7001", "DEMO-SAP-1002", 9, "216.00", "Rate billed 24 instead of 20"])
    w.writerow(["BCM-7002", "DEMO-SAP-1003", 24, "312.00", ""])
    w.writerow(["BCM-7002", "DEMO-SAP-1003", 24, "312.00", "Same LR billed again"])
    w.writerow(["BCM-7003", "DEMO-MRG-2005", 4, "52.00", ""])
    w.writerow(["BCM-9999", "", 8, "240.00", "Not in our dispatch records"])

# ---------------- QR text for the scan demo
from app.services.einvoice_service import make_demo_qr  # noqa: E402
no, d, co, cu, val, cases = D.QR_INV
comp = {c[0]: c for c in D.COMPANIES}
qr = make_demo_qr(comp[co][3], cust[cu][4], no, date(2026, 10, 17), val)
(OUT / "DEMO_einvoice_QR_for_scanning.txt").write_text(
    "# DEMONSTRATION e-invoice QR (unsigned). Type or paste this line in Dispatch desk > Scan & Gatepass.\n"
    f"# {no} - Company A Pharma (Demo) to Zeta Medical Hall (Demo), Rs {val}\n{qr}\n", encoding="utf-8")
print("files:", sorted(p.name for p in OUT.iterdir()))
