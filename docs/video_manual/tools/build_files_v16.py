"""Builds the v1.6 demonstration files (run with the project's .venv python):
the one-time master workbook (from the application's own template) and three ERP invoice exports."""
import csv
import io
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).parents[3]
sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(ROOT))
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

import demo_data_v16 as D
from app.services import master_upload_service as mus

OUT = Path(__file__).parents[1] / "v1.6" / "demo_files"
OUT.mkdir(parents=True, exist_ok=True)

wb = load_workbook(io.BytesIO(mus.template_workbook()))


def put(sheet, rows):
    ws = wb[sheet]
    keys = [c.key for c in mus.SPEC[sheet].cols]
    for r in rows:
        ws.append([r.get(k) for k in keys])


put("ERP_Systems", [dict(code=c, name=n, connection=h, active=a, notes=no) for c, n, h, a, no in D.ERPS])
put("Locations", [dict(code=c, name=n, city=ci, state="Chhattisgarh", route=r, zone=z, distance_km=k)
                  for c, n, ci, r, z, k in D.LOCATIONS])
put("Companies", [dict(code=c, name=n, marking=m, erp=e, relationship=rel, active="Yes") for c, n, m, e, rel in D.COMPANIES])
put("Depots", [dict(code=c, name=n, locality=lo, city=ci, state="Chhattisgarh", pincode=pin, address=a, marking=m,
                    companies=cos, erp=e, active="Yes") for c, n, lo, ci, pin, a, m, cos, e in D.DEPOTS]
    + [dict(code="DM-TEST", name="Depot without city (Demo)", locality="Somewhere")])  # deliberate: city missing
put("Customer_Types", [dict(code=c, name=n, marking=m or None) for c, n, m in D.CUSTOMER_TYPES])
put("Customers", [dict(code=c, name=n, ctype=t, marking=m or None, location=lo, companies=co, depots=dp or None,
                       erp=e or None, erp_code=ec or None, contact_phone=mob or None, whatsapp=mob or None,
                       city={l[0]: l[2] for l in D.LOCATIONS}.get(lo), state="Chhattisgarh", active="Yes")
                  for c, n, t, m, lo, co, dp, e, ec, mob in D.CUSTOMERS + D.BAD_CUSTOMERS])
put("Transporters", [dict(code=c, name=n, kind=k, payment=p or None, methods=m or None, usual=u or None,
                          trip=t or None, box=b or None, vehicle_no=v or None, mobile=mob, city="Raipur", active="Yes")
                     for c, n, k, p, m, u, t, b, v, mob in D.TRANSPORTERS])
put("Local_Auto_Rates", [dict(code=c, auto=a, location=lo or None, company=co or None, depot=dp or None, rate=r,
                              **{"from": "01-10-2026"}) for c, a, lo, co, dp, r in D.LOCAL_AUTO_RATES])
put("Freight_Rates", [dict(transporter=t, location=lo, rate=r, **{"from": "01-10-2026"}) for t, lo, r in D.FREIGHT_RATES])
put("Invoice_Code_Mappings", [dict(code=c, erp=e, marking=m, match=w, company=co, depot=dp or None, ctype=ct or None,
                                   notes=no) for c, e, m, w, co, dp, ct, no in D.MARKINGS])
put("ERP_Mappings", [dict(erp=e, entity=en, source_code=sc, source_name=sn, target=t) for e, en, sc, sn, t in D.ERP_MAPPINGS])
ws = wb["Instructions"]
ws.insert_rows(2)
ws["A2"] = "DEMONSTRATION WORKBOOK - every record is a demonstration entry (codes DM-..., names ending (Demo))."
ws["A2"].font = Font(bold=True, color="B45309")
wb.save(OUT / "DEMO_v1.6_Master_Workbook.xlsx")


def dmy(s):
    d, m, y = s.split("-")
    return date(int(y), int(m), int(d))


sap = Workbook()
s = sap.active
s.title = "Billing documents"
s.append(["Billing Document", "Billing Date", "Plant", "Sold-To Party", "Sold-To Name", "Net Value", "No. of Packages"])
for no, d, plant, cu, nm, val, cases in D.SAP_INV:
    s.append([no, dmy(d), plant, cu, nm, float(val), cases])
for c in s[1]:
    c.font = Font(bold=True)
for row in s.iter_rows(min_row=2, min_col=2, max_col=2):
    for c in row:
        c.number_format = "dd-mm-yyyy"
sap.save(OUT / "DEMO_v1.6_SAP_billing_ABC.xlsx")

with open(OUT / "DEMO_v1.6_MARG_sales_XYZ.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["Bill No", "Bill Date", "Party Code", "Party Name", "Net Amount", "Cases"])
    for no, d, pc, pn, val, cases in D.MARG_INV:
        w.writerow([no, d, pc, pn, f"{val:.2f}", cases])

with open(OUT / "DEMO_v1.6_CBO_invoices_PQR.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["Invoice No", "Invoice Date", "Customer Code", "Customer Name", "Net Amount", "No of Cases"])
    for no, d, cc, nm, val, cases in D.CBO_INV:
        w.writerow([no, d, cc, nm, f"{val:.2f}", cases])
print("built:", sorted(p.name for p in OUT.iterdir()))
