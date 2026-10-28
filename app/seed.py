"""Realistic demo data generator.

Everything is created through the real services (import pipeline, freight engine, dispatch workflow,
override approvals, reconciliation) with a simulated clock, so the demo database looks like ~3 months
of genuine operation. Dates are relative to the day the seed is run.

    python -m app.cli seed-demo --reset
"""
from __future__ import annotations

import csv
import io
import random
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import BASE_DIR, get_settings
from app.models import (
    CFA, Company, Customer, DispatchMode, DispatchSlip, ERPCodeMapping, ERPSystem, ImportTemplate, Invoice,
    InvoiceStatus, Location, OutsideAuto, PickupMode, Role, Transporter, User,
)
from app.models.base import set_fake_now
from app.security.auth import hash_password
from app.security.permissions import DEFAULT_ROLES
from app.services import dispatch_service as ds
from app.services import import_service as imp
from app.services import invoice_service, rate_service, recon_service
from app.services import settings_service as cfg

SAMPLE_DIR = BASE_DIR / "sample_data"
TZ = ZoneInfo("Asia/Kolkata")
RND = random.Random(20260926)


def at(d: date, hh: int, mm: int = 0) -> datetime:
    """Local (IST) date+time -> naive UTC for the fake clock."""
    from datetime import timezone
    return datetime.combine(d, time(hh, mm)).replace(tzinfo=TZ).astimezone(timezone.utc).replace(tzinfo=None)


_DB: Session | None = None


def tick(when: datetime) -> None:
    """Advance the simulated clock. Pending rows are flushed first so their default timestamps
    (evaluated at INSERT time) belong to the moment they were created."""
    if _DB is not None:
        _DB.flush()
    set_fake_now(when)


def clock(d: date, hh: int, mm: int = 0) -> None:
    tick(at(d, hh, mm))


# ----------------------------------------------------------------------------- master data
ERPS = [
    ("SAP", "SAP S/4HANA", "sap_odata", "ODATA", "Company A billing in SAP. OData credentials pending from Company A IT."),
    ("MARG", "MARG ERP 9+", "marg", "FILE", "Company B billing in MARG. Daily sales-register export; DB access TBD."),
    ("CBO", "CBO", "cbo", "API", "Company C billing in CBO. API documentation awaited from vendor."),
    ("INTRECT", "Intrect", "intrect", "API", "Own distribution billing in Intrect. API key to be issued by Intrect."),
    ("GENERIC", "Other ERP / manual Excel", "file", "FILE", "Standard CTDMS Excel/CSV template for any other system."),
    ("DEMO", "Demo ERP API (simulator)", "demo", "API", "Simulated API used to demonstrate the automated pull pipeline."),
]

CFAS = [
    ("RPR", "Guma Depot", "Raipur Central Warehouse", "Plot 14, Guma", "Raipur", "Chhattisgarh", "493221"),
    ("BSP", "Bilaspur Depot", "Bilaspur Depot", "Vyapar Vihar, Link Road", "Bilaspur", "Chhattisgarh", "495001"),
    ("DRG", "Durg Depot", "Durg Depot", "GE Road, Industrial Estate", "Durg", "Chhattisgarh", "491001"),
    # v1.6: a second depot in the same city (one city -> several depots)
    ("KVR", "Kamal Vihar Depot", "Kamal Vihar Godown", "Sector 4, Kamal Vihar", "Raipur", "Chhattisgarh", "492001"),
]

LOCATIONS = [  # code, name, state, pincode, route, zone, km (from Raipur)
    ("RPR", "Raipur", "Chhattisgarh", "492001", "R-LOCAL", "Central", 0),
    ("BSP", "Bilaspur", "Chhattisgarh", "495001", "R-NORTH", "North", 118),
    ("DRG", "Durg", "Chhattisgarh", "491001", "R-WEST", "West", 38),
    ("BHL", "Bhilai", "Chhattisgarh", "490001", "R-WEST", "West", 30),
    ("KRB", "Korba", "Chhattisgarh", "495677", "R-NORTH", "North", 200),
    ("NGP", "Nagpur", "Maharashtra", "440001", "R-MH", "Maharashtra", 290),
    ("RJN", "Rajnandgaon", "Chhattisgarh", "491441", "R-WEST", "West", 70),
    ("RGH", "Raigarh", "Chhattisgarh", "496001", "R-EAST", "East", 250),
    # v1.3: areas served by the real transporters on the sample bills
    ("JNJ", "Janjgir", "Chhattisgarh", "495668", "R-EAST", "East", 170),
    ("KWD", "Kawardha", "Chhattisgarh", "491995", "R-WEST", "West", 120),
    ("MDGR", "Manendragarh", "Chhattisgarh", "497442", "R-NORTH", "North", 390),
    ("BKPR", "Baikunthpur", "Chhattisgarh", "497335", "R-NORTH", "North", 340),
    ("AMB", "Ambikapur", "Chhattisgarh", "497001", "R-NORTH", "North", 330),
]

COMPANIES = [  # code, name, erp, type, cfas
    ("COA", "Company A Pharma Ltd", "SAP", "PRINCIPAL", ["RPR", "BSP", "DRG"]),
    ("COB", "Company B Healthcare Ltd", "MARG", "PRINCIPAL", ["RPR", "BSP", "KVR"]),
    ("COC", "Company C Lifesciences Ltd", "CBO", "PRINCIPAL", ["RPR", "DRG"]),
    ("OWN", "Own Distribution Division", "INTRECT", "OWN", ["RPR"]),
]

# code, name, location, contact, phone, companies, erp codes {erp: code}
CUSTOMERS = [
    ("C-ABC01", "ABC Pharma", "BSP", "Rakesh Agrawal", "9826011111", ["COA", "COB", "COC", "OWN"],
     {"SAP": "100245", "MARG": "ABC001", "CBO": "ABCPH01", "INTRECT": "P-0012"}),
    ("C-XYZ01", "XYZ Medical", "DRG", "Suresh Jain", "9826022222", ["COA", "COB", "COC", "OWN"],
     {"SAP": "100311", "MARG": "XYZ002", "CBO": "XYZMD01", "INTRECT": "P-0019"}),
    ("C-PQR01", "PQR Distributors", "NGP", "Anil Deshmukh", "9823033333", ["COA", "COB", "COC"],
     {"SAP": "100402", "MARG": "PQR003", "CBO": "PQRDS01"}),
    ("C-SUN01", "Sunrise Agencies", "KRB", "Manoj Sahu", "9827044444", ["COA", "COB", "COC", "OWN"],
     {"SAP": "100518", "MARG": "SUN004", "CBO": "SUNAG01", "INTRECT": "P-0027"}),
    ("C-MET01", "Metro Pharma", "RPR", "Vikas Tiwari", "9826055555", ["COA", "COB", "COC", "OWN"],
     {"SAP": "100623", "MARG": "MET005", "CBO": "METPH01", "INTRECT": "P-0031"}),
    ("C-SGM01", "Shree Ganesh Medicos", "BHL", "Pankaj Verma", "9826066666", ["COA", "COB", "COC", "OWN"],
     {"SAP": "100731", "MARG": "SGM006", "CBO": "SGMED01", "INTRECT": "P-0036"}),
    ("C-JMH01", "Jeevan Medical Hall", "RJN", "Deepak Chandrakar", "9826077777", ["COA", "COB", "COC"],
     {"SAP": "100844", "MARG": "JMH007", "CBO": "JEEVN01"}),
    ("C-LDH01", "Lifeline Drug House", "RGH", "Sanjay Patel", "9826088888", ["COA", "COB"],
     {"SAP": "100952", "MARG": "LDH008"}),
    ("C-CCC01", "City Care Chemists", "RPR", "Neha Sharma", "9826099999", ["COA", "COB", "COC", "OWN"],
     {"SAP": "101067", "MARG": "CCC009", "CBO": "CITYC01", "INTRECT": "P-0044"}),
    ("C-VDP01", "Vardaan Pharma", "BSP", "Ramesh Kashyap", "9826010101", ["COA", "COB", "OWN"],
     {"SAP": "101178", "MARG": "VDP010", "INTRECT": "P-0052"}),
    # v1.3: customers on the real transporter bills (stockists, retailers) and hospitals
    ("C-UTP01", "Uttam Pharma", "BKPR", "Uttam Gupta", "9826120001", ["COA", "COB", "COC", "OWN"],
     {"SAP": "101201", "MARG": "UTP011", "CBO": "UTTAM01", "INTRECT": "P-0061"}),
    ("C-SMD01", "Singh Medicose", "MDGR", "Gurpreet Singh", "9826120002", ["COA", "COB", "OWN"],
     {"SAP": "101202", "MARG": "SMD012", "INTRECT": "P-0062"}),
    ("C-SAI01", "Sai Traders", "MDGR", "Sai Prasad", "9826120003", ["COA", "COB"],
     {"SAP": "101203", "MARG": "SAI013"}),
    ("C-SBV01", "Sambhav Medicals", "KWD", "Rahul Jain", "9826120004", ["COA", "COB", "COC", "OWN"],
     {"SAP": "101204", "MARG": "SBV014", "CBO": "SAMBH01", "INTRECT": "P-0064"}),
    ("C-MDV01", "Shri Mahadev Pharma", "KWD", "Mahadev Sahu", "9826120005", ["COA", "COC"],
     {"SAP": "101205", "CBO": "MAHAD01"}),
    ("C-DHR01", "District Hospital Raigarh", "RGH", "Store Officer", "9826120006", ["COA", "COB", "COC"],
     {"SAP": "101206", "MARG": "DHR016", "CBO": "DHRGH01"}),
    ("C-JNH01", "Janjgir Nursing Home", "JNJ", "Dr. A. Kashyap", "9826120007", ["COA", "COB", "COC"],
     {"SAP": "101207", "MARG": "JNH017", "CBO": "JNJNH01"}),
    ("C-AIM01", "AIIMS Raipur Pharmacy", "RPR", "Purchase Cell", "9826120008", ["COA", "COB", "COC", "OWN"],
     {"SAP": "101208", "MARG": "AIM018", "CBO": "AIIMS01", "INTRECT": "P-0068"}),
    ("C-SSM01", "Shri Sai Medicose", "AMB", "Sanjay Agrawal", "9826120009", ["COA", "COB", "OWN"],
     {"SAP": "101209", "MARG": "SSM019", "INTRECT": "P-0069"}),
    ("C-HRM01", "Hiramani Medical Agencies", "BSP", "Hiramani Sahu", "9826120010", ["COA", "COB", "COC", "OWN"],
     {"SAP": "101210", "MARG": "HRM020", "CBO": "HIRAM01", "INTRECT": "P-0070"}),
]
CUSTOMER_TYPES = {  # everything else is a stockist
    "C-XYZ01": "RETAILER", "C-MET01": "RETAILER", "C-SGM01": "RETAILER", "C-JMH01": "RETAILER",
    "C-CCC01": "RETAILER", "C-SMD01": "RETAILER", "C-MDV01": "RETAILER", "C-SSM01": "RETAILER",
    "C-HRM01": "RETAILER", "C-DHR01": "HOSPITAL", "C-JNH01": "HOSPITAL", "C-AIM01": "HOSPITAL",
}
COMPANY_GSTIN = {"COA": "22AAACA1234A1Z5", "COB": "22AABCB5678B1Z3", "COC": "22AAACC9012C1Z1",
                 "OWN": "22AAJCT1963H1ZC"}

TRANSPORTERS = [  # code, name, contact, mobile, gstin, area, vehicles, lr prefix, payment mode, direct pickup
    ("CGR", "Chhattisgarh Roadlines", "Harish Yadav", "9425201234", "22AABCC1234F1Z5", "Bilaspur, Korba, Raigarh",
     ["CG10AB2211", "CG10AB2212", "CG12C7788"], "CGR", "TBB", None),
    ("SCM", "Speed Cargo Movers", "Imran Khan", "9425205678", "22AADCS5678G1Z2", "Raipur local, Bhilai, Durg",
     ["CG04LM4455", "CG04LM4456"], "SCM", "TBB", None),
    ("SBT", "Shree Balaji Transport", "Mahesh Sahu", "9425209012", "22AAFFS9012H1Z9", "Durg, Bhilai, Rajnandgaon",
     ["CG07BT1010", "CG07BT2020"], "SBT", "TBB", None),
    ("VDL", "Vidarbha Logistics", "Prakash Wankhede", "9422103456", "27AAGCV3456J1Z1", "Nagpur, Maharashtra",
     ["MH31FK9090", "MH40AS1212"], "VDL", "TBB", None),
    # the six transporters whose bills management shared (v1.3)
    ("PGC", "Prakash Goods Carrier", "Prakash Agrawal", "9425211001", "22ABYPV9378F1ZP", "Manendragarh, Baikunthpur",
     [], "", "TBB", None),
    ("BJR", "Bhaiji Roadlines", "Bhaiji", "9425211002", None, "Ambikapur", [], "", "TBB", None),
    ("OSHO", "Osho Express", "Office Raipur", "9993233974", None, "Bilaspur (daily service)", [], "", "TBB", None),
    ("ORL", "Orient Logistics Service", "Office Bhaisthan", "9425211004", "22ACIPA4333N1ZU", "Janjgir, Raigarh",
     [], "", "TBB", None),
    ("PRK", "Parakh Transport", "Parakh", "9425295353", None, "Durg, Bhilai, Rajnandgaon", [], "", "TBB",
     ("150.00", "PER_TRIP")),
    ("JBG", "Jai Bhole Garage (Goods Transport)", "Manager", "7024032306", None, "Kawardha, Pondi, Bodla", [], "",
     "CASH", None),
]

# carrier per delivery area; "AUTO" = local outside auto, "SELF" = customer collects
TRANSPORTER_BY_LOCATION = {"RPR": "AUTO", "BHL": "PRK", "DRG": "PRK", "RJN": "PRK", "BSP": "OSHO", "KRB": "CGR",
                           "RGH": "ORL", "NGP": "VDL", "JNJ": "ORL", "KWD": "JBG", "MDGR": "PGC", "BKPR": "PGC",
                           "AMB": "BJR"}

# transporter rate cards: rate per case per area (the bills show ₹35 Parakh, ₹38/₹50 Orient, ₹40 Jai Bhole ...)
TRANSPORTER_RATES = {
    "CGR": {"BSP": "30.00", "KRB": "45.00", "RGH": "42.00"},
    "SCM": {"RPR": "15.00", "BHL": "25.00", "DRG": "25.00"},
    "SBT": {"DRG": "30.00", "BHL": "28.00", "RJN": "35.00", "RPR": "20.00"},
    "VDL": {"NGP": "60.00"},
    "PGC": {"MDGR": "50.00", "BKPR": "48.00"},
    "BJR": {"AMB": "100.00"},
    "OSHO": {"BSP": "25.00"},
    "ORL": {"JNJ": "50.00", "RGH": "38.00"},
    "PRK": {"DRG": "35.00", "BHL": "35.00", "RJN": "35.00"},
    "JBG": {"KWD": "40.00"},
}
AUTOS = [("AUTO-RAM", "Ramesh Yadav (auto)", "9826230011", "PER_TRIP", "250.00", "Raipur local deliveries"),
         ("AUTO-SONU", "Sonu Tempo", "9826230012", "PER_CASE", "12.00", "To transporter booking offices")]
# v1.6: a local delivery vehicle that may be charged either way (chosen on the gatepass) + its per-box rates
AUTO_BOTH = ("AUTO-KISH", "Kishore Mini Truck", "9826230013", "CG04-LT-5521", "600.00", "18.00")

# (company, cfa, location, rate) - standard CFA-location rates, effective from FY start
RATES = [
    ("COA", "RPR", "BSP", "8.00"), ("COA", "RPR", "DRG", "6.00"), ("COA", "RPR", "NGP", "12.00"),
    ("COA", "RPR", "RPR", "3.00"), ("COA", "RPR", "BHL", "5.00"), ("COA", "RPR", "KRB", "10.00"),
    ("COA", "RPR", "RJN", "7.00"), ("COA", "RPR", "RGH", "11.00"),
    ("COA", "BSP", "BSP", "3.00"), ("COA", "BSP", "KRB", "6.00"), ("COA", "BSP", "RGH", "7.00"),
    ("COA", "BSP", "RPR", "8.00"),
    ("COA", "DRG", "DRG", "3.00"), ("COA", "DRG", "BHL", "3.50"), ("COA", "DRG", "RJN", "5.00"),
    ("COA", "DRG", "RPR", "6.00"),
    ("COB", "RPR", "BSP", "9.00"), ("COB", "RPR", "DRG", "7.00"), ("COB", "RPR", "NGP", "14.00"),
    ("COB", "RPR", "RPR", "3.50"), ("COB", "RPR", "BHL", "6.00"), ("COB", "RPR", "RJN", "8.00"),
    ("COB", "RPR", "RGH", "12.00"),
    ("COB", "BSP", "BSP", "3.50"), ("COB", "BSP", "RGH", "8.00"), ("COB", "BSP", "RPR", "9.00"),
    ("COC", "RPR", "RPR", "4.00"), ("COC", "RPR", "BSP", "8.50"), ("COC", "RPR", "DRG", "6.50"),
    ("COC", "RPR", "BHL", "6.00"), ("COC", "RPR", "NGP", "13.00"), ("COC", "RPR", "RJN", "7.50"),
    ("COC", "DRG", "DRG", "3.00"), ("COC", "DRG", "BHL", "3.50"), ("COC", "DRG", "RJN", "5.50"),
    ("OWN", "RPR", "RPR", "2.50"), ("OWN", "RPR", "BSP", "7.50"), ("OWN", "RPR", "DRG", "5.50"),
    ("OWN", "RPR", "BHL", "5.00"), ("OWN", "RPR", "KRB", "9.00"),
]
# NOTE: deliberately NO rate for Company C -> Korba (demonstrates RATE MISSING) and no OWN/Raigarh
# (covered by the explicit OWN company default rate).

USERS = [  # username, full name, role, cfas (None = all)
    ("admin", "System Administrator", "ADMIN", None, "Admin@123"),
    ("management", "Arvind Kumar", "MANAGEMENT", None, "Demo@123"),
    ("accounts", "Kavita Joshi", "ACCOUNTS", None, "Demo@123"),
    ("cfa.raipur", "Rohit Mishra", "CFA_MANAGER", ["RPR"], "Demo@123"),
    ("operator.raipur", "Sunil Dewangan", "DISPATCH_OPERATOR", ["RPR"], "Demo@123"),
    ("operator.bilaspur", "Ajay Kaushik", "DISPATCH_OPERATOR", ["BSP"], "Demo@123"),
    ("operator.durg", "Tarun Sinha", "DISPATCH_OPERATOR", ["DRG"], "Demo@123"),
    ("checker", "Meena Gupta", "CHECKER", None, "Demo@123"),
    ("viewer", "Audit Viewer", "VIEWER", None, "Demo@123"),
]


@dataclass
class Ctx:
    db: Session
    today: date
    start: date
    users: dict
    erps: dict
    companies: dict
    cfas: dict
    locations: dict
    customers: dict
    transporters: dict


def reset_database() -> None:
    from app.db import Base, engine
    import app.models  # noqa: F401
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)


def seed_core(db: Session) -> dict:
    """Roles + settings (needed in every installation)."""
    roles = {}
    for code, d in DEFAULT_ROLES.items():
        r = db.scalar(select(Role).where(Role.code == code))
        if not r:
            r = Role(code=code, name=d["name"], description=d["description"], permissions=d["permissions"],
                     is_system=True)
            db.add(r)
        roles[code] = r
    cfg.ensure_defaults(db)
    ensure_customer_types(db)
    db.flush()
    return roles


def ensure_customer_types(db: Session) -> None:
    """v1.6: the configurable customer types (Retailer, Stockist, Distributor, ...) with their recognition codes."""
    from app.models import CustomerType, CustomerTypeMaster
    have = {c.code for c in db.scalars(select(CustomerTypeMaster))}
    for n, (code, name, marking) in enumerate(CustomerType.DEFAULTS, 1):
        if code not in have:
            db.add(CustomerTypeMaster(code=code, name=name, marking=marking, sort_order=n, is_active=True))


def seed_masters(db: Session, today: date) -> Ctx:
    roles = seed_core(db)
    erps = {}
    for code, name, adapter, mode, notes in ERPS:
        e = ERPSystem(code=code, name=name, adapter=adapter, integration_mode=mode, notes=notes,
                      import_format="CSV/XLSX", config={"records_per_run": "4"} if adapter == "demo" else {})
        db.add(e)
        erps[code] = e
    db.flush()
    cfas = {}
    for code, name, wh, addr, city, state, pin in CFAS:
        c = CFA(code=code, name=name, warehouse=wh, address=addr, city=city, state=state, pincode=pin,
                locality={"RPR": "Guma", "BSP": "Vyapar Vihar", "DRG": "Industrial Estate", "KVR": "Kamal Vihar"}[code],
                marking={"RPR": "GUMA", "KVR": "KVIHAR"}.get(code),
                contact_person="Warehouse In-charge", contact_phone="0771-40" + str(RND.randint(10000, 99999)))
        db.add(c)
        cfas[code] = c
    locations = {}
    for code, name, state, pin, route, zone, km in LOCATIONS:
        lo = Location(code=code, name=name, city=name, state=state, pincode=pin, route=route, zone=zone,
                      distance_km=Decimal(km))
        db.add(lo)
        locations[code] = lo
    companies = {}
    for code, name, erp, rtype, cfa_codes in COMPANIES:
        co = Company(code=code, name=name, erp_system=erps[erp], relationship_type=rtype,
                     contact_person="Distribution Manager", contact_phone="022-6" + str(RND.randint(1000000, 9999999)),
                     email=f"logistics@{code.lower()}.example.com", cfas=[cfas[c] for c in cfa_codes],
                     gstin=COMPANY_GSTIN.get(code))
        db.add(co)
        companies[code] = co
    db.flush()
    customers = {}
    for n, (code, name, loc, person, phone, cos, erp_codes) in enumerate(CUSTOMERS):
        lo = locations[loc]
        letters = "".join(ch for ch in name.upper() if ch.isalpha())[:3].ljust(3, "X")
        gstin = (f"{'22' if lo.state == 'Chhattisgarh' else '27'}AA{letters}{4100 + n * 37:04d}"
                 f"{chr(65 + n % 26)}1Z{(n * 7) % 10}")
        cu = Customer(code=code, name=name, location=lo, city=lo.name, state=lo.state, pincode=lo.pincode,
                      party_type=CUSTOMER_TYPES.get(code, "STOCKIST"), whatsapp=phone,
                      billing_address=f"{RND.randint(1, 250)}, Main Market, {lo.name}",
                      delivery_address=f"Godown {RND.randint(1, 30)}, Transport Nagar, {lo.name}",
                      contact_person=person, contact_phone=phone, email=f"orders@{name.split()[0].lower()}.example.in",
                      gstin=gstin, drug_license_no=f"CG-DL-{RND.randint(10000, 99999)}",
                      companies=[companies[c] for c in cos], cfas=list(cfas.values()))
        db.add(cu)
        customers[code] = cu
    db.flush()
    # ERP code mappings: the same party has a different code in every ERP
    for code, name, loc, person, phone, cos, erp_codes in CUSTOMERS:
        for erp, src in erp_codes.items():
            db.add(ERPCodeMapping(erp_system=erps[erp], entity_type="CUSTOMER", source_code=src, source_name=name.upper(),
                                  target_id=customers[code].id))
    # company / CFA codes inside the ERPs
    for erp, src, tgt in [("SAP", "1000", "COA"), ("CBO", "CLS", "COC"), ("INTRECT", "OWN-DIST", "OWN")]:
        db.add(ERPCodeMapping(erp_system=erps[erp], entity_type="COMPANY", source_code=src, target_id=companies[tgt].id,
                              source_name=companies[tgt].name))
    for erp, src, tgt in [("SAP", "1101", "RPR"), ("SAP", "1102", "BSP"), ("SAP", "1103", "DRG"),
                          ("MARG", "RAIPUR", "RPR"), ("MARG", "BILASPUR", "BSP"),
                          ("CBO", "D-RPR", "RPR"), ("CBO", "D-DRG", "DRG"), ("INTRECT", "RAIPUR-HO", "RPR")]:
        db.add(ERPCodeMapping(erp_system=erps[erp], entity_type="CFA", source_code=src, target_id=cfas[tgt].id))
    db.add(ERPCodeMapping(erp_system=erps["MARG"], entity_type="LOCATION", source_code="BHILAI NAGAR",
                          target_id=locations["BHL"].id, source_name="Bhilai Nagar"))
    transporters = {}
    for code, name, person, mobile, gst, area, veh, lrp, pay, pickup in TRANSPORTERS:
        t = Transporter(code=code, name=name, contact_person=person, mobile=mobile, gstin=gst,
                        pan=gst[2:12] if gst else None, service_area=area,
                        payment_terms="Cash at dispatch" if pay == "CASH" else "Monthly bill, 30 days",
                        credit_days=None if pay == "CASH" else 30, payment_mode=pay,
                        direct_pickup_charge=Decimal(pickup[0]) if pickup else None,
                        direct_pickup_basis=pickup[1] if pickup else "PER_TRIP",
                        address=f"Transport Nagar, {'Nagpur' if code == 'VDL' else 'Raipur'}",
                        city="Nagpur" if code == "VDL" else "Raipur",
                        state="Maharashtra" if code == "VDL" else "Chhattisgarh")
        db.add(t)
        transporters[code] = t
    for code, name, mobile, basis, rate, notes in AUTOS:
        db.add(OutsideAuto(code=code, name=name, mobile=mobile, charge_basis=basis, rate=Decimal(rate), notes=notes))
    code, name, mobile, veh, trip, box = AUTO_BOTH
    kish = OutsideAuto(code=code, name=name, mobile=mobile, vehicle_no=veh, allowed_methods="BOTH",
                       charge_basis="PER_TRIP", default_trip_charge=Decimal(trip), default_box_rate=Decimal(box),
                       notes="Bigger local deliveries; per trip or per box")
    kish.sync_legacy()
    db.add(kish)
    db.flush()
    from app.models import LocalAutoRate
    db.add(LocalAutoRate(code="KISH-BHL", outside_auto_id=kish.id, location_id=locations["BHL"].id,
                         rate_per_box=Decimal("22.00"), effective_from=date(today.year if today.month >= 4 else today.year - 1, 4, 1),
                         remarks="Bhilai is further"))
    users = {}
    for uname, full, role, cfa_codes, pwd in USERS:
        u = User(username=uname, full_name=full, role=roles[role], password_hash=hash_password(pwd),
                 all_companies=True, all_cfas=cfa_codes is None,
                 email=f"{uname.replace('.', '_')}@example.com")
        if cfa_codes:
            u.cfas = [cfas[c] for c in cfa_codes]
        db.add(u)
        users[uname] = u
    db.flush()
    # Import templates (one per ERP export format)
    templates = [
        ("SAP - Billing documents (Company A)", "SAP", None, {
            "company_code": "Sales Org", "cfa_code": "Plant", "invoice_no": "Billing Document",
            "invoice_date": "Billing Date", "customer_code": "Sold-To Party", "customer_name": "Sold-To Name",
            "location": "Ship-To City", "sales_value": "Net Value", "boxes": "No. of Packages",
            "source_invoice_id": "Billing Document"}, "AUTO"),
        ("MARG - Sales register (Company B)", "MARG", "COB", {
            "cfa_code": "Branch", "invoice_no": "Bill No", "invoice_date": "Bill Date", "customer_code": "Party Code",
            "customer_name": "Party Name", "location": "Station", "sales_value": "Net Amount", "boxes": "Cases",
            "remarks": "Narration"}, "%d-%m-%Y"),
        ("CBO - Invoice export (Company C)", "CBO", None, {
            "company_code": "Division", "cfa_code": "Depot", "invoice_no": "InvoiceNo", "invoice_date": "InvoiceDate",
            "customer_code": "CustomerCode", "customer_name": "CustomerName", "location": "City",
            "sales_value": "NetAmount", "boxes": "NoOfCases", "source_invoice_id": "InvoiceId"}, "%d/%m/%Y"),
        ("Intrect - Item-wise invoice export (Own distribution)", "INTRECT", None, {
            "company_code": "principal", "cfa_code": "branch", "invoice_no": "inv_no", "invoice_date": "inv_date",
            "customer_code": "party_id", "customer_name": "party_name", "location": "delivery_city",
            "sales_value": "net_value", "boxes": "boxes", "source_invoice_id": "doc_id",
            "sku_code": "item_code", "sku_name": "item_name", "qty": "qty", "item_value": "line_value"}, "AUTO"),
        ("Standard CTDMS template", "GENERIC", None, {k: k for k in imp.TARGET_FIELDS}, "AUTO"),
    ]
    for name, erp, co, cmap, dfmt in templates:
        db.add(ImportTemplate(name=name, erp_system=erps[erp], company_id=companies[co].id if co else None,
                              column_map=cmap, date_format=dfmt, created_by_id=users["admin"].id))
    db.flush()
    return Ctx(db, today, today, users, erps, companies, cfas, locations, customers, transporters)


def seed_rates(c: Ctx, fy_from: date, revision_from: date) -> None:
    db, admin = c.db, c.users["admin"]
    for co, cfa, loc, rate in RATES:
        rate_service.create_rate(db, admin, company_id=c.companies[co].id, cfa_id=c.cfas[cfa].id,
                                 location_id=c.locations[loc].id, rate_per_box=Decimal(rate), effective_from=fy_from,
                                 remarks="Annual freight agreement", change_reason="Initial rate card", auto_approve=True)
    # Level 1 - customer-specific: ABC Pharma negotiated ₹7 (vs ₹8) for Company A Raipur -> Bilaspur
    rate_service.create_rate(db, admin, company_id=c.companies["COA"].id, cfa_id=c.cfas["RPR"].id,
                             customer_id=c.customers["C-ABC01"].id, location_id=c.locations["BSP"].id,
                             rate_per_box=Decimal("7.00"), effective_from=fy_from, auto_approve=True,
                             remarks="Customer-specific: volume customer", change_reason="Negotiated volume rate")
    # Level 3 - company-wide location rate: Company B Korba from any CFA
    rate_service.create_rate(db, admin, company_id=c.companies["COB"].id, location_id=c.locations["KRB"].id,
                             rate_per_box=Decimal("11.00"), effective_from=fy_from, auto_approve=True,
                             remarks="Company-wide Korba rate (any CFA)", change_reason="Initial rate card")
    # Level 4 - explicitly configured default for own distribution
    rate_service.create_rate(db, admin, company_id=c.companies["OWN"].id, is_default=True,
                             rate_per_box=Decimal("10.00"), effective_from=fy_from, auto_approve=True,
                             remarks="Own distribution default for unlisted locations", change_reason="Default rate")
    # v1.3 transporter rate cards (rate per case per area) - what gatepasses charge and bills are checked against
    from app.services import transporter_rate_service as trs
    for tcode, areas in TRANSPORTER_RATES.items():
        for loc, rate in areas.items():
            trs.propose(db, admin, transporter_id=c.transporters[tcode].id, location_id=c.locations[loc].id,
                        rate_per_case=Decimal(rate), effective_from=fy_from, reason="Rate card at start of year",
                        auto_approve=True)


def rate_revision(c: Ctx, revision_from: date) -> None:
    """Company A / Raipur CFA / Durg: ₹6.00 -> ₹6.50 from the 1st of this month (maker: CFA manager, checker approves)."""
    db = c.db
    from app.models import FreightRate, RateStatus
    prev = db.scalar(select(FreightRate).where(FreightRate.company_id == c.companies["COA"].id,
                                               FreightRate.cfa_id == c.cfas["RPR"].id,
                                               FreightRate.location_id == c.locations["DRG"].id,
                                               FreightRate.customer_id.is_(None), FreightRate.status == RateStatus.APPROVED))
    clock(revision_from - timedelta(days=3), 11, 5)
    new = rate_service.revise_rate(db, c.users["cfa.raipur"], prev, rate_per_box=Decimal("6.50"),
                                   effective_from=revision_from, change_reason="Diesel price increase - revised by Company A")
    clock(revision_from - timedelta(days=2), 16, 40)
    rate_service.approve_rate(db, c.users["checker"], new, remarks="Verified against Company A circular")


def pending_rate_proposal(c: Ctx) -> None:
    from app.models import FreightRate, RateStatus
    db = c.db
    prev = db.scalar(select(FreightRate).where(FreightRate.company_id == c.companies["COB"].id,
                                               FreightRate.cfa_id == c.cfas["RPR"].id,
                                               FreightRate.location_id == c.locations["NGP"].id,
                                               FreightRate.status == RateStatus.APPROVED))
    nxt = (c.today.replace(day=1) + timedelta(days=32)).replace(day=1)
    clock(c.today, 10, 15)
    rate_service.revise_rate(db, c.users["cfa.raipur"], prev, rate_per_box=Decimal("15.00"), effective_from=nxt,
                             change_reason="Vidarbha Logistics revised Nagpur tariff")


# ----------------------------------------------------------------------------- invoice generation
@dataclass
class Rec:
    erp: str
    company: str
    cfa: str
    customer: str
    location: str
    inv_no: str
    inv_date: date
    sales: Decimal
    boxes: int | None


def _serial(prefix: str, n: int) -> str:
    return f"{prefix}{n:05d}"


def generate_invoices(c: Ctx, start: date, end: date) -> list[Rec]:
    recs: list[Rec] = []
    counters = {"SAP": 90012000, "MARG": 1200, "CBO": 4500, "INTRECT": 7800}
    by_company = {"COA": "SAP", "COB": "MARG", "COC": "CBO", "OWN": "INTRECT"}
    weights = {"COA": 5, "COB": 4, "COC": 3, "OWN": 2}
    d = start
    while d <= end:
        if d.weekday() != 6 and d.day in (4, 14, 24):
            # Anchor invoices for the spec examples: Company A / Raipur CFA to Durg (rate revised this month),
            # Bilaspur (customer-specific ABC Pharma rate + standard rate), Nagpur; Company B / Raipur likewise.
            for co, cu_code in [("COA", "C-XYZ01"), ("COA", "C-ABC01"), ("COA", "C-VDP01"), ("COA", "C-PQR01"),
                                ("COB", "C-ABC01"), ("COB", "C-XYZ01"), ("COB", "C-PQR01")]:
                erp = by_company[co]
                cu = next(x for x in CUSTOMERS if x[0] == cu_code)
                counters[erp] += 1
                n = counters[erp]
                inv_no = {"SAP": str(n), "MARG": f"MB/{d:%y}/{n:05d}"}[erp]
                boxes = 36 if (cu_code == "C-ABC01" and co == "COA") else max(4, int(RND.gauss(24, 8)))
                sales = Decimal(boxes * RND.randint(700, 1800)).quantize(Decimal("1"))
                recs.append(Rec(erp, co, "RPR", cu_code, cu[2], inv_no, d, sales, boxes))
        if d.weekday() != 6:  # no Sunday dispatches
            n_today = RND.randint(2, 4) if d < end - timedelta(days=1) else RND.randint(8, 11)
            for _ in range(n_today):
                co = RND.choices(list(weights), weights=list(weights.values()))[0]
                erp = by_company[co]
                cust_rows = [x for x in CUSTOMERS if co in x[5]]
                cu = RND.choice(cust_rows)
                cfa_codes = next(x[4] for x in COMPANIES if x[0] == co)
                # nearest CFA for the customer's location
                pref = {"BSP": "BSP", "KRB": "BSP", "RGH": "BSP", "DRG": "DRG", "BHL": "DRG", "RJN": "DRG"}.get(cu[2], "RPR")
                cfa = pref if pref in cfa_codes and RND.random() < 0.8 else "RPR"
                if co == "COA" and cfa == "BSP" and cu[2] not in ("BSP", "KRB", "RGH", "RPR"):
                    cfa = "RPR"
                if co == "COA" and cfa == "DRG" and cu[2] not in ("DRG", "BHL", "RJN", "RPR"):
                    cfa = "RPR"
                if co == "COB" and cfa == "BSP" and cu[2] not in ("BSP", "RGH", "RPR", "KRB"):
                    cfa = "RPR"
                if co == "COC" and cfa == "DRG" and cu[2] not in ("DRG", "BHL", "RJN"):
                    cfa = "RPR"
                counters[erp] += RND.randint(1, 3)
                n = counters[erp]
                inv_no = {"SAP": str(n), "MARG": f"MB/{d:%y}/{n:05d}", "CBO": f"CBO/{cfa}/{n:05d}",
                          "INTRECT": f"INV{d:%y}-{n:05d}"}[erp]
                boxes = max(1, int(RND.gauss(22, 11)))
                sales = Decimal(boxes * RND.randint(650, 1900) + RND.randint(0, 999)).quantize(Decimal("1"))
                recs.append(Rec(erp, co, cfa, cu[0], cu[2], inv_no, d, sales, boxes))
        d += timedelta(days=1)
    recs.sort(key=lambda r: (r.inv_date, r.erp, r.inv_no))
    return recs


def _erp_row(r: Rec) -> dict:
    cu = next(x for x in CUSTOMERS if x[0] == r.customer)
    loc_name = next(x[1] for x in LOCATIONS if x[0] == r.location)
    if r.erp == "SAP":
        plant = {"RPR": "1101", "BSP": "1102", "DRG": "1103"}[r.cfa]
        return {"Billing Document": r.inv_no, "Billing Date": r.inv_date, "Sales Org": "1000", "Plant": plant,
                "Sold-To Party": cu[6]["SAP"], "Sold-To Name": cu[1].upper(), "Ship-To City": loc_name.upper(),
                "Net Value": float(r.sales), "Currency": "INR", "No. of Packages": r.boxes}
    if r.erp == "MARG":
        branch = {"RPR": "RAIPUR", "BSP": "BILASPUR"}[r.cfa]
        station = "BHILAI NAGAR" if r.location == "BHL" else loc_name.upper()
        return {"Bill No": r.inv_no, "Bill Date": r.inv_date.strftime("%d-%m-%Y"), "Branch": branch,
                "Party Code": cu[6]["MARG"], "Party Name": cu[1].upper(), "Station": station,
                "Net Amount": f"{r.sales:.2f}", "Cases": "" if r.boxes is None else r.boxes, "Narration": ""}
    if r.erp == "CBO":
        return {"InvoiceId": f"CBO-{r.inv_no.split('/')[-1]}", "InvoiceNo": r.inv_no,
                "InvoiceDate": r.inv_date.strftime("%d/%m/%Y"), "Division": "CLS", "Depot": f"D-{r.cfa}",
                "CustomerCode": cu[6]["CBO"], "CustomerName": cu[1], "City": loc_name, "NetAmount": f"{r.sales:.2f}",
                "NoOfCases": r.boxes}
    return {"doc_id": f"INTR-{r.inv_no}", "inv_no": r.inv_no, "inv_date": r.inv_date.isoformat(),
            "principal": "OWN-DIST", "branch": "RAIPUR-HO", "party_id": cu[6]["INTRECT"], "party_name": cu[1],
            "delivery_city": loc_name, "net_value": float(r.sales), "boxes": r.boxes}


SKUS = [("OWN-PCM650", "Paracetamol 650 mg Tab 15s", 18.5), ("OWN-AZI500", "Azithromycin 500 mg Tab 3s", 62.0),
        ("OWN-PAN40", "Pantoprazole 40 mg Tab 15s", 48.0), ("OWN-CET10", "Cetirizine 10 mg Tab 10s", 12.0),
        ("OWN-ORS", "ORS Sachet 21 g", 19.0), ("OWN-MVT", "Multivitamin Syrup 200 ml", 96.0),
        ("OWN-AMX625", "Amoxicillin + Clav 625 Tab 6s", 118.0), ("OWN-DCL", "Diclofenac Gel 30 g", 64.0)]


def _erp_rows(r: Rec) -> list[dict]:
    """Intrect exports one row per SKU (item-wise register); the others one row per invoice."""
    row = _erp_row(r)
    if r.erp != "INTRECT":
        return [row]
    import zlib
    rnd = random.Random(zlib.crc32(r.inv_no.encode()))
    picks = rnd.sample(SKUS, k=rnd.randint(2, 4))
    weights = [rnd.randint(1, 5) for _ in picks]
    total = float(r.sales)
    out, used = [], 0.0
    for k, ((code, name, price), w) in enumerate(zip(picks, weights)):
        value = round(total * w / sum(weights), 2) if k < len(picks) - 1 else round(total - used, 2)
        used += value
        out.append({**row, "item_code": code, "item_name": name, "qty": max(1, int(value / price)),
                    "line_value": value})
    return out


def _to_bytes(rows: list[dict], kind: str) -> tuple[bytes, str]:
    headers = list(rows[0].keys())
    if kind == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append([r.get(h) for h in headers])
        from openpyxl.styles import Font
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, date):
                    cell.number_format = "dd.mm.yyyy"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), "xlsx"
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=headers)
    w.writeheader()
    for r in rows:
        w.writerow({k: (v.strftime("%d-%m-%Y") if isinstance(v, date) else v) for k, v in r.items()})
    return buf.getvalue().encode("utf-8-sig"), "csv"


FORMAT = {"SAP": "xlsx", "MARG": "csv", "CBO": "csv", "INTRECT": "xlsx"}
TEMPLATE = {"SAP": "SAP - Billing documents (Company A)", "MARG": "MARG - Sales register (Company B)",
            "CBO": "CBO - Invoice export (Company C)",
            "INTRECT": "Intrect - Item-wise invoice export (Own distribution)"}
IMPORTER = {"SAP": "admin", "MARG": "admin", "CBO": "admin", "INTRECT": "admin"}


def import_files(c: Ctx, recs: list[Rec], label: str, when: datetime, extra_rows: dict | None = None) -> list:
    db = c.db
    batches = []
    for erp in ("SAP", "MARG", "CBO", "INTRECT"):
        rows = [x for r in recs if r.erp == erp for x in _erp_rows(r)]
        rows += (extra_rows or {}).get(erp, [])
        if not rows:
            continue
        content, ext = _to_bytes(rows, FORMAT[erp])
        tick(when)
        tpl = db.scalar(select(ImportTemplate).where(ImportTemplate.name == TEMPLATE[erp]))
        user = c.users[IMPORTER[erp]]
        batch, _ = imp.create_batch_from_file(db, user, content, f"{erp}_{label}.{ext}", template_id=tpl.id)
        imp.validate_batch(db, batch, user)
        tick(when + timedelta(minutes=4))
        imp.commit_batch(db, user, batch)
        batches.append(batch)
    db.flush()
    return batches


# ----------------------------------------------------------------------------- dispatch simulation
VEH_IDX: dict[str, int] = {}
LR_SEQ: dict[str, int] = {}


def _operator(c: Ctx, inv: Invoice) -> User:
    return {"RPR": c.users["operator.raipur"], "BSP": c.users["operator.bilaspur"],
            "DRG": c.users["operator.durg"]}[inv.cfa.code]


def _transporter(c: Ctx, inv: Invoice) -> Transporter:
    code = TRANSPORTER_BY_LOCATION[inv.location.code]
    if code in ("AUTO", "SELF"):
        code = "SCM"
    return c.transporters[code]


def _carrier(inv: Invoice) -> str:
    """Who takes this invoice: a transporter code, AUTO (local outside auto) or SELF (customer collects)."""
    loc = inv.location.code
    code = TRANSPORTER_BY_LOCATION.get(loc, "SCM")
    r = RND.random()
    if loc == "RPR":
        return "SELF" if r < 0.15 else ("AUTO" if r < 0.75 else "SCM")
    if loc in ("DRG", "BHL", "RJN") and r < 0.3:
        return "SBT"
    if loc == "BSP" and r < 0.25:
        return "CGR"
    return code


LR_START = {"PGC": 24600, "BJR": 2850, "OSHO": 6500, "ORL": 900, "PRK": 3530, "JBG": 6750, "CGR": 5000,
            "SCM": 5000, "SBT": 5000, "VDL": 5000}
LR_PREFIX = {"CGR": "CGR", "SCM": "SCM", "SBT": "SBT", "VDL": "VDL"}
LR_NEXT: dict[str, int] = {}  # running LR numbers per transporter (reset by run())


def make_gatepass(c: Ctx, invs: list[Invoice], carrier: str, d: date, stage: str = "DISPATCHED",
                  reprint: bool = False) -> DispatchSlip | None:
    db = c.db
    op = _operator(c, invs[0])
    hh = RND.randint(13, 18) if d == invs[0].invoice_date else RND.randint(9, 16)
    mm = RND.choice([0, 10, 15, 25, 30, 40, 45])
    clock(d, hh, mm)
    autos = {a.code: a for a in db.scalars(select(OutsideAuto))}
    req = ds.GatepassRequest(dispatch_date=d, dispatch_time=time(hh, mm),
                             cases={i.id: i.boxes or RND.randint(2, 8) for i in invs})
    if carrier == "SELF":
        req.mode = DispatchMode.SELF
    elif carrier == "AUTO":
        req.mode = DispatchMode.OUTSIDE_AUTO
        req.outside_auto_id = autos["AUTO-RAM"].id
    else:
        req.transporter_id = c.transporters[carrier].id
        req.mode = DispatchMode.DIRECT if carrier == "PRK" and RND.random() < 0.3 else DispatchMode.TRANSPORTER
        if carrier in ("ORL", "PGC", "BJR", "CGR") and RND.random() < 0.35:
            req.pickup_mode = PickupMode.OUTSIDE_AUTO
            req.outside_auto_id = autos["AUTO-SONU"].id
    req.group_cross_company = True  # demonstration data: the grouping is by customer / area, whatever the company
    try:
        slip = ds.create_gatepass(db, op, invs, req)
    except ds.BusinessError:
        return None  # e.g. an invoice with a box change awaiting approval stays in the queue
    if stage == "GENERATED":
        return slip
    clock(d, hh, min(mm + 4, 59))
    ds.print_gatepass(db, op, slip, RND.choice(["PRINT_A5", "PRINT_A5", "PRINT_A4", "PDF_A5"]))
    if reprint:
        clock(d, min(hh + 1, 21), 5)
        ds.print_gatepass(db, op, slip, "PRINT_A5", RND.choice(["Printer jam - first copy torn",
                                                                "Transporter lost the copy",
                                                                "Customer requested duplicate for gate pass"]))
    if stage == "PRINTED":
        return slip
    clock(d, min(hh + 1, 22), 20)
    _pawatis_printed(db, op, slip)
    ds.gate_out(db, op, slip)
    return slip


def dispatch_in_gatepasses(c: Ctx, invs: list[Invoice], stages: list[str] | None = None) -> list[DispatchSlip]:
    """Group the day's invoices per (dispatch date, CFA, carrier) into gatepasses - one per pickup."""
    groups: dict[tuple, list[Invoice]] = {}
    for inv in invs:
        if inv.status not in (InvoiceStatus.PENDING, InvoiceStatus.VERIFIED, InvoiceStatus.EXCEPTION):
            continue
        if inv.exception_code and inv.exception_code != "BOXES_MISSING":
            continue
        same_day = RND.random() < 0.7 or inv.invoice_date >= c.today
        d = inv.invoice_date if same_day else inv.invoice_date + timedelta(days=1)
        if d.weekday() == 6:
            d += timedelta(days=1)
        d = min(d, c.today)
        groups.setdefault((d, inv.cfa_id, _carrier(inv)), []).append(inv)
    slips = []
    for k, ((d, _cfa, carrier), group) in enumerate(sorted(groups.items(), key=lambda x: (x[0][0], x[0][1], x[0][2]))):
        stage = stages[k % len(stages)] if stages else "DISPATCHED"
        if stage == "NONE":
            continue
        slip = make_gatepass(c, group, carrier, d, stage, reprint=RND.random() < 0.04)
        if slip:
            slips.append(slip)
    return slips


def assign_lrs(c: Ctx, slips: list[DispatchSlip], leave_open: float = 0.0) -> None:
    """The transporter's LR / bilty numbers arrive a day or two after dispatch (one per consignee)."""
    db = c.db
    seq = LR_NEXT
    for slip in sorted(slips, key=lambda s: (s.dispatch_date, s.id)):
        if not slip.transporter_id or slip.is_void:
            continue
        tcode = slip.transporter.code
        by_cust: dict[int, list[Invoice]] = {}
        for inv in ds.slip_invoices(db, slip):
            if inv.status != InvoiceStatus.CANCELLED and not inv.lr_no:
                by_cust.setdefault(inv.customer_id, []).append(inv)
        for group in by_cust.values():
            if RND.random() < leave_open:
                continue
            seq[tcode] = seq.get(tcode, 5000) + RND.randint(1, 3)
            lr = f"{LR_PREFIX.get(tcode, '')}{seq[tcode]}"
            when = slip.dispatch_date + timedelta(days=RND.choice([0, 1, 1, 2]))
            clock(min(when, c.today), 11, RND.randint(0, 59))
            ds.set_lr(db, _operator(c, group[0]), group, lr, slip.dispatch_date)


def dispatch_invoice(c: Ctx, inv: Invoice, stop_at: str = "DISPATCHED", reprint: bool = False) -> None:
    """Run the real workflow with a simulated clock. stop_at: PENDING_DETAILS/VERIFIED/SLIP_GENERATED/SLIP_PRINTED/DISPATCHED"""
    db = c.db
    if inv.status not in (InvoiceStatus.PENDING, InvoiceStatus.VERIFIED):
        return
    op = _operator(c, inv)
    same_day = RND.random() < 0.7 or inv.invoice_date >= c.today
    d = inv.invoice_date if same_day else inv.invoice_date + timedelta(days=1)
    if d.weekday() == 6:  # no Sunday dispatch
        d += timedelta(days=1)
    if d > c.today:
        d = c.today
    tr = _transporter(c, inv)
    veh_pool = next(x[6] for x in TRANSPORTERS if x[0] == tr.code)
    VEH_IDX[tr.code] = VEH_IDX.get(tr.code, 0) + 1
    veh = veh_pool[VEH_IDX[tr.code] % len(veh_pool)]
    LR_SEQ[tr.code] = LR_SEQ.get(tr.code, 5000) + RND.randint(1, 4)
    lr = f"{tr.code}{LR_SEQ[tr.code]}"
    hh = RND.randint(13, 18) if d == inv.invoice_date else RND.randint(9, 16)
    mm = RND.choice([0, 10, 15, 25, 30, 40, 45, 50])
    clock(d, hh, mm)
    ds.save_details(db, op, inv, transporter_id=tr.id, vehicle_no=veh, lr_no=lr, dispatch_date=d,
                    dispatch_time=time(hh, mm), remarks=None)
    if stop_at == "PENDING_DETAILS":
        return
    clock(d, hh, mm + 2 if mm < 55 else mm)
    ds.verify(db, op if RND.random() < 0.8 else c.users["checker"], inv)
    if stop_at == "VERIFIED":
        return
    clock(d, hh, min(mm + 4, 59))
    ds.generate_slip(db, op, inv)
    if stop_at == "SLIP_GENERATED":
        return
    ds.record_print(db, op, inv, RND.choice(["PRINT_A5", "PRINT_A5", "PRINT_A4", "PDF_A5"]))
    if reprint:
        clock(d, min(hh + 1, 22), 5)
        ds.record_print(db, op, inv, "PRINT_A5", RND.choice(["Printer jam - first copy torn",
                                                             "Transporter lost the copy",
                                                             "Customer requested duplicate for gate pass"]))
    if stop_at == "SLIP_PRINTED":
        return
    clock(d, min(hh + 1, 22), 20)
    ds.complete(db, op, inv)


# ----------------------------------------------------------------------------- main seed
def _pawatis_printed(db, user, slip):
    """Garage Pawatis travel with the auto: print them before the trip leaves."""
    from app.services import pawati_service as pws
    rows = pws.active(db, slip)
    if rows:
        pws.print_sheet(db, user, rows)


def run(reset: bool = True, quiet: bool = False) -> dict:
    from app.db import SessionLocal
    if reset:
        reset_database()
    now = datetime.now(TZ)
    real_today = now.date()
    # The simulated working day ends at 19:30. Before that, simulate up to yesterday so that no demo event
    # carries a timestamp later than the real current time.
    today = real_today if now.time() >= time(19, 30) else real_today - timedelta(days=1)
    m0 = today.replace(day=1)                                   # this month
    m1 = (m0 - timedelta(days=1)).replace(day=1)                # last month
    m2 = (m1 - timedelta(days=1)).replace(day=1)                # two months ago
    fy_from = date(today.year if today.month >= 4 else today.year - 1, 4, 1)
    db = SessionLocal()
    global _DB
    _DB = db
    LR_NEXT.clear()
    LR_NEXT.update(LR_START)
    log = (lambda *a: None) if quiet else print
    try:
        clock(fy_from if fy_from < m2 else m2, 9, 0)
        c = seed_masters(db, today)
        seed_rates(c, fy_from, m0)
        db.commit()
        log("masters & rates ✓")

        # ---- daily ERP exports imported at 12:30 each working day; rate revision before this month
        all_recs = generate_invoices(c, m2, today)
        d = m2
        revised = False
        while d <= today:
            if not revised and d >= m0 - timedelta(days=3):
                rate_revision(c, m0)
                revised = True
            day_recs = [r for r in all_recs if r.inv_date == d]
            extra = _problem_rows(c, day_recs, [r for r in all_recs if r.inv_date < m0]) if d == today else None
            if day_recs or extra:
                import_files(c, day_recs, f"{d:%Y%m%d}", at(d, 12, 30), extra)
            d += timedelta(days=1)
        if not revised:
            rate_revision(c, m0)
        _qr_captures(c, today)
        db.commit()
        log(f"imports ✓ ({len(all_recs)} invoices)")

        # ---- dispatch simulation: one gatepass per pickup (carrier) per day, LR numbers arriving later
        invs = list(db.scalars(select(Invoice).order_by(Invoice.invoice_date, Invoice.id)))
        recent_cut = today - timedelta(days=3)
        old = [i for i in invs if i.invoice_date < recent_cut]
        recent = [i for i in invs if i.invoice_date >= recent_cut]
        slips = dispatch_in_gatepasses(c, old)
        db.flush()
        assign_lrs(c, slips)
        # recent invoices in every stage so the desk, the gate and the LR-pending list are realistic
        recent_slips = dispatch_in_gatepasses(c, recent, ["DISPATCHED", "PRINTED", "DISPATCHED", "GENERATED", "NONE",
                                                          "DISPATCHED", "NONE"])
        db.flush()
        assign_lrs(c, [s for s in recent_slips if s.gate_out_at], leave_open=0.5)
        db.commit()
        log(f"dispatch workflow ✓ ({len(slips) + len(recent_slips)} gatepasses)")

        _overrides_and_cancellations(c, invs, today)
        db.commit()
        log("overrides / cancellations ✓")

        _bills(c, m2, m1)
        db.commit()
        log("transporter bills ✓")

        pending_rate_proposal(c)
        _deliveries(c, today)
        db.commit()
        log("proofs of delivery ✓")
        # a demo API pull so the integration log is not empty
        from app.integrations.runner import run_sync
        clock(today, 8, 45)
        run_sync(db, c.erps["DEMO"], c.users["admin"], trigger="SCHEDULED",
                 since=datetime.combine(today - timedelta(days=1), time.min))
        # ERP write-back: older dispatches already pulled back into the ERPs
        for inv in db.scalars(select(Invoice).where(Invoice.status == InvoiceStatus.DISPATCHED,
                                                    Invoice.dispatch_date < m1)):
            inv.erp_sync_status = "SYNCED"
            inv.erp_synced_at = inv.completed_at
        db.commit()
        set_fake_now(None)
        if get_settings().environment != "test":
            write_sample_files(c, real_today)
        counts = {s_: db.query(Invoice).filter(Invoice.status == s_).count() for s_ in InvoiceStatus.ALL}
        log("status counts:", counts)
        return counts
    finally:
        set_fake_now(None)
        _DB = None
        db.close()


def _deliveries(c: Ctx, today: date) -> None:
    """Proofs of delivery for most consignments that left more than a day ago: transit 1-6 days, a few short /
    damaged, some confirmed by the customer on the tracking link; the rest stay 'in transit' (some overdue)."""
    from app.services import delivery_service as dsv
    rnd = random.Random(20261015)  # separate stream: the rest of the demo data stays as it was
    invs = list(_DB.scalars(select(Invoice).where(Invoice.status == InvoiceStatus.DISPATCHED,
                                                  Invoice.dispatch_slip_id.isnot(None),
                                                  Invoice.dispatch_date <= today - timedelta(days=2))
                            .order_by(Invoice.dispatch_date, Invoice.id)))
    clock(today, 18, 0)
    for group in dsv.consignments_of(invs):
        first = group[0]
        if rnd.random() > 0.85:
            continue  # still on the road / POD not back yet
        d = min(first.dispatch_date + timedelta(days=rnd.choice([1, 1, 2, 2, 2, 3, 3, 4, 5, 6])), today)
        roll = rnd.random()
        status, note = ("SHORT", "1 case short on arrival") if roll < 0.03 else \
            ("DAMAGED", "one case wet / torn") if roll < 0.06 else ("DELIVERED", None)
        customer = rnd.random() < 0.15 and status == "DELIVERED"
        dsv.record(_DB, None if customer else _operator(c, first), group, status=status, delivered_on=d,
                   received_by=rnd.choice(["Owner", "Store keeper", "Accounts", "Pharmacist", "Godown in-charge"]),
                   note=note, source="CUSTOMER" if customer else "STAFF")


QR_DEMO = [  # (company, customer, invoice no, value, cases) - captured at the desk by scanning the e-invoice QR
    ("COA", "C-UTP01", "QR-SAP-5501", "48210.00", 7), ("COB", "C-SMD01", "QR-MB-5502", "23655.00", 4),
    ("COA", "C-DHR01", "QR-SAP-5503", "112840.00", 18),
]


def demo_qr(c: Ctx, co: str, cu: str, inv_no: str, value, d: date) -> str:
    from app.services.einvoice_service import make_demo_qr
    return make_demo_qr(c.companies[co].gstin, c.customers[cu].gstin, inv_no, d, Decimal(value))


def _qr_captures(c: Ctx, today: date) -> None:
    from app.services import einvoice_service
    for k, (co, cu, no, value, cases) in enumerate(QR_DEMO):
        clock(today, 15, 10 + k * 3)
        einvoice_service.scan(c.db, c.users["operator.raipur"], demo_qr(c, co, cu, no, value, today),
                              cfa_id=c.cfas["RPR"].id, cases=cases)


def _problem_rows(c: Ctx, part: list[Rec], hist: list[Rec]) -> dict:
    """Rows that must fail validation in the latest MARG file: unknown customer, unknown location, duplicate."""
    d = c.today
    dup = next((r for r in hist if r.erp == "MARG"), None)
    rows = [
        {"Bill No": f"MB/{d:%y}/09901", "Bill Date": d.strftime("%d-%m-%Y"), "Branch": "RAIPUR", "Party Code": "NEW777",
         "Party Name": "NAVJEEVAN MEDICAL AGENCIES", "Station": "RAIPUR", "Net Amount": "48250.00", "Cases": 14,
         "Narration": "New party - not yet in master"},
        {"Bill No": f"MB/{d:%y}/09902", "Bill Date": d.strftime("%d-%m-%Y"), "Branch": "RAIPUR", "Party Code": "ABC 001",
         "Party Name": "A.B.C. PHARMA", "Station": "BILASPUR", "Net Amount": "22110.00", "Cases": 9,
         "Narration": "Code typed with a space in MARG"},
        {"Bill No": f"MB/{d:%y}/09903", "Bill Date": d.strftime("%d-%m-%Y"), "Branch": "RAIPUR", "Party Code": "MET005",
         "Party Name": "METRO PHARMA", "Station": "MAHASAMUND", "Net Amount": "15400.00", "Cases": 6,
         "Narration": "Location not in master"},
        {"Bill No": f"MB/{d:%y}/09904", "Bill Date": d.strftime("%d-%m-%Y"), "Branch": "RAIPUR", "Party Code": "SGM006",
         "Party Name": "SHREE GANESH MEDICOS", "Station": "BHILAI NAGAR", "Net Amount": "30980.00", "Cases": "",
         "Narration": "Cases not entered in MARG"},
    ]
    if dup:
        r = _erp_row(dup)
        r["Narration"] = "Re-exported by mistake"
        rows.append(r)
    return {"MARG": rows}


def _overrides_and_cancellations(c: Ctx, invs: list[Invoice], today: date) -> None:
    db = c.db
    open_invs = [i for i in invs if i.status in (InvoiceStatus.PENDING,) and (i.boxes or 0) >= 6]
    # 1) operator requests a box change, checker approves
    if open_invs:
        inv = open_invs[0]
        clock(today, 10, 5)
        ov = invoice_service.request_override(db, c.users["operator.raipur" if inv.cfa.code == "RPR" else
                                                           "operator.bilaspur" if inv.cfa.code == "BSP" else
                                                           "operator.durg"], inv, "BOXES", inv.boxes + 2,
                                              "Physical count at dock: 2 extra boxes (split carton)")
        clock(today, 10, 40)
        invoice_service.decide_override(db, c.users["checker"], ov, True, "Checked physically with warehouse in-charge")
    # 2) pending request - stays in the approval queue (invoice on hold)
    if len(open_invs) > 1:
        inv = open_invs[1]
        clock(today, 11, 20)
        invoice_service.request_override(db, _operator(c, inv), inv, "BOXES", max(1, inv.boxes - 3),
                                         "3 boxes short-shipped from principal; confirm before dispatch")
    # 3) checker applies a special freight rate directly (audited)
    if len(open_invs) > 2:
        inv = open_invs[2]
        clock(today, 12, 5)
        invoice_service.request_override(db, c.users["checker"], inv, "RATE", "5.00",
                                         "Shared vehicle with another consignment - special rate agreed by management")
    # 4) boxes missing (from MARG row) -> operator enters them (initial entry, no approval needed)
    miss = [i for i in invs if i.exception_code == "BOXES_MISSING"]
    for inv in miss[:0]:  # keep it open for the demo
        pass
    # cancellations: one open invoice cancelled (ERP credit note), one historic slip voided & regenerated
    cand = [i for i in invs if i.status in (InvoiceStatus.PENDING, InvoiceStatus.VERIFIED)]
    if len(cand) > 4:
        inv = cand[-1]
        clock(today, 12, 30)
        ds.cancel(db, c.users["cfa.raipur"] if inv.cfa.code == "RPR" else c.users["checker"], inv,
                  "Invoice cancelled in ERP - credit note issued (order returned)")
    slipped = list(db.scalars(select(DispatchSlip).where(DispatchSlip.is_void.is_(False),
                                                         DispatchSlip.print_count == 0,
                                                         DispatchSlip.transporter_id.isnot(None))
                              .order_by(DispatchSlip.id)))
    if slipped:
        slip = slipped[0]
        group = ds.slip_invoices(db, slip)
        clock(today, 13, 10)
        ds.void_gatepass(db, c.users["checker"], slip, "Wrong invoice picked - gatepass made again")
        clock(today, 13, 20)
        ds.create_gatepass(db, _operator(c, group[0]), group, ds.GatepassRequest(
            mode=slip.dispatch_mode, transporter_id=slip.transporter_id, dispatch_date=today,
            cases={i.id: i.boxes for i in group}, remarks="Re-made after correction", group_cross_company=True))
    hist_dispatched = [i for i in invs if i.status == InvoiceStatus.DISPATCHED]
    if hist_dispatched:
        inv = hist_dispatched[len(hist_dispatched) // 3]
        clock(inv.dispatch_date + timedelta(days=2), 15, 0)
        ds.cancel(db, c.users["admin"], inv, "Consignment returned undelivered - customer shop closed (admin cancellation)")


def _bills(c: Ctx, m2: date, m1: date) -> None:
    db = c.db
    acc, chk = c.users["accounts"], c.users["checker"]

    def period(m):
        return m, (m + timedelta(days=32)).replace(day=1) - timedelta(days=1)

    def make(tcode, m, suffix, tweak=None):
        pf, pt = period(m)
        clock(pt + timedelta(days=4), 11, 0)
        b = recon_service.create_bill(db, acc, transporter_id=c.transporters[tcode].id, bill_no=f"{tcode}/{m:%y%m}/{suffix}",
                                      bill_date=pt + timedelta(days=3), bill_amount="0", period_from=pf, period_to=pt,
                                      remarks=f"Monthly freight bill {m:%b %Y}")
        recon_service.link_dispatches(db, acc, b)
        amounts = {}
        for ln in b.lines:
            amt = ln.calculated_amount or Decimal("0")
            if tweak:
                amt = tweak(ln, amt)
            amounts[ln.id] = str(amt)
        recon_service.set_line_amounts(db, acc, b, amounts)
        b.bill_amount = sum((ln.billed_amount or 0) for ln in b.lines)
        recon_service.rematch(db, b)
        return b

    def approve_pay(b, m, pay=True):
        if not b.lines or b.status not in ("MATCHED", "VARIANCE"):
            return
        clock(period(m)[1] + timedelta(days=5), 12, 0)
        recon_service.apply_rate_review(db, acc, b, [], "Rates as per card")
        clock(period(m)[1] + timedelta(days=6), 15, 30)
        recon_service.approve(db, chk, b, "Matched with gatepasses")
        if pay:
            clock(period(m)[1] + timedelta(days=12), 12, 0)
            recon_service.mark_paid(db, acc, b, f"UTR{RND.randint(10**11, 10**12 - 1)}")

    # M-2: every TBB transporter's bill matched, approved and paid
    for tcode in ("CGR", "SBT", "SCM", "PRK", "OSHO", "PGC", "ORL", "BJR"):
        approve_pay(make(tcode, m2, "01"), m2)
    # M-1: some approved, one with variances (awaiting approval), scanned bills with irregularities below
    approve_pay(make("SBT", m1, "01"), m1, pay=False)
    approve_pay(make("PGC", m1, "01"), m1, pay=False)
    b = make("PRK", m1, "01")
    clock(period(m1)[1] + timedelta(days=5), 12, 0)
    recon_service.apply_rate_review(db, acc, b, [], "Parakh rates unchanged")
    k = {"n": 0}

    def overbill(ln, amt):
        k["n"] += 1
        if k["n"] % 4 == 0:
            return amt + Decimal("60.00")  # detention / extra charges billed
        if k["n"] % 7 == 0:
            return amt + Decimal("25.00")
        return amt
    b2 = make("SCM", m1, "01", overbill)
    # a line the transporter billed that the system cannot find
    from app.models import TransporterBillLine
    b2.lines.append(TransporterBillLine(lr_no="SCM99999", invoice_no="UNKNOWN-1", billed_boxes=12,
                                        billed_amount=Decimal("420.00"), remarks="LR not in dispatch records"))
    db.flush()
    b2.bill_amount = sum((ln.billed_amount or 0) for ln in b2.lines)
    recon_service.rematch(db, b2)
    make("CGR", m1, "01")
    make("VDL", m2, "01")
    _scanned_bills(c, m1, period(m1))


def _bill_image(title: str, header: list[str], rows: list[list], footer: str) -> bytes:
    """A plain 'photographed bill' picture for the demo (real bills are phone photos / scans)."""
    from PIL import Image, ImageDraw, ImageFont
    w, rh = 1100, 34
    h = 200 + rh * (len(rows) + 2) + 80
    img = Image.new("RGB", (w, h), (250, 248, 240))
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype(str(BASE_DIR / "app" / "static" / "fonts" / "DejaVuSans.ttf"), 20)
        big = ImageFont.truetype(str(BASE_DIR / "app" / "static" / "fonts" / "DejaVuSans-Bold.ttf"), 30)
    except OSError:
        font = big = ImageFont.load_default()
    d.text((40, 30), title, fill=(30, 30, 90), font=big)
    d.text((40, 80), "TRANSPORTATION BILL", fill=(60, 60, 60), font=font)
    cols = [40, 180, 330, 600, 760, 880, 980]
    y = 140
    for x, txt in zip(cols, header):
        d.text((x, y), txt, fill=(0, 0, 0), font=font)
    d.line((30, y + rh - 4, w - 30, y + rh - 4), fill=(80, 80, 80), width=2)
    for r in rows:
        y += rh
        for x, txt in zip(cols, r):
            d.text((x, y), "" if txt is None else str(txt), fill=(20, 20, 120), font=font)
    y += rh + 10
    d.line((30, y, w - 30, y), fill=(80, 80, 80), width=2)
    d.text((40, y + 15), footer, fill=(0, 0, 0), font=big)
    import io as _io
    buf = _io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _store_scan(content: bytes, label: str) -> list[dict]:
    folder = get_settings().upload_dir / "bills"
    folder.mkdir(parents=True, exist_ok=True)
    name = f"demo-{label}.png"
    (folder / name).write_bytes(content)
    return [{"name": name, "original": f"{label}.jpg", "type": "image/png"}]


def _scanned_bills(c: Ctx, m1: date, per: tuple[date, date]) -> None:
    """Bills as they arrive on paper, scanned and read into lines, showing every kind of irregularity."""
    db = c.db
    acc = c.users["accounts"]
    pf, pt = per

    def consignments(tcode, d_from, d_to, billed=False):
        invs = list(db.scalars(select(Invoice).where(
            Invoice.transporter_id == c.transporters[tcode].id, Invoice.status == InvoiceStatus.DISPATCHED,
            Invoice.dispatch_date >= d_from, Invoice.dispatch_date <= d_to,
            Invoice.transporter_bill_id.isnot(None) if billed else Invoice.transporter_bill_id.is_(None))
            .order_by(Invoice.dispatch_date, Invoice.id)))
        return recon_service.consignments(invs)

    def scanned(tcode, bill_no, rows, title, total=None):
        tr = c.transporters[tcode]
        clock(pt + timedelta(days=3), 10, 30)
        lines_total = sum(Decimal(str(r.get("billed_amount") or 0)) for r in rows)
        b = recon_service.create_bill(db, acc, transporter_id=tr.id, bill_no=bill_no, bill_date=pt + timedelta(days=1),
                                      bill_amount=total if total is not None else lines_total, period_from=pf,
                                      period_to=pt, source="SCAN", remarks=f"Paper bill {m1:%b %Y}, scanned")
        img = _bill_image(title, ["LR No", "Date", "Party", "Station", "Cases", "Rate", "Amount"],
                          [[r.get("lr_no"), r["lr_date"].strftime("%d.%m.%y") if r.get("lr_date") else "",
                            (r.get("party_name") or "")[:18], r.get("station"), r.get("billed_boxes"),
                            r.get("billed_rate"), r.get("billed_amount")] for r in rows],
                          f"TOTAL  Rs. {b.bill_amount:,.2f}")
        b.scan_files = _store_scan(img, f"{tcode}-{m1:%y%m}")
        b.extraction = {"demo": True, "lines": len(rows)}
        recon_service.save_lines(db, acc, b, [{k: ("" if v is None else str(v)) for k, v in r.items()} for r in rows])
        return b

    def row(grp, **over):
        first = grp[0]
        cases = sum(i.boxes or 0 for i in grp)
        rate = first.rate_per_box or Decimal("0")
        r = {"lr_no": first.lr_no, "lr_date": first.lr_date or first.dispatch_date, "party_name": first.customer.name,
             "station": first.location.code if RND.random() < 0.5 else first.location.name, "billed_boxes": cases,
             "billed_rate": rate, "billed_amount": sum((recon_service.transporter_due(db, i) for i in grp), Decimal("0"))}
        r.update(over)
        return r

    # 1) Orient Logistics (like the printed bill no. 79): a rate increase to ₹40 for Raigarh, an LR we did not have,
    #    one consignment billed twice, more cases than dispatched, an arithmetic slip and an LR that is not ours
    grps = [g for g in consignments("ORL", pf, pt)]
    lrless = grps.pop() if len(grps) >= 4 else None
    rows = []
    for k, g in enumerate(grps[:12]):
        first = g[0]
        if first.location.code == "RGH" and k % 2 == 0:
            cases = sum(i.boxes or 0 for i in g)
            rows.append(row(g, billed_rate=Decimal("40.00"), billed_amount=Decimal("40.00") * cases))
        elif k == 3:
            rows.append(row(g, billed_boxes=sum(i.boxes or 0 for i in g) + 2,
                            billed_amount=(sum(i.boxes or 0 for i in g) + 2) * (first.rate_per_box or 0)))
        elif k == 5:
            r = row(g)
            r["billed_amount"] = r["billed_amount"] + Decimal("100")  # 3 x 50 written as 250
            rows.append(r)
        else:
            rows.append(row(g))
    # an LR the transporter gave us but we never typed in: clear ours so the bill line has to be matched
    if lrless:
        lr = lrless[0].lr_no or "999"
        for i in lrless:
            i.lr_no = None
            i.lr_date = None
        db.flush()
        rows.append(row(lrless, lr_no=lr))
    paid = consignments("ORL", pf - timedelta(days=40), pf - timedelta(days=1), billed=True)
    if paid:
        rows.append(row(paid[0]))  # already paid on last month's bill
    rows.append({"lr_no": "1189", "lr_date": pt - timedelta(days=3), "party_name": "Ganga Medical", "station": "Champa",
                 "billed_boxes": 6, "billed_rate": Decimal("50.00"), "billed_amount": Decimal("300.00")})
    if rows:
        scanned("ORL", "79", rows, "ORIENT LOGISTICS SERVICE · Bill No 79")

    # 2) Jai Bhole Garage is paid in cash at every dispatch - their bill must not be paid again
    jbg = consignments("JBG", pf, pt)
    if jbg:
        scanned("JBG", "2615", [row(g) for g in jbg[:5]], "JAI BHOLE GARAGE (Goods Transport) · 2615")

    # 3) Osho Express charged ₹28 instead of ₹25 per case all month -> rate review proposes the new rate
    osho = consignments("OSHO", pf, pt)
    if osho:
        rows = [row(g, billed_rate=Decimal("28.00"), billed_amount=Decimal("28.00") * sum(i.boxes or 0 for i in g))
                for g in osho[:10]]
        scanned("OSHO", "336", rows, "OSHO EXPRESS · Bilaspur - Raipur daily service · 336")
    db.flush()


def write_sample_files(c: Ctx, today: date) -> None:
    """Fresh ERP exports dated today (not yet imported) that users can upload to try the import screen."""
    SAMPLE_DIR.mkdir(exist_ok=True)
    recs = []
    base = {"SAP": 91000000, "MARG": 20000, "CBO": 80000, "INTRECT": 60000}
    plan = [("SAP", "COA", "RPR", "C-ABC01", "BSP", 36), ("SAP", "COA", "RPR", "C-XYZ01", "DRG", 18),
            ("SAP", "COA", "RPR", "C-PQR01", "NGP", 42), ("SAP", "COA", "BSP", "C-SUN01", "KRB", 25),
            ("SAP", "COA", "DRG", "C-SGM01", "BHL", 12), ("SAP", "COA", "RPR", "C-MET01", "RPR", 8),
            ("MARG", "COB", "RPR", "C-ABC01", "BSP", 30), ("MARG", "COB", "RPR", "C-XYZ01", "DRG", 16),
            ("MARG", "COB", "RPR", "C-PQR01", "NGP", 28), ("MARG", "COB", "BSP", "C-VDP01", "BSP", 11),
            ("MARG", "COB", "RPR", "C-SUN01", "KRB", 21),
            ("CBO", "COC", "RPR", "C-SUN01", "KRB", 19), ("CBO", "COC", "RPR", "C-MET01", "RPR", 10),
            ("CBO", "COC", "DRG", "C-JMH01", "RJN", 14),
            ("INTRECT", "OWN", "RPR", "C-CCC01", "RPR", 6), ("INTRECT", "OWN", "RPR", "C-SUN01", "KRB", 17),
            ("INTRECT", "OWN", "RPR", "C-VDP01", "BSP", 9)]
    for i, (erp, co, cfa, cu, loc, bx) in enumerate(plan):
        n = base[erp] + i
        inv_no = {"SAP": str(n), "MARG": f"MB/{today:%y}/{n}", "CBO": f"CBO/{cfa}/{n}", "INTRECT": f"INV{today:%y}-{n}"}[erp]
        recs.append(Rec(erp, co, cfa, cu, loc, inv_no, today, Decimal(bx * 1150 + i * 17), bx))
    names = {"SAP": "SAP_billing_documents_sample", "MARG": "MARG_sales_register_sample",
             "CBO": "CBO_invoice_export_sample", "INTRECT": "Intrect_invoice_export_sample"}
    # e-invoice QR payloads for invoices not in CTDMS yet: paste one into "Scan & Gatepass" (a 2D scanner types the same)
    qr_lines = ["# Demo e-invoice QR codes (GST IRP format, unsigned demo). Paste one line into Scan & Gatepass.", ""]
    for k, (co, cu, val, cases) in enumerate([("COA", "C-SBV01", "36240.00", 6), ("COB", "C-JNH01", "58900.00", 9),
                                             ("COC", "C-HRM01", "17480.00", 3)]):
        no = f"QR{today:%y%m%d}{k + 1:02d}"
        qr_lines += [f"# {c.companies[co].name} -> {c.customers[cu].name}, invoice {no}, Rs {val}, {cases} cases",
                     demo_qr(c, co, cu, no, val, today), ""]
    (SAMPLE_DIR / "einvoice_qr_samples.txt").write_text("\n".join(qr_lines), encoding="utf-8")
    for erp in names:
        rows = [x for r in recs if r.erp == erp for x in _erp_rows(r)]
        if erp == "MARG":
            rows.append({"Bill No": f"MB/{today:%y}/29999", "Bill Date": today.strftime("%d-%m-%Y"), "Branch": "RAIPUR",
                         "Party Code": "NEW888", "Party Name": "SANJIVANI MEDICAL STORES", "Station": "RAIPUR",
                         "Net Amount": "18900.00", "Cases": 7, "Narration": "New party"})
        content, ext = _to_bytes(rows, FORMAT[erp])
        (SAMPLE_DIR / f"{names[erp]}.{ext}").write_bytes(content)
    # standard template
    std = [{"company_code": "COA", "cfa_code": "RPR", "invoice_no": f"STD-{today:%y%m%d}-001",
            "invoice_date": today.strftime("%d-%m-%Y"), "customer_code": "C-ABC01", "customer_name": "ABC Pharma",
            "location": "Bilaspur", "sales_value": "115000.00", "boxes": 36, "source_invoice_id": "", "remarks": ""},
           {"company_code": "COB", "cfa_code": "RPR", "invoice_no": f"STD-{today:%y%m%d}-002",
            "invoice_date": today.strftime("%d-%m-%Y"), "customer_code": "C-PQR01", "customer_name": "PQR Distributors",
            "location": "Nagpur", "sales_value": "88400.00", "boxes": 24, "source_invoice_id": "", "remarks": ""},
           {"company_code": "COC", "cfa_code": "RPR", "invoice_no": f"STD-{today:%y%m%d}-003",
            "invoice_date": today.strftime("%d-%m-%Y"), "customer_code": "C-SUN01", "customer_name": "Sunrise Agencies",
            "location": "Korba", "sales_value": "41250.00", "boxes": 15, "source_invoice_id": "",
            "remarks": "No Company C Korba rate -> RATE MISSING"}]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(std[0]))
    w.writeheader()
    w.writerows(std)
    (SAMPLE_DIR / "standard_invoice_template.csv").write_text(buf.getvalue(), encoding="utf-8")
    # freight rate upload template
    (SAMPLE_DIR / "freight_rates_template.csv").write_text(
        "company,cfa,customer,location,rate_per_box,effective_from,effective_to,is_default,reason\n"
        f"COC,RPR,,Korba,10.50,{today.replace(day=1):%d-%m-%Y},,,Korba rate for Company C\n"
        f"COA,RPR,C-MET01,Raipur,2.50,{today:%d-%m-%Y},,,Customer-specific local rate\n", encoding="utf-8")
    # transporter statement for this month's Vidarbha dispatches (with two variances)
    db = c.db
    invs = list(db.scalars(select(Invoice).where(Invoice.transporter_id == c.transporters["VDL"].id,
                                                 Invoice.status == InvoiceStatus.DISPATCHED,
                                                 Invoice.dispatch_date >= c.today.replace(day=1)).order_by(Invoice.dispatch_date)))
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["LR No", "Invoice No", "Boxes", "Amount", "Remarks"])
    for i, inv in enumerate(invs):
        amt = inv.freight_amount + (Decimal("75.00") if i == 1 else Decimal("0"))
        bx = inv.boxes + (2 if i == 2 else 0)
        w.writerow([inv.lr_no, inv.invoice_no, bx, f"{amt:.2f}", "Loading charges" if i == 1 else ""])
    w.writerow(["VDL77777", "", 10, "1400.00", "Not in our dispatch records"])
    (SAMPLE_DIR / "transporter_bill_sample.csv").write_text(buf.getvalue(), encoding="utf-8")
