"""v1.6 one-time MASTER WORKBOOK upload (first-day set-up).

One cleaned, standardised Excel workbook compiled from MARG / SAP / CBO / Intrect / other ERPs creates the masters:
ERP systems, locations, companies, depots, customer types, customers, transporters (long-distance and local autos),
local-auto box rates, transporter freight rates, invoice code mappings (existing ERP invoice markings) and ERP code
mappings.

    upload -> validate every row -> preview (per sheet: new / already in system / rejected, with reasons)
           -> import the valid rows -> history (+ Excel error report)

Nothing is created silently: a code that already exists is reported and skipped (never overwritten, never duplicated);
a row with a missing mandatory field, a bad value or an unknown reference is rejected with the reason. Later
additions are made one record at a time under Master Settings. Freight rates still go to the checker for approval.
"""
from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import (
    CFA, ChargeBasis, Company, Customer, CustomerTypeMaster, ERPCodeMapping, ERPSystem, InvoiceMarking,
    LocalAutoRate, Location, MasterUpload, OutsideAuto, PaymentMode, Transporter, User, utcnow,
)
from app.services import audit, storage
from app.services import transporter_rate_service as trs
from app.services.errors import BusinessError

MODULE = "MASTER_UPLOAD"
ADAPTERS = {"FILE": "file", "EXCEL": "file", "CSV": "file", "FOLDER": "file_drop", "DROP FOLDER": "file_drop",
            "API": "rest", "REST": "rest", "MARG": "marg", "SAP": "sap_odata", "SAP ODATA": "sap_odata",
            "CBO": "cbo", "DATABASE": "database"}


@dataclass
class Col:
    header: str
    key: str
    required: bool = False
    kind: str = "text"   # text upper int decimal date bool choice list
    choices: tuple = ()
    help: str = ""
    example: str = ""


@dataclass
class SheetSpec:
    name: str
    title: str
    purpose: str
    cols: list[Col]
    key: str = "code"                      # unique key column (None = no uniqueness)
    aliases: tuple = ()
    optional: bool = True


YES = ("YES", "Y", "TRUE", "1", "ACTIVE")
NO = ("NO", "N", "FALSE", "0", "INACTIVE")

SHEETS: list[SheetSpec] = [
    SheetSpec("ERP_Systems", "Source ERP systems", "Billing systems invoices come from (MARG, SAP, CBO, Intrect ...).", [
        Col("ERP Code", "code", True, "upper", help="Short unique code", example="MARG"),
        Col("ERP Name", "name", True, help="Name shown on screens", example="MARG ERP 9+ (Company B)"),
        Col("How data arrives", "connection", False, "upper",
            choices=("FILE", "FOLDER", "API", "MARG", "SAP", "CBO", "DATABASE"),
            help="FILE = Excel/CSV upload (default) · FOLDER = drop folder · API = pushed by the ERP", example="FILE"),
        Col("Active", "active", False, "bool", help="Yes / No (default Yes)", example="Yes"),
        Col("Notes", "notes", help="Status of the link, contact person ...", example="Exported every evening"),
    ], aliases=("ERP", "ERPS", "ERP_SYSTEM")),
    SheetSpec("Locations", "Delivery locations / areas", "Every town / area goods are delivered to. Freight is priced per location.", [
        Col("Location Code", "code", True, "upper", example="RPR-CITY"),
        Col("Location Name", "name", True, example="Raipur City"),
        Col("City", "city", example="Raipur"), Col("State", "state", example="Chhattisgarh"),
        Col("PIN Code", "pincode", example="492001"), Col("Route", "route", example="Local"),
        Col("Zone", "zone", example="Central"), Col("Distance km", "distance_km", kind="decimal", example="0"),
        Col("Active", "active", kind="bool", example="Yes"),
    ], aliases=("LOCATION", "AREAS", "AREA")),
    SheetSpec("Companies", "Companies (principals)", "Every company whose goods are dispatched.", [
        Col("Company Code", "code", True, "upper", help="Existing company code / marking from the ERP", example="ABC"),
        Col("Company Name", "name", True, example="ABC Pharma Ltd"),
        Col("Recognition Code", "marking", kind="upper",
            help="Short code used in the Invoice Recognition Code (blank = the company code)", example="ABC"),
        Col("Source ERP", "erp", kind="upper", help="ERP Code from the ERP_Systems sheet / system", example="SAP"),
        Col("Relationship", "relationship", kind="upper", choices=("PRINCIPAL", "OWN"), example="PRINCIPAL"),
        Col("GSTIN", "gstin", kind="upper"), Col("Contact Person", "contact_person"),
        Col("Contact Number", "contact_phone"), Col("Email", "email"), Col("Address", "address"),
        Col("Active", "active", kind="bool", example="Yes"),
    ], aliases=("COMPANY", "PRINCIPALS")),
    SheetSpec("Depots", "Depots (godowns / dispatch locations)",
              "Each physical godown goods are dispatched from. A city may have several depots - each needs its own code.", [
        Col("Depot Code", "code", True, "upper", help="Unique, also within the same city", example="GUMA"),
        Col("Depot Name", "name", True, example="Guma Depot"),
        Col("Locality / Area", "locality", example="Guma"),
        Col("City", "city", True, example="Raipur"), Col("State", "state", example="Chhattisgarh"),
        Col("PIN Code", "pincode", example="493111"), Col("Address", "address", example="Plot 7, Guma"),
        Col("Recognition Code", "marking", kind="upper", help="Blank = the depot code", example="GUMA"),
        Col("Companies Handled", "companies", kind="list", help="Company codes separated by commas", example="ABC, XYZ"),
        Col("Source ERP", "erp", kind="upper", example="MARG"),
        Col("GSTIN", "gstin", kind="upper"), Col("Contact Person", "contact_person"), Col("Contact Number", "contact_phone"),
        Col("Active", "active", kind="bool", example="Yes"),
    ], aliases=("DEPOT", "CFA", "CFAS", "GODOWNS")),
    SheetSpec("Customer_Types", "Customer types", "Retailer, Stockist, Distributor, Hospital ... with their short codes.", [
        Col("Type Key", "code", True, "upper", help="Used in the Customers sheet", example="RETAILER"),
        Col("Type Name", "name", True, example="Retailer"),
        Col("Recognition Code", "marking", kind="upper", help="Short code for the Invoice Recognition Code", example="RET"),
        Col("Description", "description"), Col("Active", "active", kind="bool", example="Yes"),
    ], aliases=("CUSTOMER TYPES", "CUSTOMERTYPES", "TYPES")),
    SheetSpec("Customers", "Customers", "Stockists, retailers, hospitals ... (one row per customer).", [
        Col("Customer Code", "code", True, "upper", help="Central code", example="C025"),
        Col("Customer Name", "name", True, example="Sharma Medical Store"),
        Col("Customer Type", "ctype", True, "upper", help="Type Key or Type Name from Customer_Types", example="RETAILER"),
        Col("Recognition Code", "marking", kind="upper", help="Blank = the customer code", example="C025"),
        Col("Billing Address", "billing_address"), Col("Delivery Address", "delivery_address"),
        Col("City", "city", example="Raipur"), Col("State", "state", example="Chhattisgarh"),
        Col("PIN Code", "pincode"),
        Col("Delivery Location", "location", True, help="Location Code or Location Name", example="RPR-CITY"),
        Col("Companies", "companies", kind="list", help="Company codes, comma separated", example="ABC"),
        Col("Depots", "depots", kind="list", help="Depot codes, comma separated", example="GUMA"),
        Col("Source ERP", "erp", kind="upper", example="MARG"),
        Col("Source ERP Customer Code", "erp_code", kind="upper", help="What that ERP calls the customer", example="SH025"),
        Col("Contact Person", "contact_person"), Col("Mobile", "contact_phone"), Col("WhatsApp", "whatsapp"),
        Col("Email", "email"), Col("GSTIN", "gstin", kind="upper"), Col("Drug Licence No", "drug_license_no"),
        Col("Active", "active", kind="bool", example="Yes"),
    ], aliases=("CUSTOMER", "PARTIES")),
    SheetSpec("Transporters", "Transporters (long-distance and local autos)",
              "Both kinds in one list: LONG_DISTANCE (rate card, LR, bill) or LOCAL_AUTO (per trip / per box).", [
        Col("Transporter Code", "code", True, "upper", example="SRL"),
        Col("Transporter Name", "name", True, example="Shree Roadlines"),
        Col("Transporter Type", "kind", True, "upper", choices=("LONG_DISTANCE", "LOCAL_AUTO"), example="LONG_DISTANCE"),
        Col("Payment Mode", "payment", kind="upper", choices=("TBB", "CASH"),
            help="Long-distance: TBB = bills later (default) · CASH = paid at dispatch", example="TBB"),
        Col("Charging Method", "methods", kind="upper", choices=("PER_TRIP", "PER_BOX", "BOTH"),
            help="Local auto only", example="PER_TRIP"),
        Col("Usual Method", "usual", kind="upper", choices=("PER_TRIP", "PER_BOX"), help="Local auto with BOTH"),
        Col("Usual Trip Charge", "trip", kind="decimal", help="Local auto per trip: only a suggestion", example="500"),
        Col("Default Rate per Box", "box", kind="decimal", help="Local auto per box", example="20"),
        Col("Vehicle No", "vehicle_no", kind="upper"), Col("Contact Person", "contact_person"), Col("Mobile", "mobile"),
        Col("GSTIN", "gstin", kind="upper"), Col("PAN", "pan", kind="upper"), Col("City", "city"), Col("State", "state"),
        Col("Address", "address"), Col("Areas Served", "service_area"),
        Col("Direct Pickup Charge", "pickup", kind="decimal", help="Long-distance: extra when they collect from our depot"),
        Col("Direct Pickup Basis", "pickup_basis", kind="upper", choices=("PER_TRIP", "PER_BOX")),
        Col("Active", "active", kind="bool", example="Yes"),
    ], aliases=("TRANSPORTER", "CARRIERS")),
    SheetSpec("Local_Auto_Rates", "Local auto box rates", "Per-box rates of local autos (blank location / company / depot = all).", [
        Col("Rate Code", "code", True, "upper", example="LA-B-RPR"),
        Col("Local Auto Code", "auto", True, "upper", help="Transporter Code of a LOCAL_AUTO", example="LA-B"),
        Col("Location", "location", help="Location Code or Name (blank = all)"),
        Col("Company", "company", kind="upper", help="Company Code (blank = all)"),
        Col("Depot", "depot", kind="upper", help="Depot Code (blank = all)"),
        Col("Rate per Box", "rate", True, "decimal", example="20"),
        Col("Effective From", "from", True, "date", example="01-10-2026"),
        Col("Effective To", "to", kind="date"), Col("Remarks", "remarks"),
    ], aliases=("LOCAL_TRANSPORT_RATES", "LOCAL AUTO RATES", "AUTO RATES")),
    SheetSpec("Freight_Rates", "Transporter freight rates (rate cards)",
              "Rate per case per location for long-distance transporters. Sent to the checker for approval.", [
        Col("Transporter Code", "transporter", True, "upper", example="SRL"),
        Col("Location", "location", True, help="Location Code or Name", example="BSP-CITY"),
        Col("Rate per Case", "rate", True, "decimal", example="22"),
        Col("Effective From", "from", True, "date", example="01-10-2026"),
    ], key=None, aliases=("FREIGHT RATES", "RATE_CARDS", "RATES")),
    SheetSpec("Invoice_Code_Mappings", "Invoice code mappings (existing ERP invoice markings)",
              "Markings already inside ERP invoice numbers that show the company (and optionally depot / customer type).", [
        Col("Mapping Code", "code", True, "upper", example="MARG-XYZ"),
        Col("Source ERP", "erp", kind="upper", help="Blank = any ERP", example="MARG"),
        Col("Marking", "marking", True, "upper", help="e.g. XY/ at the start of the ERP invoice number", example="XY/"),
        Col("Where", "match", kind="upper", choices=("PREFIX", "CONTAINS"), example="PREFIX"),
        Col("Company", "company", True, "upper", example="XYZ"),
        Col("Depot", "depot", kind="upper"), Col("Customer Type", "ctype", kind="upper"), Col("Notes", "notes"),
    ], aliases=("INVOICE CODE MAPPINGS", "INVOICE_MARKINGS", "MARKINGS")),
    SheetSpec("ERP_Mappings", "ERP code mappings", "Codes an ERP uses for a central record (other than the customer's own row).", [
        Col("Source ERP", "erp", True, "upper", example="CBO"),
        Col("Entity", "entity", True, "upper", choices=("CUSTOMER", "COMPANY", "DEPOT", "LOCATION", "TRANSPORTER"),
            example="CUSTOMER"),
        Col("ERP Code", "source_code", True, "upper", example="CB-0025"),
        Col("Name in ERP", "source_name"),
        Col("Central Code", "target", True, "upper", help="Code of the record in this system", example="C025"),
    ], key=None, aliases=("ERP MAPPINGS", "MAPPINGS")),
]
ORDER = [s.name for s in SHEETS]
SPEC = {s.name: s for s in SHEETS}


def _norm(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


# ----------------------------------------------------------------------------- template
def template_workbook() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="1F4E79")
    ws.append(["CTDMS v1.6 - INITIAL MASTER WORKBOOK"])
    ws["A1"].font = Font(bold=True, size=14)
    for line in [
        "Fill this workbook ONCE at the first-day set-up with the cleaned master data compiled from all ERPs.",
        "Upload it under Master Settings > Initial Master Upload. Nothing is saved until you have seen the preview.",
        "Rules: one row per record; codes must be unique; columns marked * are mandatory; leave a sheet empty if not needed.",
        "A code that already exists in CTDMS is skipped (never duplicated or overwritten). Later additions: Master Settings.",
        "Freight rates go to the checker for approval. Dates as DD-MM-YYYY. Lists (companies, depots) comma separated.",
        "Product / SKU sheets are NOT part of v1.6 (postponed until ERP integration).",
        ""]:
        ws.append([line])
    ws.append(["Sheet", "Column", "Mandatory", "What it means", "Allowed values", "Example"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
    for sp in SHEETS:
        ws.append([sp.name, "", "", sp.purpose, "", ""])
        ws.cell(ws.max_row, 1).font = bold
        for col in sp.cols:
            ws.append(["", col.header, "Yes" if col.required else "", col.help, ", ".join(col.choices) or
                       {"bool": "Yes / No", "date": "DD-MM-YYYY", "decimal": "number", "int": "whole number",
                        "list": "codes, comma separated"}.get(col.kind, ""), col.example])
    for letter, w in zip("ABCDEF", (24, 26, 11, 70, 34, 24)):
        ws.column_dimensions[letter].width = w
    for row in ws.iter_rows(min_row=9):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for sp in SHEETS:
        s = wb.create_sheet(sp.name)
        s.append([c.header + (" *" if c.required else "") for c in sp.cols])
        for i, c in enumerate(sp.cols, 1):
            cell = s.cell(1, i)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            s.column_dimensions[cell.column_letter].width = max(14, len(c.header) + 4)
            if c.choices:
                dv = DataValidation(type="list", formula1='"' + ",".join(c.choices) + '"', allow_blank=True)
                s.add_data_validation(dv)
                dv.add(f"{cell.column_letter}2:{cell.column_letter}5000")
        s.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------- reading
def read_workbook(content: bytes, filename: str) -> dict[str, list[dict]]:
    """{sheet name: [ {"_row": n, key: raw value} ]} for the sheets this workbook has."""
    if not (filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise BusinessError("Upload the master workbook as an Excel file (.xlsx)")
    if len(content) > get_settings().max_upload_mb * 1024 * 1024:
        raise BusinessError("File too large")
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise BusinessError("The file could not be read as an Excel workbook - open it in Excel and save as .xlsx") from e
    by_norm = {}
    for sp in SHEETS:
        for n in (sp.name, *sp.aliases):
            by_norm[_norm(n)] = sp
    out: dict[str, list[dict]] = {}
    for ws in wb.worksheets:
        sp = by_norm.get(_norm(ws.title))
        if not sp:
            continue
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [_norm(str(h or "").replace("*", "")) for h in rows[0]]
        idx = {}
        for c in sp.cols:
            for i, h in enumerate(headers):
                if h and h == _norm(c.header) and c.key not in idx:
                    idx[c.key] = i
        recs = []
        for n, r in enumerate(rows[1:], start=2):
            if r is None or all(v is None or str(v).strip() == "" for v in r):
                continue
            rec = {"_row": n}
            for c in sp.cols:
                i = idx.get(c.key)
                v = r[i] if i is not None and i < len(r) else None
                if isinstance(v, datetime):
                    v = v.date()
                if isinstance(v, date):
                    v = v.isoformat()
                elif isinstance(v, float) and v.is_integer():
                    v = int(v)
                rec[c.key] = None if v is None or str(v).strip() == "" else str(v).strip()
            recs.append(rec)
        out[sp.name] = recs
        out.setdefault("_missing_columns", {})
        miss = [c.header for c in sp.cols if c.required and c.key not in idx]
        if miss:
            out["_missing_columns"][sp.name] = miss
    if not any(k for k in out if not k.startswith("_")):
        raise BusinessError("No master sheet found. Use the template: Companies, Depots, Customers, Transporters ...")
    return out


# ----------------------------------------------------------------------------- validation
@dataclass
class Known:
    """Codes that exist in the system or will be created by earlier sheets of this workbook."""
    sets: dict = field(default_factory=dict)
    names: dict = field(default_factory=dict)  # location / customer type name -> code

    def has(self, ent: str, code: str | None) -> bool:
        return bool(code) and code.upper() in self.sets.get(ent, set())

    def add(self, ent: str, code: str, name: str | None = None):
        self.sets.setdefault(ent, set()).add(code.upper())
        if name:
            self.names.setdefault(ent, {})[name.upper()] = code.upper()

    def resolve(self, ent: str, value: str | None) -> str | None:
        if not value:
            return None
        v = value.strip().upper()
        if v in self.sets.get(ent, set()):
            return v
        return self.names.get(ent, {}).get(v)


def _load_known(db: Session) -> Known:
    k = Known()
    for ent, model in (("erp", ERPSystem), ("location", Location), ("company", Company), ("depot", CFA),
                       ("ctype", CustomerTypeMaster), ("customer", Customer), ("transporter", Transporter),
                       ("auto", OutsideAuto), ("auto_rate", LocalAutoRate), ("marking", InvoiceMarking)):
        name_col = model.__table__.c.get("name")
        cols = [model.code, name_col] if name_col is not None else [model.code]
        for row in db.execute(select(*cols)):
            k.add(ent, row[0], row[1] if len(row) > 1 else None)
    return k


ENTITY_OF = {"ERP_Systems": "erp", "Locations": "location", "Companies": "company", "Depots": "depot",
             "Customer_Types": "ctype", "Customers": "customer", "Local_Auto_Rates": "auto_rate",
             "Invoice_Code_Mappings": "marking"}


def _typed(col: Col, raw: str | None, errors: list[str]):
    if raw is None:
        return None
    if col.kind == "upper":
        v = raw.upper()
    elif col.kind in ("decimal", "int"):
        try:
            v = Decimal(raw.replace(",", "").replace("₹", ""))
        except InvalidOperation:
            errors.append(f"{col.header}: '{raw}' is not a number")
            return None
        if v < 0:
            errors.append(f"{col.header} cannot be negative")
        if col.kind == "int":
            v = int(v)
    elif col.kind == "date":
        from app.services.utils import parse_date
        try:
            v = parse_date(raw)
        except ValueError:
            errors.append(f"{col.header}: '{raw}' is not a date (DD-MM-YYYY)")
            return None
    elif col.kind == "bool":
        if raw.upper() in YES:
            v = True
        elif raw.upper() in NO:
            v = False
        else:
            errors.append(f"{col.header}: write Yes or No")
            return None
    elif col.kind == "list":
        v = [x.strip().upper() for x in re.split(r"[,;/]", raw) if x.strip()]
    else:
        v = raw
    if col.choices and v not in col.choices:
        errors.append(f"{col.header}: '{raw}' - use one of {', '.join(col.choices)}")
    return v


def validate(db: Session, sheets: dict[str, list[dict]]) -> tuple[list[dict], dict]:
    """Every row with status NEW / EXISTS / REJECTED / DUPLICATE and its messages, plus counts per sheet."""
    known = _load_known(db)
    out: list[dict] = []
    counts: dict = {}
    missing_cols = sheets.get("_missing_columns", {})
    for sname in ORDER:
        if sname not in sheets:
            continue
        sp = SPEC[sname]
        ent = ENTITY_OF.get(sname)
        seen: set = set()
        c = counts.setdefault(sname, {"rows": 0, "new": 0, "exists": 0, "rejected": 0, "warnings": 0})
        for rec in sheets[sname]:
            c["rows"] += 1
            errors: list[str] = []
            warnings: list[str] = []
            if sname in missing_cols:
                errors.append("Column(s) missing in this sheet: " + ", ".join(missing_cols[sname]))
            vals = {}
            for col in sp.cols:
                raw = rec.get(col.key)
                if col.required and raw is None:
                    errors.append(f"{col.header} is mandatory")
                vals[col.key] = _typed(col, raw, errors)
            key = (vals.get(sp.key) if sp.key else None)
            status = "NEW"
            if sname == "Freight_Rates":
                key = f"{vals.get('transporter')}/{(vals.get('location') or '').upper()}/{vals.get('from')}"
            elif sname == "ERP_Mappings":
                key = f"{vals.get('erp')}/{vals.get('entity')}/{vals.get('source_code')}"
            if key:
                if key in seen:
                    errors.append(f"{key} appears more than once in this sheet")
                    status = "DUPLICATE"
                seen.add(key)
            # references to other masters
            _check_refs(db, sname, vals, known, errors, warnings)
            already = (known.has(ent, key) if ent else
                       (known.has("transporter", key) or known.has("auto", key)) if sname == "Transporters" else False)
            if key and not errors and already:
                status = "EXISTS"
                warnings.append(f"{key} already exists in CTDMS - skipped (edit it under Master Settings if it changed)")
            if errors and status != "DUPLICATE":
                status = "REJECTED"
            if status == "NEW" and key:
                if ent:
                    known.add(ent, key, vals.get("name"))
                elif sname == "Transporters":
                    known.add("auto" if vals.get("kind") == "LOCAL_AUTO" else "transporter", key, vals.get("name"))
            c[{"NEW": "new", "EXISTS": "exists"}.get(status, "rejected")] += 1
            if warnings:
                c["warnings"] += 1
            out.append({"sheet": sname, "row": rec["_row"], "key": key, "status": status, "errors": errors,
                        "warnings": warnings, "values": {k: v for k, v in rec.items() if k != "_row"}})
    return out, counts


def _check_refs(db, sname, v, known: Known, errors, warnings):
    def need(ent, val, label, required=False):
        if not val:
            if required:
                errors.append(f"{label} is mandatory")
            return None
        code = known.resolve(ent, val)
        if not code:
            errors.append(f"{label} '{val}' not found (not in this workbook and not in CTDMS)")
        return code

    if sname == "Companies":
        need("erp", v.get("erp"), "Source ERP")
    elif sname == "Depots":
        need("erp", v.get("erp"), "Source ERP")
        for co in v.get("companies") or []:
            need("company", co, "Company")
    elif sname == "Customers":
        need("ctype", v.get("ctype"), "Customer Type", True)
        need("location", v.get("location"), "Delivery Location", True)
        need("erp", v.get("erp"), "Source ERP")
        for co in v.get("companies") or []:
            need("company", co, "Company")
        for d in v.get("depots") or []:
            need("depot", d, "Depot")
        if v.get("erp_code") and not v.get("erp"):
            errors.append("Source ERP Customer Code given without Source ERP")
        if v.get("gstin") and len(v["gstin"]) != 15:
            warnings.append(f"GSTIN {v['gstin']} is not 15 characters - check it")
    elif sname == "Transporters":
        kind = v.get("kind")
        if kind == "LOCAL_AUTO":
            m = v.get("methods") or "PER_TRIP"
            if m in ("PER_BOX", "BOTH") and not v.get("box"):
                warnings.append("No default rate per box - add Local_Auto_Rates rows or the rate later")
            if m == "BOTH" and v.get("usual") not in (None, "PER_TRIP", "PER_BOX"):
                errors.append("Usual Method must be PER_TRIP or PER_BOX")
            if v.get("payment") and v.get("payment") != "CASH":
                warnings.append("Local autos are paid at dispatch - Payment Mode ignored")
        elif kind == "LONG_DISTANCE" and v.get("methods"):
            warnings.append("Charging Method applies to local autos only - ignored for a long-distance transporter")
    elif sname == "Local_Auto_Rates":
        need("auto", v.get("auto"), "Local Auto Code", True)
        need("location", v.get("location"), "Location")
        need("company", v.get("company"), "Company")
        need("depot", v.get("depot"), "Depot")
        if v.get("from") and v.get("to") and v["to"] < v["from"]:
            errors.append("Effective To is before Effective From")
        if v.get("rate") is not None and v["rate"] <= 0:
            errors.append("Rate per Box must be more than 0")
    elif sname == "Freight_Rates":
        code = known.resolve("transporter", v.get("transporter"))
        if v.get("transporter") and not code:
            if known.resolve("auto", v.get("transporter")):
                errors.append(f"{v['transporter']} is a local auto - its rates go in Local_Auto_Rates")
            else:
                errors.append(f"Transporter '{v['transporter']}' not found (not in this workbook and not in CTDMS)")
        loc = need("location", v.get("location"), "Location", True)
        if v.get("rate") is not None and v["rate"] <= 0:
            errors.append("Rate per Case must be more than 0")
        if code and loc and v.get("from") and not errors:
            from app.models import TransporterRate
            dup = db.scalar(select(TransporterRate.id).join(Transporter, Transporter.id == TransporterRate.transporter_id)
                            .join(Location, Location.id == TransporterRate.location_id)
                            .where(func.upper(Transporter.code) == code, func.upper(Location.code) == loc,
                                   TransporterRate.effective_from == v["from"],
                                   TransporterRate.status.in_(("APPROVED", "PENDING_APPROVAL"))))
            if dup:
                errors.append("This transporter already has a rate for this location from the same date")
    elif sname == "Invoice_Code_Mappings":
        need("erp", v.get("erp"), "Source ERP")
        need("company", v.get("company"), "Company", True)
        need("depot", v.get("depot"), "Depot")
        need("ctype", v.get("ctype"), "Customer Type")
    elif sname == "ERP_Mappings":
        need("erp", v.get("erp"), "Source ERP", True)
        ent = {"CUSTOMER": "customer", "COMPANY": "company", "DEPOT": "depot", "LOCATION": "location",
               "TRANSPORTER": "transporter"}.get(v.get("entity") or "")
        if ent:
            need(ent, v.get("target"), "Central Code", True)


# ----------------------------------------------------------------------------- upload / commit
def preview(db: Session, user: User, content: bytes, filename: str) -> MasterUpload:
    sheets = read_workbook(content, filename)
    rows, counts = validate(db, sheets)
    digest = hashlib.sha256(content).hexdigest()
    folder = storage.upload_dir() / "masters"
    folder.mkdir(parents=True, exist_ok=True)
    up = MasterUpload(file_name=Path(filename).name[:200], file_hash=digest, status="PREVIEW", counts=counts,
                      rows=rows, uploaded_by_id=user.id, uploaded_at=utcnow())
    db.add(up)
    db.flush()
    path = folder / f"MASTER-{up.id:04d}.xlsx"
    path.write_bytes(content)
    up.file_path = str(path)
    audit.log(db, user, "MASTER_UPLOAD_PREVIEW", MODULE, record_type="MasterUpload", record_id=up.id,
              ref=up.file_name, details={k: v for k, v in counts.items()})
    return up


def totals(up: MasterUpload) -> dict:
    t = {"rows": 0, "new": 0, "exists": 0, "rejected": 0, "warnings": 0}
    for c in (up.counts or {}).values():
        for k in t:
            t[k] += c.get(k, 0)
    return t


def commit(db: Session, user: User, up: MasterUpload) -> dict:
    if up.status != "PREVIEW":
        raise BusinessError("This upload has already been imported or cancelled")
    # re-read the stored workbook and re-validate against the database as it is now
    content = Path(up.file_path).read_bytes()
    rows, counts = validate(db, read_workbook(content, up.file_name))
    created = {s: 0 for s in ORDER}
    for r in rows:
        if r["status"] != "NEW":
            continue
        _create(db, user, r["sheet"], _values(r), up)
        r["status"] = "IMPORTED"
        created[r["sheet"]] += 1
        db.flush()
    up.rows = rows
    up.counts = {s: {**c, "imported": created.get(s, 0)} for s, c in counts.items()}
    up.status = "IMPORTED"
    up.committed_at = utcnow()
    up.committed_by_id = user.id
    audit.log(db, user, "MASTER_UPLOAD_IMPORTED", MODULE, record_type="MasterUpload", record_id=up.id,
              ref=up.file_name, details={"created": {k: v for k, v in created.items() if v}})
    return created


def _values(r: dict) -> dict:
    sp = SPEC[r["sheet"]]
    errs: list = []
    return {c.key: _typed(c, r["values"].get(c.key), errs) for c in sp.cols}


def _by_code(db, model, code):
    if not code:
        return None
    return db.scalar(select(model).where(func.upper(model.code) == code.upper()))


def _by_code_or_name(db, model, value):
    if not value:
        return None
    return _by_code(db, model, value) or db.scalar(select(model).where(func.upper(model.name) == value.strip().upper()))


def _act(v) -> bool:
    return True if v.get("active") is None else bool(v["active"])


def _create(db: Session, user: User, sheet: str, v: dict, up: MasterUpload) -> None:
    ref = f"master upload #{up.id}"
    if sheet == "ERP_Systems":
        obj = ERPSystem(code=v["code"], name=v["name"], adapter=ADAPTERS.get(v.get("connection") or "FILE", "file"),
                        is_active=_act(v), notes=v.get("notes"),
                        integration_mode="API" if (v.get("connection") or "") in ("API", "REST") else "FILE")
    elif sheet == "Locations":
        obj = Location(code=v["code"], name=v["name"], city=v.get("city"), state=v.get("state"),
                       pincode=v.get("pincode"), route=v.get("route"), zone=v.get("zone"),
                       distance_km=v.get("distance_km"), is_active=_act(v))
    elif sheet == "Companies":
        erp = _by_code_or_name(db, ERPSystem, v.get("erp"))
        obj = Company(code=v["code"], name=v["name"], marking=v.get("marking"), erp_system_id=erp.id if erp else None,
                      relationship_type=v.get("relationship") or "PRINCIPAL", gstin=v.get("gstin"),
                      contact_person=v.get("contact_person"), contact_phone=v.get("contact_phone"),
                      email=v.get("email"), address=v.get("address"), is_active=_act(v))
    elif sheet == "Depots":
        erp = _by_code_or_name(db, ERPSystem, v.get("erp"))
        obj = CFA(code=v["code"], name=v["name"], locality=v.get("locality"), city=v.get("city"),
                  state=v.get("state"), pincode=v.get("pincode"), address=v.get("address"), marking=v.get("marking"),
                  default_erp_id=erp.id if erp else None, gstin=v.get("gstin"),
                  contact_person=v.get("contact_person"), contact_phone=v.get("contact_phone"), is_active=_act(v))
        obj.companies = [c for c in (_by_code(db, Company, x) for x in v.get("companies") or []) if c]
    elif sheet == "Customer_Types":
        n = (db.scalar(select(func.max(CustomerTypeMaster.sort_order))) or 0) + 1
        obj = CustomerTypeMaster(code=v["code"], name=v["name"], marking=v.get("marking"),
                                 description=v.get("description"), sort_order=n, is_active=_act(v))
    elif sheet == "Customers":
        ct = _by_code_or_name(db, CustomerTypeMaster, v.get("ctype"))
        loc = _by_code_or_name(db, Location, v.get("location"))
        erp = _by_code_or_name(db, ERPSystem, v.get("erp"))
        obj = Customer(code=v["code"], name=v["name"], party_type=ct.code, marking=v.get("marking"),
                       billing_address=v.get("billing_address"), delivery_address=v.get("delivery_address"),
                       city=v.get("city"), state=v.get("state"), pincode=v.get("pincode"), location_id=loc.id,
                       contact_person=v.get("contact_person"), contact_phone=v.get("contact_phone"),
                       whatsapp=v.get("whatsapp"), email=v.get("email"), gstin=v.get("gstin"),
                       drug_license_no=v.get("drug_license_no"), source_erp_id=erp.id if erp else None,
                       source_erp_code=v.get("erp_code"), is_active=_act(v))
        obj.companies = [c for c in (_by_code(db, Company, x) for x in v.get("companies") or []) if c]
        obj.cfas = [c for c in (_by_code(db, CFA, x) for x in v.get("depots") or []) if c]
        db.add(obj)
        db.flush()
        if erp and v.get("erp_code"):
            db.add(ERPCodeMapping(erp_system_id=erp.id, entity_type="CUSTOMER", source_code=v["erp_code"],
                                  source_name=v["name"], target_id=obj.id, created_by_id=user.id))
    elif sheet == "Transporters":
        if v["kind"] == "LOCAL_AUTO":
            m = {"PER_BOX": ChargeBasis.PER_CASE, "BOTH": ChargeBasis.BOTH}.get(v.get("methods") or "", ChargeBasis.PER_TRIP)
            usual = {"PER_BOX": ChargeBasis.PER_CASE, "PER_TRIP": ChargeBasis.PER_TRIP}.get(v.get("usual") or "")
            obj = OutsideAuto(code=v["code"], name=v["name"], mobile=v.get("mobile"), vehicle_no=v.get("vehicle_no"),
                              allowed_methods=m, charge_basis=usual or (ChargeBasis.PER_CASE if m == ChargeBasis.PER_CASE
                                                                       else ChargeBasis.PER_TRIP),
                              default_trip_charge=v.get("trip"), default_box_rate=v.get("box"),
                              notes=v.get("address"), is_active=_act(v))
            obj.sync_legacy()
        else:
            obj = Transporter(code=v["code"], name=v["name"], payment_mode=v.get("payment") or PaymentMode.TBB,
                              contact_person=v.get("contact_person"), mobile=v.get("mobile"), gstin=v.get("gstin"),
                              pan=v.get("pan"), city=v.get("city"), state=v.get("state"), address=v.get("address"),
                              service_area=v.get("service_area"), direct_pickup_charge=v.get("pickup"),
                              direct_pickup_basis=ChargeBasis.PER_CASE if v.get("pickup_basis") == "PER_BOX"
                              else ChargeBasis.PER_TRIP, is_active=_act(v))
    elif sheet == "Local_Auto_Rates":
        auto = _by_code(db, OutsideAuto, v["auto"])
        loc = _by_code_or_name(db, Location, v.get("location"))
        co = _by_code(db, Company, v.get("company"))
        dp = _by_code(db, CFA, v.get("depot"))
        obj = LocalAutoRate(code=v["code"], outside_auto_id=auto.id, location_id=loc.id if loc else None,
                            company_id=co.id if co else None, cfa_id=dp.id if dp else None, rate_per_box=v["rate"],
                            effective_from=v["from"], effective_to=v.get("to"), remarks=v.get("remarks"))
    elif sheet == "Freight_Rates":
        tr = _by_code_or_name(db, Transporter, v["transporter"])
        loc = _by_code_or_name(db, Location, v["location"])
        trs.propose(db, user, transporter_id=tr.id, location_id=loc.id, rate_per_case=v["rate"],
                    effective_from=v["from"], reason=f"Initial master workbook ({ref})", source="IMPORT")
        return
    elif sheet == "Invoice_Code_Mappings":
        erp = _by_code_or_name(db, ERPSystem, v.get("erp"))
        co = _by_code(db, Company, v["company"])
        dp = _by_code(db, CFA, v.get("depot"))
        ct = _by_code_or_name(db, CustomerTypeMaster, v.get("ctype"))
        obj = InvoiceMarking(code=v["code"], erp_system_id=erp.id if erp else None, marking=v["marking"],
                             match_type=v.get("match") or "PREFIX", company_id=co.id, cfa_id=dp.id if dp else None,
                             customer_type=ct.code if ct else None, notes=v.get("notes"))
    elif sheet == "ERP_Mappings":
        erp = _by_code_or_name(db, ERPSystem, v["erp"])
        ent = {"DEPOT": "CFA"}.get(v["entity"], v["entity"])
        model = {"CUSTOMER": Customer, "COMPANY": Company, "CFA": CFA, "LOCATION": Location,
                 "TRANSPORTER": Transporter}[ent]
        tgt = _by_code_or_name(db, model, v["target"])
        exists = db.scalar(select(ERPCodeMapping).where(ERPCodeMapping.erp_system_id == erp.id,
                                                        ERPCodeMapping.entity_type == ent,
                                                        func.upper(ERPCodeMapping.source_code) == v["source_code"]))
        if exists:
            return
        obj = ERPCodeMapping(erp_system_id=erp.id, entity_type=ent, source_code=v["source_code"],
                             source_name=v.get("source_name"), target_id=tgt.id, created_by_id=user.id)
    else:
        return
    if sheet != "Customers":
        db.add(obj)
        db.flush()
    audit.log(db, user, "MASTER_CREATED", "MASTER", record_type=type(obj).__name__, record_id=obj.id,
              ref=f"{sheet} {getattr(obj, 'code', '') or getattr(obj, 'source_code', '')}", reason=f"Initial master workbook ({ref})")


def error_report(up: MasterUpload) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Rejected and skipped"
    ws.append(["Sheet", "Excel row", "Code / key", "Result", "Reason(s)", "Warnings", "Row values"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for r in up.rows or []:
        if r["status"] in ("REJECTED", "DUPLICATE", "EXISTS") or r["warnings"]:
            ws.append([r["sheet"], r["row"], r["key"] or "", r["status"], "; ".join(r["errors"]),
                       "; ".join(r["warnings"]),
                       "; ".join(f"{k}={v}" for k, v in (r.get("values") or {}).items() if v is not None)])
    for letter, w in zip("ABCDEFG", (20, 10, 22, 12, 60, 50, 80)):
        ws.column_dimensions[letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
