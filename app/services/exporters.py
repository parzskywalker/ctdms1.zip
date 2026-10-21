"""Export ReportResult objects to CSV / Excel / PDF."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import InvoiceStatus
from app.services.pdf_service import table_pdf
from app.services.reports import ReportResult
from app.services.utils import csv_safe, fmt_date, fmt_dt, fmt_money, to_local


def _plain(v, kind):
    if v is None:
        return ""
    if kind == "datetime" and isinstance(v, datetime):
        return to_local(v)
    if kind == "status":
        return InvoiceStatus.LABELS.get(v, v)
    if kind in ("money", "rate", "pct", "num") and not isinstance(v, str):
        return float(v)
    if kind == "int" and not isinstance(v, str):
        return int(v)
    return v


def _training_mark() -> str | None:
    from app.training import runtime
    return runtime.WATERMARK if runtime.active() else None


def to_csv(res: ReportResult) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    if _training_mark():
        w.writerow([_training_mark()])
    w.writerow([c.label for c in res.columns])
    for r in res.rows:
        row = []
        for c in res.columns:
            v = _plain(r.get(c.key), c.kind)
            if isinstance(v, datetime):
                v = v.strftime("%d-%m-%Y %H:%M")
            elif isinstance(v, date):
                v = v.strftime("%d-%m-%Y")
            row.append(csv_safe(v))
        w.writerow(row)
    if res.totals:
        w.writerow([_plain(res.totals.get(c.key), c.kind) if res.totals.get(c.key) is not None else ""
                    for c in res.columns])
    return buf.getvalue().encode("utf-8-sig")


XLSX_CELL_LIMIT = 32_767  # Excel cannot hold more text than this in one cell


def _fit_cell(v):
    """Excel silently cuts text longer than 32,767 characters (e.g. a transporter's LR list for a whole year).
    Cut it ourselves at an item boundary and say so, instead of losing the tail without warning."""
    if not isinstance(v, str) or len(v) <= XLSX_CELL_LIMIT:
        return v
    note_tpl = " ... (+{n:,} more not shown: Excel allows 32,767 characters per cell - use the CSV export or " \
               "the Dispatch Register report for the full list)"
    room = XLSX_CELL_LIMIT - len(note_tpl.format(n=10 ** 9))
    cut = v.rfind(",", 0, room)
    cut = cut if cut > 0 else room
    hidden = v[cut:].count(",") if "," in v[cut:] else 1
    return v[:cut] + note_tpl.format(n=hidden)

FORMATS = {"money": '#,##0.00', "rate": '#,##0.00', "pct": '0.00', "int": '#,##0', "date": 'dd-mm-yyyy',
           "datetime": 'dd-mm-yyyy hh:mm', "num": '#,##0.00'}


def to_xlsx(res: ReportResult, title: str, subtitle: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31].replace("/", "-")
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13, color="1F3A5F")
    ws.append([subtitle])
    ws["A2"].font = Font(italic=True, size=9, color="5B6675")
    ws.append([_training_mark()])  # Guided Training: "TRAINING - ... NOT VALID FOR BUSINESS USE" (else empty)
    ws["A3"].font = Font(bold=True, size=12, color="B45309")
    head_row = 4
    ws.append([c.label for c in res.columns])
    fill = PatternFill("solid", fgColor="1F3A5F")
    for i in range(1, len(res.columns) + 1):
        cell = ws.cell(row=head_row, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in res.rows:
        ws.append([_fit_cell(csv_safe(_plain(r.get(c.key), c.kind))) for c in res.columns])
    if res.totals:
        ws.append([_plain(res.totals.get(c.key), c.kind) if res.totals.get(c.key) is not None else None
                   for c in res.columns])
        for i in range(1, len(res.columns) + 1):
            ws.cell(row=ws.max_row, column=i).font = Font(bold=True)
    for i, c in enumerate(res.columns, start=1):
        fmt = FORMATS.get(c.kind)
        if fmt:
            for row in range(head_row + 1, ws.max_row + 1):
                ws.cell(row=row, column=i).number_format = fmt
        width = max([len(c.label)] + [len(str(r.get(c.key) or "")) for r in res.rows[:300]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 9), 45)
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(res.columns))}{ws.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def display(v, kind) -> str:
    if v is None or v == "":
        return ""
    if kind == "money":
        return fmt_money(v)
    if kind in ("rate", "num"):
        return f"{Decimal(str(v)):,.2f}"
    if kind == "pct":
        return f"{Decimal(str(v)):.2f}%"
    if kind == "int":
        return f"{int(v):,}"
    if kind == "date":
        return fmt_date(v)
    if kind == "datetime":
        return fmt_dt(v)
    if kind == "status":
        return InvoiceStatus.LABELS.get(v, str(v).replace("_", " ").title())
    return str(v)


def to_pdf(res: ReportResult, title: str, subtitle: str, org: str) -> bytes:
    cols = [(c.label, c.kind) for c in res.columns]
    rows = [[display(r.get(c.key), c.kind) if c.kind != "invoice" else r.get(c.key) for c in res.columns]
            for r in res.rows]
    totals = [display(res.totals.get(c.key), c.kind) if res.totals and res.totals.get(c.key) is not None else ""
              for c in res.columns] if res.totals else None
    return table_pdf(title, subtitle, cols, rows, totals, org)
