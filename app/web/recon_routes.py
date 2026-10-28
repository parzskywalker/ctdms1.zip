from __future__ import annotations

import io
import uuid
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, Response
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import get_db
from app.models import BillStatus, Transporter, TransporterBill, User
from app.security.auth import csrf_protect, require
from app.services import audit, storage
from app.services import bill_scan_service as scan_svc
from app.services import recon_service as rs
from app.services import settings_service as cfg
from app.services.errors import BusinessError
from app.services.reports import Filters
from app.services.scope import apply_bill_scope, can_change
from app.services.utils import csv_safe, local_now, local_today, money, parse_date, parse_decimal, safe_int
from app.web.templating import filter_options, flash, get_form, read_upload, redirect, render

router = APIRouter(dependencies=[Depends(csrf_protect)], include_in_schema=False)


def _bill(db, bill_id, user: User) -> TransporterBill:
    b = db.get(TransporterBill, bill_id)
    if not b or not user.scope_is_open or not can_change(user, b.company_id, b.cfa_id):
        raise HTTPException(404)  # out-of-scope bills behave as if they do not exist
    return b


@router.get("/recon")
def bill_list(request: Request, db: Session = Depends(get_db), user: User = Depends(require("recon.view"))):
    f = Filters.from_params(request.query_params, default_range="fy")
    stmt = apply_bill_scope(select(TransporterBill), user, TransporterBill)
    if f.date_from:
        stmt = stmt.where(TransporterBill.bill_date >= f.date_from)
    if f.date_to:
        stmt = stmt.where(TransporterBill.bill_date <= f.date_to)
    if f.transporter_id:
        stmt = stmt.where(TransporterBill.transporter_id == f.transporter_id)
    if f.status:
        stmt = stmt.where(TransporterBill.status == f.status)
    bills = list(db.scalars(stmt.order_by(TransporterBill.bill_date.desc(), TransporterBill.id.desc())))
    tot = {"calc": sum(b.calculated_amount for b in bills), "bill": sum(b.bill_amount for b in bills),
           "var": sum(b.variance for b in bills)}
    return render(request, "recon/list.html", {"title": "Transporter Bill Reconciliation", "bills": bills, "f": f,
                                               "opts": filter_options(db, user), "statuses": BillStatus.ALL, "tot": tot})


@router.get("/recon/new")
def new_bill(request: Request, db: Session = Depends(get_db), user: User = Depends(require("recon.edit"))):
    return render(request, "recon/new.html", {"title": "New transporter bill", "opts": filter_options(db, user)})


@router.post("/recon/new")
def create_bill(request: Request, transporter_id: int = Form(...), bill_no: str = Form(...), bill_date: str = Form(...),
                bill_amount: str = Form(...), company_id: str = Form(""), cfa_id: str = Form(""),
                period_from: str = Form(""), period_to: str = Form(""), remarks: str = Form(""),
                link: str = Form(""), db: Session = Depends(get_db), user: User = Depends(require("recon.edit"))):
    try:
        bd = parse_date(bill_date)
        pf = parse_date(period_from) if period_from else None
        pt = parse_date(period_to) if period_to else None
    except ValueError as e:
        raise BusinessError(str(e)) from e
    co_id, cf_id = safe_int(company_id), safe_int(cfa_id)
    if (company_id and co_id is None) or (cfa_id and cf_id is None):
        raise BusinessError("Invalid company / depot")
    if not can_change(user, co_id, cf_id):
        raise BusinessError("Select a company and depot you are assigned to")
    b = rs.create_bill(db, user, transporter_id=transporter_id, bill_no=bill_no, bill_date=bd, bill_amount=bill_amount,
                       company_id=co_id, cfa_id=cf_id,
                       period_from=pf, period_to=pt, remarks=remarks or None)
    n = rs.link_dispatches(db, user, b) if link and pf and pt else 0
    db.commit()
    flash(request, f"Bill {b.bill_no} created" + (f" with {n} dispatches linked" if n else "") + ".", "success")
    return redirect(f"/recon/{b.id}")


@router.get("/recon/{bill_id:int}")
def bill_detail(bill_id: int, request: Request, db: Session = Depends(get_db), user: User = Depends(require("recon.view"))):
    b = _bill(db, bill_id, user)
    rep = rs.report(db, b)
    blocking = [ln for ln in b.lines if ln.match_status in rs.BLOCKING_LINES]
    review = rs.rate_review(db, b) if b.lines and not blocking else []
    return render(request, "recon/detail.html", {"title": f"Bill {b.bill_no}", "b": b, "lines_total": rep["lines_total"],
                                                 "editable": b.status in rs.EDITABLE, "rep": rep, "review": review,
                                                 "blocking": blocking, "spare_rows": 3 if b.status in rs.EDITABLE else 0,
                                                 "tol_amt": cfg.get(db, "recon.tolerance_amount"),
                                                 "tol_pct": cfg.get(db, "recon.tolerance_pct"),
                                                 "ai_ready": scan_svc.available() and cfg.get(db, "ai.bill_reader") == "CLAUDE",
                                                 "ai_locked": scan_svc.locked()})


# ----------------------------------------------------------------------------- scanned bills
BILL_DIR = "bills"
MEDIA = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp", ".gif": "image/gif",
         ".pdf": "application/pdf"}


@router.get("/recon/scan")
def scan_form(request: Request, db: Session = Depends(get_db), user: User = Depends(require("recon.edit"))):
    return render(request, "recon/scan.html", {"title": "Scan a transporter bill", "opts": filter_options(db, user),
                                               "ai_ready": scan_svc.available() and cfg.get(db, "ai.bill_reader") == "CLAUDE",
                                                 "ai_locked": scan_svc.locked(),
                                               "today": local_today()})


@router.post("/recon/scan")
async def scan_upload(request: Request, db: Session = Depends(get_db), user: User = Depends(require("recon.edit"))):
    form = await request.form()
    return _scan_create(request, db, user, form, [f for f in form.getlist("files") if getattr(f, "filename", None)])


def _scan_create(request: Request, db: Session, user: User, form, uploads: list):
    from app.config import get_settings
    tid = safe_int(form.get("transporter_id"))
    tr = db.get(Transporter, tid) if tid else None
    if not tr:
        raise BusinessError("Choose whose bill this is (transporter)")
    if not uploads:
        raise BusinessError("Attach the photo(s) or PDF of the bill")
    if len(uploads) > scan_svc.MAX_PAGES:
        raise BusinessError(f"At most {scan_svc.MAX_PAGES} pages per bill")
    co_id, cf_id = safe_int(form.get("company_id")), safe_int(form.get("cfa_id"))
    if not can_change(user, co_id, cf_id):
        raise BusinessError("Select a company and depot you are assigned to")
    files, stored = [], []
    folder = storage.upload_dir() / BILL_DIR
    folder.mkdir(parents=True, exist_ok=True)
    for up in uploads:
        ext = Path(up.filename or "").suffix.lower()
        if ext not in MEDIA:
            raise BusinessError(f"'{up.filename}' is not a photo (JPG/PNG/WEBP) or PDF")
        content = read_upload(up)
        name = f"{local_today():%Y%m%d}-{uuid.uuid4().hex[:10]}{ext}"
        (folder / name).write_bytes(content)
        files.append((content, MEDIA[ext]))
        stored.append({"name": name, "original": (up.filename or "")[:120], "type": MEDIA[ext]})
    ext_data, read_error = None, None
    if scan_svc.available() and cfg.get(db, "ai.bill_reader") == "CLAUDE" and form.get("read") != "manual":
        try:
            ext_data = scan_svc.read_bill(files, hint=tr.name)
        except BusinessError as e:
            read_error = e.message
    try:
        bd = parse_date(form.get("bill_date")) if form.get("bill_date") else (
            scan_svc.to_date(ext_data.bill_date) if ext_data and ext_data.bill_date else None) or local_today()
        pf = parse_date(form.get("period_from")) if form.get("period_from") else (
            scan_svc.to_date(ext_data.period_from) if ext_data else None)
        pt = parse_date(form.get("period_to")) if form.get("period_to") else (
            scan_svc.to_date(ext_data.period_to) if ext_data else None)
    except ValueError as e:
        raise BusinessError(str(e)) from e
    lines = []
    if ext_data:
        for sl in ext_data.lines:
            lines.append({"lr_no": sl.lr_no, "lr_date": scan_svc.to_date(sl.lr_date), "party_name": sl.party_name,
                          "station": sl.station, "billed_boxes": sl.cases, "billed_rate": sl.rate,
                          "billed_amount": sl.amount, "remarks": ("[unclear] " if sl.unclear else "") + (sl.remarks or "")})
        dates = [x["lr_date"] for x in lines if x["lr_date"]]
        if dates and not pf:
            pf = min(dates)
        if dates and not pt:
            pt = max(dates)
    bill_no = (form.get("bill_no") or "").strip() or (ext_data.bill_no if ext_data and ext_data.bill_no else "") or \
        f"SCAN-{local_today():%y%m%d}-{uuid.uuid4().hex[:4].upper()}"
    amount = form.get("bill_amount") or (ext_data.total_amount if ext_data and ext_data.total_amount else None) or \
        sum((x["billed_amount"] or 0) for x in lines) or 0
    b = rs.create_bill(db, user, transporter_id=tr.id, bill_no=str(bill_no), bill_date=bd, bill_amount=amount,
                       company_id=co_id, cfa_id=cf_id, period_from=pf, period_to=pt, source="SCAN",
                       total_cases=ext_data.total_cases if ext_data else None,
                       remarks=(ext_data.notes if ext_data and ext_data.notes else None))
    b.scan_files = stored
    b.extraction = ext_data.model_dump() if ext_data else {}
    rows = [{k: ("" if v is None else str(v)) for k, v in x.items()} for x in lines]
    for r in rows:
        if r.get("lr_date"):
            r["lr_date"] = r["lr_date"][:10]
    if rows:
        rs.save_lines(db, user, b, rows)
    db.commit()
    if ext_data:
        msg = f"Bill {b.bill_no}: {len(lines)} line(s) read from the scan. Check them against the photo, then review the flags."
        if ext_data.transporter_name and not rs.party_matches(ext_data.transporter_name, tr.name):
            flash(request, f"The bill is printed by '{ext_data.transporter_name}' - is it really {tr.name}?", "warn")
        flash(request, msg, "success")
    else:
        flash(request, (f"Automatic reading did not work ({read_error}). " if read_error else "") +
              "Type the bill lines next to the photo below.", "warn" if read_error else "info")
    return redirect(f"/recon/{b.id}#lines")


@router.get("/recon/{bill_id:int}/file/{n:int}")
def scan_file(bill_id: int, n: int, request: Request, db: Session = Depends(get_db),
              user: User = Depends(require("recon.view"))):
    from app.config import get_settings
    b = _bill(db, bill_id, user)
    files = b.scan_files or []
    if not (0 <= n < len(files)):
        raise HTTPException(404)
    path = storage.upload_dir() / BILL_DIR / Path(files[n]["name"]).name
    if not path.exists():
        raise HTTPException(404)
    return FileResponse(path, media_type=files[n].get("type") or "application/octet-stream")


@router.post("/recon/{bill_id:int}/lines")
def save_lines(bill_id: int, request: Request, form=Depends(get_form), db: Session = Depends(get_db),
               user: User = Depends(require("recon.edit"))):
    b = _bill(db, bill_id, user)
    rows: dict[str, dict] = {}
    for key, val in form.multi_items():
        if not key.startswith("ln_"):
            continue
        _, idx, field = key.split("_", 2)
        rows.setdefault(idx, {})[field] = val
    n = rs.save_lines(db, user, b, [rows[k] for k in sorted(rows, key=lambda x: int(x) if x.isdigit() else 0)])
    if form.get("bill_amount"):
        try:
            new = money(parse_decimal(form.get("bill_amount")))
        except ValueError as e:
            raise BusinessError(f"Bill amount: {e}") from e
        if new != b.bill_amount:
            audit.log(db, user, "BILL_AMOUNT_CHANGED", rs.MODULE, record_type="TransporterBill", record_id=b.id,
                      ref=b.bill_no, field="bill_amount", old=b.bill_amount, new=new)
            b.bill_amount = new
            rs.rematch(db, b)
    db.commit()
    flash(request, f"{n} line change(s) saved and re-checked. Bill status: {b.status}.", "success")
    return redirect(f"/recon/{b.id}#lines")


@router.post("/recon/{bill_id:int}/header")
def save_header(bill_id: int, request: Request, form=Depends(get_form), db: Session = Depends(get_db),
                user: User = Depends(require("recon.edit"))):
    b = _bill(db, bill_id, user)
    rs.update_header(db, user, b, bill_no=form.get("bill_no"), bill_date=form.get("bill_date"),
                     period_from=form.get("period_from"), period_to=form.get("period_to"),
                     total_cases=form.get("total_cases"))
    db.commit()
    flash(request, "Bill details saved and re-checked.", "success")
    return redirect(f"/recon/{b.id}")


@router.post("/recon/{bill_id:int}/lines/{line_id:int}/link/{choice:int}")
def link_suggestion(bill_id: int, line_id: int, choice: int, request: Request, db: Session = Depends(get_db),
                    user: User = Depends(require("recon.edit"))):
    b = _bill(db, bill_id, user)
    rs.link_suggestion(db, user, b, line_id, choice)
    db.commit()
    flash(request, "Line linked to our consignment; its LR number is now on our gatepass.", "success")
    return redirect(f"/recon/{b.id}#lines")


@router.post("/recon/{bill_id:int}/rates")
def rate_review(bill_id: int, request: Request, form=Depends(get_form), db: Session = Depends(get_db),
                user: User = Depends(require("recon.edit"))):
    b = _bill(db, bill_id, user)
    changes = []
    for key in form.keys():
        if key.startswith("accept_"):
            loc = key[7:]
            changes.append({"location_id": loc, "rate": form.get(f"newrate_{loc}"), "effective_from": form.get(f"from_{loc}")})
    n = rs.apply_rate_review(db, user, b, changes, form.get("note"))
    db.commit()
    flash(request, f"{n} rate change(s) sent for approval." if n else "Recorded: no rate changes on this bill.", "success")
    return redirect(f"/recon/{b.id}")


@router.get("/recon/{bill_id:int}/report")
def irregularity_report(bill_id: int, request: Request, db: Session = Depends(get_db),
                        user: User = Depends(require("recon.view"))):
    b = _bill(db, bill_id, user)
    rep = rs.report(db, b)
    if request.query_params.get("export") == "xlsx" and user.can("data.export"):
        return _report_xlsx(b, rep)
    return render(request, "recon/report.html", {"title": f"Irregularities · {b.bill_no}", "b": b, "rep": rep,
                                                 "printed_at": local_now()})


def _report_xlsx(b: TransporterBill, rep: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Bill lines"
    ws.append([f"{b.transporter.name} - bill {b.bill_no} dated {b.bill_date:%d-%m-%Y}"])
    ws.append(["LR no", "LR date", "Party", "Station", "Billed cases", "Our cases", "Billed rate", "Card rate",
               "Billed amount", "Our freight", "Difference", "Match", "Irregularities", "Our invoices", "Gatepass"])
    for c in ws[2]:
        c.font = Font(bold=True)
    for ln in b.lines:
        ws.append([csv_safe(ln.lr_no), ln.lr_date, csv_safe(ln.party_name), csv_safe(ln.station), ln.billed_boxes,
                   ln.system_cases, float(ln.billed_rate) if ln.billed_rate is not None else None,
                   float(ln.card_rate) if ln.card_rate is not None else None,
                   float(ln.billed_amount) if ln.billed_amount is not None else None,
                   float(ln.calculated_amount) if ln.calculated_amount is not None else None,
                   float(ln.variance) if ln.variance is not None else None, ln.match_status,
                   csv_safe("; ".join(i["msg"] for i in (ln.issues or []))), csv_safe(ln.invoice_no), ln.dispatch_no])
    ws2 = wb.create_sheet("Not on this bill")
    ws2.append(["Gatepass", "Date", "Customer", "Area", "Invoices", "LR", "Cases", "Our freight"])
    for u in rep["unbilled"]:
        ws2.append([u["dispatch_no"], u["date"], u["customer"], u["location"], u["invoices"], u["lr_no"], u["cases"],
                    float(u["freight"])])
    buf = io.BytesIO()
    wb.save(buf)
    name = f"irregularities-{b.transporter.code}-{b.bill_no}".replace("/", "-")
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})


@router.post("/recon/{bill_id:int}/link")
def link(bill_id: int, request: Request, db: Session = Depends(get_db), user: User = Depends(require("recon.edit"))):
    b = _bill(db, bill_id, user)
    n = rs.link_dispatches(db, user, b)
    db.commit()
    flash(request, f"{n} dispatched invoice(s) linked from {b.period_from:%d-%m-%Y} to {b.period_to:%d-%m-%Y}.", "success")
    return redirect(f"/recon/{b.id}")


@router.post("/recon/{bill_id:int}/import")
def import_lines(bill_id: int, request: Request, file: UploadFile = File(...), db: Session = Depends(get_db),
                 user: User = Depends(require("recon.edit"))):
    b = _bill(db, bill_id, user)
    n, errors = rs.import_lines(db, user, b, read_upload(file), file.filename or "bill.csv")
    db.commit()
    flash(request, f"{n} bill line(s) imported and matched." + (f" {len(errors)} error(s): " + "; ".join(errors[:5]) if errors else ""),
          "success" if not errors else "warn")
    return redirect(f"/recon/{b.id}")


@router.post("/recon/{bill_id:int}/amounts")
def amounts(bill_id: int, request: Request, form=Depends(get_form), db: Session = Depends(get_db),
            user: User = Depends(require("recon.edit"))):
    b = _bill(db, bill_id, user)
    vals = {k[4:]: v for k, v in form.items() if k.startswith("amt_")}
    rs._editable(b)  # approved / paid bills are frozen
    if form.get("bill_amount"):
        from app.services.utils import money, parse_decimal
        try:
            new = money(parse_decimal(form.get("bill_amount")))
        except ValueError as e:
            raise BusinessError(f"Bill amount: {e}") from e
        if new is None or new <= 0:
            raise BusinessError("Bill amount must be greater than zero")
        if new != b.bill_amount:
            from app.services import audit
            audit.log(db, user, "BILL_AMOUNT_CHANGED", rs.MODULE, record_type="TransporterBill", record_id=b.id,
                      ref=b.bill_no, field="bill_amount", old=b.bill_amount, new=new)
            b.bill_amount = new
    rs.set_line_amounts(db, user, b, vals)
    db.commit()
    flash(request, f"Amounts saved. Bill status: {b.status}.", "success")
    return redirect(f"/recon/{b.id}")


@router.post("/recon/{bill_id:int}/allocate")
def allocate(bill_id: int, request: Request, db: Session = Depends(get_db), user: User = Depends(require("recon.edit"))):
    b = _bill(db, bill_id, user)
    rs.allocate_header(db, user, b)
    db.commit()
    flash(request, "Bill amount allocated to lines in proportion to system freight.", "success")
    return redirect(f"/recon/{b.id}")


@router.post("/recon/{bill_id:int}/lines/{line_id:int}/delete")
def delete_line(bill_id: int, line_id: int, request: Request, db: Session = Depends(get_db),
                user: User = Depends(require("recon.edit"))):
    b = _bill(db, bill_id, user)
    rs.delete_line(db, user, b, line_id)
    db.commit()
    return redirect(f"/recon/{b.id}")


@router.post("/recon/{bill_id:int}/approve")
def approve(bill_id: int, request: Request, remarks: str = Form(""), db: Session = Depends(get_db),
            user: User = Depends(require("recon.approve"))):
    b = _bill(db, bill_id, user)
    rs.approve(db, user, b, remarks)
    db.commit()
    flash(request, f"Bill {b.bill_no} approved.", "success")
    return redirect(f"/recon/{b.id}")


@router.post("/recon/{bill_id:int}/pay")
def pay(bill_id: int, request: Request, payment_ref: str = Form(""), db: Session = Depends(get_db),
        user: User = Depends(require("recon.pay"))):
    b = _bill(db, bill_id, user)
    rs.mark_paid(db, user, b, payment_ref)
    db.commit()
    flash(request, f"Bill {b.bill_no} marked PAID ({payment_ref}).", "success")
    return redirect(f"/recon/{b.id}")


@router.post("/recon/{bill_id:int}/reject")
def reject(bill_id: int, request: Request, reason: str = Form(""), db: Session = Depends(get_db),
           user: User = Depends(require("recon.approve"))):
    b = _bill(db, bill_id, user)
    rs.reject(db, user, b, reason)
    db.commit()
    flash(request, f"Bill {b.bill_no} rejected.", "warn")
    return redirect(f"/recon/{b.id}")
