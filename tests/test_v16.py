"""v1.6 - Depots, customer types, Invoice Recognition Code, ERP invoice markings, local auto per trip / per box,
one-time master workbook upload, reports by city and depot."""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models import (
    CFA, Company, Customer, CustomerTypeMaster, DispatchMode, Invoice, InvoiceMarking, LocalAutoRate, Location,
    MasterUpload, OutsideAuto, PickupMode, Transporter, TransporterRate, User,
)
from app.services import dispatch_service as ds
from app.services import import_service as imp
from app.services import master_upload_service as mus
from app.services import recognition
from app.services import settings_service as cfg
from app.services.errors import BusinessError
from app.services.utils import local_today
from tests.helpers import get_inv, make_invoice, new_invoice_no


def _u(db, name) -> User:
    return db.scalar(select(User).where(User.username == name))


def _fresh(api, customer, boxes=4, company="COA", cfa="RPR"):
    return get_inv(make_invoice(api("admin"), company=company, cfa=cfa, customer=customer, boxes=boxes)["id"])


# ----------------------------------------------------------------------------- terminology
def test_depot_terminology_replaces_cfa(web):
    w = web("admin")
    html = w.get("/").text
    assert "Depot" in html and "Master Settings" in html and "CFA" not in html.replace("CFA_", "")
    assert "Depot Manager" in w.get("/admin/roles").text
    page = w.get("/masters/depots")
    assert page.status_code == 200 and "Depot Master" in page.text and "Kamal Vihar" in page.text
    assert w.get("/masters/cfas").status_code == 200  # v1.5 links still work


def test_two_depots_in_one_city_and_city_report(web, db):
    raipur = list(db.scalars(select(CFA).where(CFA.city == "Raipur")))
    assert {c.code for c in raipur} >= {"RPR", "KVR"} and len({c.code for c in raipur}) == len(raipur)
    m = web("management")
    r = m.get("/reports/city-freight?date_from=2000-01-01")
    assert r.status_code == 200 and "Raipur" in r.text and "Guma Depot" in r.text
    d = m.get("/reports/cfa-freight?date_from=2000-01-01")
    assert "City" in d.text and "Guma Depot" in d.text
    f = m.get("/invoices?city=Raipur&date_from=2000-01-01")
    assert f.status_code == 200 and "Guma" in f.text


# ----------------------------------------------------------------------------- customer types
def test_customer_types_are_a_master(web, db):
    codes = {c.code for c in db.scalars(select(CustomerTypeMaster))}
    assert {"RETAILER", "STOCKIST", "DISTRIBUTOR", "HOSPITAL", "INSTITUTION", "CHAIN_PHARMACY", "OTHER"} <= codes
    w = web("admin")
    w.post("/masters/customer-types/save", {"code": "CLINIC", "name": "Clinic", "marking": "CLN", "is_active": "1"})
    lo = db.scalar(select(Location).where(Location.code == "RPR"))
    w.post("/masters/customers/save", {"code": "C-V16CL", "name": "Test Clinic", "location_id": str(lo.id),
                                       "party_type": "CLINIC", "is_active": "1"})
    cu = db.scalar(select(Customer).where(Customer.code == "C-V16CL"))
    assert cu.party_type == "CLINIC" and cu.type_label == "Clinic"
    bad = w.post("/masters/customers/save", {"code": "C-V16BAD", "name": "Bad", "location_id": str(lo.id),
                                             "party_type": "NOPE", "is_active": "1"})
    assert "not in Master Settings" in bad.text
    assert db.scalar(select(Customer).where(Customer.code == "C-V16BAD")) is None


# ----------------------------------------------------------------------------- recognition code
def test_serial_part_rules():
    assert recognition.serial_part("001245") == "001245"
    assert recognition.serial_part("MB/26/1245") == "001245"
    assert recognition.serial_part("MB/26/1245", "ALL_DIGITS", 0) == "261245"
    assert recognition.serial_part("CBO/RPR/00045", "FULL", 0) == "CBORPR00045"
    assert recognition.serial_part("NO-DIGITS") is None


def test_recognition_code_made_and_erp_number_untouched(api, db):
    inv = _fresh(api, "C-MET01", boxes=3)
    full = db.get(Invoice, inv.id)
    assert full.invoice_no == inv.invoice_no  # the ERP number as received
    ct = db.scalar(select(CustomerTypeMaster).where(CustomerTypeMaster.code == full.customer.party_type))
    parts = full.recognition_code.split("-")
    assert parts[0] == "COA" and parts[1] == ct.marking and parts[3] == "GUMA"
    assert parts[-1] == recognition.serial_part(full.invoice_no) and full.serial_no == parts[-1]
    assert full.recognition_note is None
    js = api("admin").get(f"/api/v1/invoices/{inv.id}").json()
    assert js["original_invoice_no"] == full.invoice_no and js["recognition_code"] == full.recognition_code


def test_recognition_code_pattern_missing_code_and_regenerate(api, db):
    adm = _u(db, "admin")
    ct = db.scalar(select(CustomerTypeMaster).where(CustomerTypeMaster.code == "HOSPITAL"))
    old = ct.marking
    ct.marking = None
    db.commit()
    try:
        inv = _fresh(api, "C-AIM01", boxes=2)   # a hospital
        row = db.get(Invoice, inv.id)
        db.refresh(row)
        assert row.recognition_code is None and "Customer type Hospital has no recognition code" in row.recognition_note
        assert row.invoice_no  # dispatch does not depend on the code
    finally:
        ct = db.scalar(select(CustomerTypeMaster).where(CustomerTypeMaster.code == "HOSPITAL"))
        ct.marking = old
        db.commit()
    res = recognition.regenerate(db, adm, scope="missing")
    db.commit()
    row = db.get(Invoice, inv.id)
    db.refresh(row)
    assert row.recognition_code and row.recognition_code.split("-")[1] == old and res["checked"] >= 1
    # a different pattern
    cfg.set_value(db, "invoice_code.format", "{company}/{depot}/{serial}")
    try:
        code, serial, note = recognition.build(db, row)
        assert code == f"COA/GUMA/{serial}" and note is None
    finally:
        cfg.set_value(db, "invoice_code.format", "{company}-{ctype}-{customer}-{depot}-{serial}")
        db.commit()


def test_recognition_switched_off_makes_no_new_codes(api, db):
    cfg.set_value(db, "invoice_code.enabled", False)
    db.commit()
    try:
        inv = _fresh(api, "C-MET01", boxes=1)
        assert db.get(Invoice, inv.id).recognition_code is None
    finally:
        cfg.set_value(db, "invoice_code.enabled", True)
        db.commit()


def test_recognition_page_and_search(web, api, db):
    inv = _fresh(api, "C-MET01", boxes=2)
    code = db.get(Invoice, inv.id).recognition_code
    w = web("admin")
    page = w.get("/masters/recognition")
    assert page.status_code == 200 and "Invoice Recognition Code" in page.text and "never changed" in page.text
    r = w.get(f"/invoices?q={code}")
    assert r.status_code == 200 and inv.invoice_no in r.text  # the code finds the invoice (opens it when unique)
    prefix = "-".join(code.split("-")[:2])
    lst = w.get(f"/invoices?rcode={prefix}&date_from=2000-01-01")
    assert inv.invoice_no in lst.text and "Invoice Recognition Code" in lst.text


def test_invoice_marking_recognises_company_and_depot(db):
    adm = _u(db, "admin")
    co = db.scalar(select(Company).where(Company.code == "COB"))
    kvr = db.scalar(select(CFA).where(CFA.code == "KVR"))
    db.add(InvoiceMarking(code="T-KV", marking="KVB/", company_id=co.id, cfa_id=kvr.id, match_type="PREFIX"))
    db.flush()
    cu = db.scalar(select(Customer).where(Customer.code == "C-MET01"))
    no = "KVB/" + new_invoice_no("M")
    csv_text = ("Bill No,Bill Date,Party Code,Party Name,Net Amount,Cases\n"
                f"{no},{local_today():%d-%m-%Y},{cu.code},{cu.name},12000,3\n")
    batch, result, _ = imp.quick_import(db, adm, csv_text.encode(), "marked.csv")
    assert result["imported"] == 1, [r.errors for r in batch.rows]
    inv = db.scalar(select(Invoice).where(Invoice.invoice_no == no.upper()))
    assert inv.company_id == co.id and inv.cfa_id == kvr.id and inv.marking_id is not None
    assert inv.recognition_code.endswith("-KVIHAR-" + recognition.serial_part(inv.invoice_no))
    db.rollback()


# ----------------------------------------------------------------------------- local autos
def test_local_auto_per_trip_four_invoices_one_amount(api, db):
    adm = _u(db, "admin")
    auto = db.scalar(select(OutsideAuto).where(OutsideAuto.code == "AUTO-RAM"))   # per trip
    invs = [_fresh(api, "C-MET01", boxes=b) for b in (5, 15, 3, 7)]
    rows = [db.get(Invoice, i.id) for i in invs]
    with __import__("pytest").raises(BusinessError):  # per box not allowed for this auto
        ds.create_gatepass(db, adm, rows, ds.GatepassRequest(mode=DispatchMode.OUTSIDE_AUTO, outside_auto_id=auto.id,
                                                             auto_method="PER_CASE", cases={i.id: i.boxes for i in invs}))
    db.rollback()
    rows = [db.get(Invoice, i.id) for i in invs]
    slip = ds.create_gatepass(db, adm, rows, ds.GatepassRequest(
        mode=DispatchMode.OUTSIDE_AUTO, outside_auto_id=auto.id, trip_charge=Decimal("500"),
        cases={i.id: i.boxes for i in invs}))
    assert slip.trip_charge == Decimal("500.00") and slip.auto_method == "PER_TRIP"
    assert slip.total_charge == Decimal("500.00")          # NOT 4 x 500
    shares = [ln.freight for ln in slip.lines]
    assert sum(shares) == Decimal("500.00")
    assert shares[0] == Decimal("83.33") and shares[1] == Decimal("250.00")  # by boxes: 5/30, 15/30
    assert "500.00" in str(slip.snapshot["outside_auto"]["trip_charge"])
    db.rollback()


def test_local_auto_trip_charge_varies_and_is_required(api, db):
    adm = _u(db, "admin")
    auto = db.scalar(select(OutsideAuto).where(OutsideAuto.code == "AUTO-RAM"))
    a, b = _fresh(api, "C-MET01", boxes=5), _fresh(api, "C-MET01", boxes=15)
    slip = ds.create_gatepass(db, adm, [db.get(Invoice, a.id), db.get(Invoice, b.id)], ds.GatepassRequest(
        mode=DispatchMode.OUTSIDE_AUTO, outside_auto_id=auto.id, trip_charge=Decimal("500"),
        cases={a.id: 5, b.id: 15}))
    assert [ln.freight for ln in slip.lines] == [Decimal("125.00"), Decimal("375.00")]
    db.rollback()
    auto = db.scalar(select(OutsideAuto).where(OutsideAuto.code == "AUTO-RAM"))
    auto.default_trip_charge = None
    db.flush()
    _, _, problems = ds.plan_gatepass(db, adm, [db.get(Invoice, a.id)], ds.GatepassRequest(
        mode=DispatchMode.OUTSIDE_AUTO, outside_auto_id=auto.id, cases={a.id: 5}))
    assert any("Enter the trip charge" in p for p in problems)
    db.rollback()


def test_local_auto_per_box_rate_by_location_and_both(api, db):
    adm = _u(db, "admin")
    kish = db.scalar(select(OutsideAuto).where(OutsideAuto.code == "AUTO-KISH"))   # BOTH; 18/box, Bhilai 22
    assert kish.methods == ["PER_TRIP", "PER_CASE"]
    bhl = _fresh(api, "C-HRM01" if False else "C-BHL01", boxes=15) if db.scalar(
        select(Customer).where(Customer.code == "C-BHL01")) else None
    rpr = _fresh(api, "C-MET01", boxes=15)
    slip = ds.create_gatepass(db, adm, [db.get(Invoice, rpr.id)], ds.GatepassRequest(
        mode=DispatchMode.OUTSIDE_AUTO, outside_auto_id=kish.id, auto_method="PER_CASE", cases={rpr.id: 15}))
    assert slip.total_charge == Decimal("270.00") and slip.lines[0].rate_per_case == Decimal("18.00")  # 15 x 18
    db.rollback()
    bhl_loc = db.scalar(select(Location).where(Location.code == "BHL"))
    from app.services import local_auto_service as las
    br = las.box_rate(db, kish, location_id=bhl_loc.id, company_id=None, cfa_id=None, on=local_today())
    assert br.rate == Decimal("22.00")
    # a company-specific rate wins over the location-only one
    co = db.scalar(select(Company).where(Company.code == "COA"))
    db.add(LocalAutoRate(code="T-KISH-COA", outside_auto_id=kish.id, location_id=bhl_loc.id, company_id=co.id,
                         rate_per_box=Decimal("20.00"), effective_from=date(2000, 1, 1)))
    db.flush()
    assert las.box_rate(db, kish, location_id=bhl_loc.id, company_id=co.id, cfa_id=None,
                        on=local_today()).rate == Decimal("20.00")
    db.rollback()


def test_long_distance_transporter_unchanged_with_local_auto_pickup(api, db):
    adm = _u(db, "admin")
    sonu = db.scalar(select(OutsideAuto).where(OutsideAuto.code == "AUTO-SONU"))  # per box 12
    osho = db.scalar(select(Transporter).where(Transporter.code == "OSHO"))
    inv = _fresh(api, "C-HRM01", boxes=10)
    slip = ds.create_gatepass(db, adm, [db.get(Invoice, inv.id)], ds.GatepassRequest(
        transporter_id=osho.id, pickup_mode=PickupMode.OUTSIDE_AUTO, outside_auto_id=sonu.id, cases={inv.id: 10}))
    assert slip.freight_total == Decimal("250.00")       # rate card, as in v1.5
    assert slip.auto_charge == Decimal("120.00") and slip.auto_method == "PER_CASE"
    db.rollback()


def test_local_auto_trips_report_no_double_counting(api, web, db):
    adm = _u(db, "admin")
    auto = db.scalar(select(OutsideAuto).where(OutsideAuto.code == "AUTO-RAM"))
    invs = [_fresh(api, "C-MET01", boxes=b) for b in (2, 4, 6)]
    slip = ds.create_gatepass(db, adm, [db.get(Invoice, i.id) for i in invs], ds.GatepassRequest(
        mode=DispatchMode.OUTSIDE_AUTO, outside_auto_id=auto.id, trip_charge=Decimal("450"),
        cases={i.id: i.boxes for i in invs}))
    db.commit()
    from app.services.reports import Filters, rpt_outside_auto
    res = rpt_outside_auto(db, adm, Filters(date_from=date(2000, 1, 1)))
    row = next(r for r in res.rows if r["dispatch_no"] == slip.dispatch_no)
    assert row["trip_charge"] == Decimal("450.00") == row["allocated"] and row["invoices"] == 3
    assert all(r["trip_charge"] == r["allocated"] for r in res.rows)
    page = web("management").get("/reports/outside-auto?date_from=2000-01-01")
    assert slip.dispatch_no in page.text and "nothing" in page.text.lower() or "not ₹2,000" in page.text


def test_local_auto_master_form_and_gatepass_screen(web, db):
    w = web("admin")
    w.post("/masters/autos/save", {"code": "AUTO-V16", "name": "Both Ways Auto", "allowed_methods": "BOTH",
                                   "charge_basis": "PER_CASE", "default_trip_charge": "550",
                                   "default_box_rate": "20", "is_active": "1"})
    a = db.scalar(select(OutsideAuto).where(OutsideAuto.code == "AUTO-V16"))
    assert a.methods == ["PER_TRIP", "PER_CASE"] and a.default_method == "PER_CASE" and a.rate == Decimal("20.00")
    page = w.get("/gatepass/new")
    assert "Both Ways Auto" in page.text and "Trip charge" in page.text and "Per box" in page.text


# ----------------------------------------------------------------------------- master workbook
def _workbook(rows: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, data in rows.items():
        sp = mus.SPEC[name]
        ws = wb.create_sheet(name)
        ws.append([c.header + (" *" if c.required else "") for c in sp.cols])
        for r in data:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row(sheet: str, **vals) -> list:
    return [vals.get(c.key) for c in mus.SPEC[sheet].cols]


def test_master_template_has_instructions_and_sheets(web):
    r = web("admin").get("/masters/upload/template.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Instructions"
    assert {"Companies", "Depots", "Customers", "Customer_Types", "Transporters", "Locations", "Freight_Rates",
            "Local_Auto_Rates", "Invoice_Code_Mappings", "ERP_Mappings"} <= set(wb.sheetnames)
    assert not any("SKU" in s or "Product" in s for s in wb.sheetnames)


def test_master_workbook_validate_preview_import_history(db, web):
    adm = _u(db, "admin")
    content = _workbook({
        "Companies": [_row("Companies", code="V6ABC", name="ABC Pharma (test)", marking="ABC", erp="MARG"),
                      _row("Companies", code="COA", name="Already there"),                  # exists -> skipped
                      _row("Companies", code="V6ABC", name="Twice")],                       # duplicate in file
        "Locations": [_row("Locations", code="V6-RPR", name="V6 Raipur City", city="Raipur")],
        "Depots": [_row("Depots", code="V6GUMA", name="Guma (test)", locality="Guma", city="Raipur",
                        companies="V6ABC, COA", marking="GUMA6"),
                   _row("Depots", code="V6KV", name="Kamal Vihar (test)", locality="Kamal Vihar", city="Raipur",
                        companies="V6ABC"),
                   _row("Depots", code="V6BAD", name="No city")],                           # city mandatory
        "Customer_Types": [_row("Customer_Types", code="V6CHEM", name="Chemist (test)", marking="CHM")],
        "Customers": [_row("Customers", code="V6C025", name="Sharma Medical (test)", ctype="RETAILER",
                           location="V6-RPR", companies="V6ABC", depots="V6GUMA", erp="MARG", erp_code="SH025"),
                      _row("Customers", code="V6C026", name="Unknown place", ctype="RETAILER", location="Nowhere"),
                      _row("Customers", code="V6C027", name="Chemist", ctype="Chemist (test)", location="V6 Raipur City")],
        "Transporters": [_row("Transporters", code="V6SRL", name="Shree Roadlines (test)", kind="LONG_DISTANCE",
                              payment="TBB"),
                         _row("Transporters", code="V6LA", name="Sharma Local Auto (test)", kind="LOCAL_AUTO",
                              methods="PER_TRIP", trip="500"),
                         _row("Transporters", code="V6LB", name="Box Auto (test)", kind="LOCAL_AUTO",
                              methods="PER_BOX", box="20")],
        "Local_Auto_Rates": [_row("Local_Auto_Rates", code="V6LB-RPR", auto="V6LB", location="V6-RPR", rate="20",
                                  **{"from": "01-10-2026"})],
        "Freight_Rates": [_row("Freight_Rates", transporter="V6SRL", location="V6-RPR", rate="12",
                               **{"from": "01-10-2026"}),
                          _row("Freight_Rates", transporter="V6LA", location="V6-RPR", rate="5",
                               **{"from": "01-10-2026"})],                                   # local auto: rejected
        "Invoice_Code_Mappings": [_row("Invoice_Code_Mappings", code="V6-MARK", erp="MARG", marking="AB/",
                                       company="V6ABC", depot="V6GUMA")],
    })
    up = mus.preview(db, adm, content, "masters.xlsx")
    db.commit()
    c = up.counts
    assert c["Companies"] == {"rows": 3, "new": 1, "exists": 1, "rejected": 1, "warnings": 1}
    assert c["Depots"]["new"] == 2 and c["Depots"]["rejected"] == 1
    assert c["Customers"]["new"] == 2 and c["Customers"]["rejected"] == 1
    assert c["Freight_Rates"]["new"] == 1 and c["Freight_Rates"]["rejected"] == 1
    assert db.scalar(select(Company).where(Company.code == "V6ABC")) is None  # nothing saved before Import
    page = web("admin").get(f"/masters/upload/{up.id}")
    assert "Nowhere" in page.text and "not found" in page.text and "City is mandatory" in page.text
    created = mus.commit(db, adm, db.get(MasterUpload, up.id))
    db.commit()
    assert created["Companies"] == 1 and created["Depots"] == 2 and created["Customers"] == 2
    gm, kv = (db.scalar(select(CFA).where(CFA.code == x)) for x in ("V6GUMA", "V6KV"))
    assert gm.city == kv.city == "Raipur" and gm.locality == "Guma" and {c.code for c in gm.companies} == {"V6ABC", "COA"}
    cu = db.scalar(select(Customer).where(Customer.code == "V6C025"))
    assert cu.source_erp_code == "SH025" and imp.find_code_mapping(db, cu.source_erp_id, "CUSTOMER", "SH025")
    assert db.scalar(select(Customer).where(Customer.code == "V6C027")).party_type == "V6CHEM"
    la = db.scalar(select(OutsideAuto).where(OutsideAuto.code == "V6LA"))
    lb = db.scalar(select(OutsideAuto).where(OutsideAuto.code == "V6LB"))
    assert la.methods == ["PER_TRIP"] and la.default_trip_charge == Decimal("500.00")
    assert lb.methods == ["PER_CASE"] and db.scalar(select(func.count(LocalAutoRate.id))
                                                    .where(LocalAutoRate.outside_auto_id == lb.id)) == 1
    tr = db.scalar(select(Transporter).where(Transporter.code == "V6SRL"))
    assert db.scalar(select(TransporterRate).where(TransporterRate.transporter_id == tr.id)).status == "PENDING_APPROVAL"
    assert db.scalar(select(InvoiceMarking).where(InvoiceMarking.code == "V6-MARK")).company.code == "V6ABC"
    # cannot import twice; a second upload of the same workbook finds everything already there
    import pytest
    with pytest.raises(BusinessError):
        mus.commit(db, adm, db.get(MasterUpload, up.id))
    again = mus.preview(db, adm, content, "masters-again.xlsx")
    assert mus.totals(again)["new"] == 0
    db.rollback()
    rep = web("admin").get(f"/masters/upload/{up.id}/errors.xlsx")
    ws = load_workbook(io.BytesIO(rep.content)).active
    assert ws.max_row > 3 and ws["A1"].value == "Sheet"
    hist = web("admin").get("/masters/upload")
    assert "masters.xlsx" in hist.text and "Master Settings" in hist.text


def test_master_upload_rejects_non_excel(db):
    import pytest
    with pytest.raises(BusinessError):
        mus.preview(db, _u(db, "admin"), b"a,b\n1,2\n", "masters.csv")
    db.rollback()


def test_sku_reporting_postponed(web, db):
    w = web("management")
    assert "SKU-wise" not in w.get("/reports").text
    assert cfg.get(db, "features.sku_lines") is False


def test_approve_all_rate_cards_keeps_maker_checker(db, web):
    from app.services import transporter_rate_service as trs
    lo1 = db.scalar(select(Location).where(Location.code == "RGH"))
    lo2 = db.scalar(select(Location).where(Location.code == "NGP"))
    tr = db.scalar(select(Transporter).where(Transporter.code == "OSHO"))
    mine = trs.propose(db, _u(db, "checker"), transporter_id=tr.id, location_id=lo1.id, rate_per_case="41",
                       effective_from=date(2031, 1, 1), check_permission=False)
    other = trs.propose(db, _u(db, "cfa.raipur"), transporter_id=tr.id, location_id=lo2.id, rate_per_case="42",
                        effective_from=date(2031, 2, 1), check_permission=False)
    db.commit()
    web("checker").post("/rate-cards/approve-all", {"remarks": "test"})
    db.expire_all()
    assert db.get(TransporterRate, other.id).status == "APPROVED"
    assert db.get(TransporterRate, mine.id).status == "PENDING_APPROVAL"


def test_trip_charge_with_extra_digits_is_refused(api, web, db):
    """'500.00500' (typed after a pre-filled 500.00) must not quietly become ₹500.01."""
    auto = db.scalar(select(OutsideAuto).where(OutsideAuto.is_active.is_(True)))
    auto.allowed_methods = "PER_TRIP"
    db.commit()
    inv = _fresh(api, "C-VDP01", boxes=3)
    w = web("admin")
    r = w.post("/gatepass/new", data={"ids": str(inv.id), f"cases_{inv.id}": "3", "mode": DispatchMode.OUTSIDE_AUTO,
                                      "outside_auto_id": str(auto.id), "auto_method": "PER_TRIP",
                                      "trip_charge": "500.00500"})
    assert "is not a rupee amount" in r.text
    assert db.scalar(select(func.count()).select_from(Invoice).where(Invoice.id == inv.id,
                                                                     Invoice.dispatch_slip_id.is_not(None))) == 0
    html = w.get("/gatepass/new").text
    assert "tripTouched" in html  # the screen no longer refills an emptied trip charge


def test_local_auto_gatepass_prints_trip_amount(api, web, db):
    adm = _u(db, "admin")
    auto = db.scalar(select(OutsideAuto).where(OutsideAuto.code == "AUTO-RAM"))
    invs = [_fresh(api, "C-MET01", boxes=b) for b in (2, 3)]
    slip = ds.create_gatepass(db, adm, [db.get(Invoice, i.id) for i in invs], ds.GatepassRequest(
        mode=DispatchMode.OUTSIDE_AUTO, outside_auto_id=auto.id, trip_charge=Decimal("550"),
        cases={i.id: i.boxes for i in invs}))
    db.commit()
    w = web("admin")
    from tests.helpers import confirm_labels
    confirm_labels(slip.id)                                  # the case labels are printed before the gatepass
    r = w.post(f"/gatepass/{slip.id}/print", {"size": "A5", "mode": "print"})
    assert r.status_code == 200 and "for the whole trip" in r.text and "550.00" in r.text
    assert invs[0].invoice_no in r.text
    pdf = w.post(f"/gatepass/{slip.id}/print", {"size": "A5", "mode": "pdf", "reason": "PDF copy for the file"})
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    detail = w.get(f"/gatepass/{slip.id}")
    assert "Per trip" in detail.text and "owed once" in detail.text
