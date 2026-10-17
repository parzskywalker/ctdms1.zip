"""Acceptance tests - one test (or more) per item of the 27-point acceptance list in the brief.
See docs/ACCEPTANCE.md for the matching manual demo steps."""
from __future__ import annotations

import csv
import io
from datetime import timedelta
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models import (
    CFA, AuditLog, Company, FreightRate, ImportBatch, Invoice, InvoiceStatus, Location, RateStatus,
    SlipPrintLog, Transporter, TransporterBill, User,
)
from app.services import rate_service, recon_service
from app.services.utils import local_today
from tests.helpers import get_inv, make_invoice


# 1-5 ------------------------------------------------------------------------------------------------
def test_01_05_demo_data_breadth(db):
    imported = db.scalar(select(func.count(Invoice.id)).where(Invoice.import_batch_id.isnot(None)))
    assert imported >= 20
    assert db.scalar(select(func.count(ImportBatch.id))) >= 4
    assert db.scalar(select(func.count(func.distinct(Invoice.company_id)))) >= 3
    assert db.scalar(select(func.count(func.distinct(Invoice.customer_id)))) >= 5
    assert db.scalar(select(func.count(func.distinct(Invoice.cfa_id)))) >= 3
    assert db.scalar(select(func.count(func.distinct(Invoice.location_id)))) >= 6
    assert db.scalar(select(func.count(func.distinct(Invoice.source_erp_id)))) >= 4  # MARG, SAP, CBO, Intrect


# 6 --------------------------------------------------------------------------------------------------
def test_06_different_rates_per_company_and_location(api):
    a = api("admin")
    q = lambda co, loc: a.get("/api/v1/freight/quote", params={"company_code": co, "cfa_code": "RPR",  # noqa: E731
                                                               "location": loc, "boxes": 36}).json()
    assert q("COA", "Bilaspur")["rate_per_box"] == "8.00"
    assert q("COA", "Durg")["rate_per_box"] == "6.50"   # revised this month (was 6.00)
    assert q("COA", "Nagpur")["rate_per_box"] == "12.00"
    assert q("COB", "Bilaspur")["rate_per_box"] == "9.00"
    assert q("COB", "Durg")["rate_per_box"] == "7.00"
    assert q("COB", "Nagpur")["rate_per_box"] == "14.00"
    assert q("COA", "Bilaspur")["transportation_expense"] == "288.00"  # 36 x 8 (brief example)


# 7 --------------------------------------------------------------------------------------------------
def test_07_customer_specific_rate_wins(api):
    a = api("admin")
    std = a.get("/api/v1/freight/quote", params={"company_code": "COA", "cfa_code": "RPR", "location": "Bilaspur",
                                                 "customer_code": "C-VDP01"}).json()
    cust = a.get("/api/v1/freight/quote", params={"company_code": "COA", "cfa_code": "RPR", "location": "Bilaspur",
                                                  "customer_code": "C-ABC01"}).json()
    assert std["level"] == "CFA_LOCATION" and std["rate_per_box"] == "8.00"
    assert cust["level"] == "CUSTOMER" and cust["rate_per_box"] == "7.00"
    # company-wide (level 3) and default (level 4)
    kb = a.get("/api/v1/freight/quote", params={"company_code": "COB", "cfa_code": "BSP", "location": "Korba"}).json()
    assert kb["level"] == "LOCATION" and kb["rate_per_box"] == "11.00"
    df = a.get("/api/v1/freight/quote", params={"company_code": "OWN", "cfa_code": "RPR", "location": "Raigarh"}).json()
    assert df["level"] == "DEFAULT" and df["rate_per_box"] == "10.00"


# 8 --------------------------------------------------------------------------------------------------
def test_08_automatic_boxwise_freight(api):
    inv = make_invoice(api("admin"), company="COA", cfa="RPR", customer="C-VDP01", boxes=36)
    assert inv["status"] == "PENDING"
    assert inv["rate_per_box"] == "8.00"
    assert inv["transportation_expense"] == "288.00"


# 9 --------------------------------------------------------------------------------------------------
def test_09_missing_rate_exception_blocks_dispatch(api):
    a = api("admin")
    inv = make_invoice(a, company="COC", cfa="RPR", customer="C-SUN01", boxes=15)  # Company C has no Korba rate
    assert inv["status"] == "RATE_MISSING"
    assert inv["transportation_expense"] is None and inv["rate_per_box"] is None
    r = a.post(f"/api/v1/invoices/{inv['id']}/verify")
    assert r.status_code == 400 and "rate missing" in r.json()["detail"].lower()


# 10 -------------------------------------------------------------------------------------------------
def test_10_pending_dispatch_queue(api, web):
    rows = api("operator.raipur").get("/api/v1/dispatch/pending").json()["items"]
    assert rows and all(r["status"] in InvoiceStatus.OPEN for r in rows)
    assert all(r["cfa"]["code"] == "RPR" for r in rows)  # operator only sees own CFA
    page = web("operator.raipur").get("/dispatch/pending")
    assert page.status_code == 200 and "OPEN DISPATCH" in page.text


# 11-16 ----------------------------------------------------------------------------------------------
def test_11_16_dispatch_workflow_with_mandatory_slip(api, web, db):
    inv = make_invoice(api("admin"), company="COB", cfa="RPR", customer="C-XYZ01", boxes=20)  # 20 x 7 = 140
    w = web("operator.raipur")
    iid = inv["id"]
    tr = db.scalar(select(Transporter).where(Transporter.code == "SBT"))
    today = local_today().isoformat()
    form = {"transporter_id": tr.id, "vehicle_no": "cg07 bt 1010", "lr_no": "sbt-t-1", "dispatch_date": today,
            "dispatch_time": "15:30", "remarks": "test"}
    # 11/12 transporter, vehicle, LR saved (normalised) + verify
    r = w.post(f"/dispatch/{iid}/verify", form)
    assert r.status_code == 200
    i = get_inv(iid)
    assert i.status == "VERIFIED" and i.vehicle_no == "CG07BT1010" and i.lr_no == "SBT-T-1"
    assert i.transporter_id == tr.id
    # 15 completion blocked before slip exists / before printing
    w.post(f"/dispatch/{iid}/complete")
    assert get_inv(iid).status == "VERIFIED"
    # 13 slip generation -> unique dispatch number
    w.post(f"/dispatch/{iid}/generate", form)
    i = get_inv(iid)
    assert i.status == "SLIP_GENERATED" and i.dispatch_no.startswith("DS/RPR/")
    r = w.post(f"/dispatch/{iid}/complete")
    assert "must be generated AND printed" in r.text
    assert get_inv(iid).status == "SLIP_GENERATED"
    blocked = db.scalar(select(func.count(AuditLog.id)).where(AuditLog.action == "COMPLETE_BLOCKED",
                                                              AuditLog.record_id == str(iid)))
    assert blocked >= 1
    # 14 print (HTML) and PDF
    r = w.post(f"/dispatch/{iid}/print", {"size": "A5", "mode": "print"})
    assert r.status_code == 200 and i.dispatch_no in r.text and "ORIGINAL" in r.text
    assert get_inv(iid).status == "SLIP_PRINTED"
    r = w.post(f"/dispatch/{iid}/print", {"size": "A4", "mode": "pdf"})  # reprint without reason -> refused
    assert "Reason is mandatory" in r.text
    r = w.post(f"/dispatch/{iid}/print", {"size": "A4", "mode": "pdf", "reason": "Transporter copy lost"})
    assert r.headers["content-type"] == "application/pdf" and r.content[:4] == b"%PDF"
    logs = db.scalars(select(SlipPrintLog).where(SlipPrintLog.invoice_id == iid).order_by(SlipPrintLog.id)).all()
    assert [lg.is_reprint for lg in logs] == [False, True] and logs[1].reason == "Transporter copy lost"
    # 16 completion
    w.post(f"/dispatch/{iid}/complete")
    i = get_inv(iid)
    assert i.status == "DISPATCHED" and i.completed_by_id is not None and i.erp_sync_status == "PENDING"


# 17-20 ----------------------------------------------------------------------------------------------
def _report(a, code, **params):
    r = a.get(f"/api/v1/reports/{code}", params=params)
    assert r.status_code == 200, r.text
    return r.json()


def test_17_20_mis_reports_reconcile_to_invoice_total(api):
    a = api("management")
    params = {"date_from": (local_today() - timedelta(days=120)).isoformat(), "date_to": local_today().isoformat()}
    inv = _report(a, "invoice-freight", **params)
    total = sum(Decimal(r["freight"] or 0) for r in inv["rows"])
    for code in ("customer-freight", "company-freight", "cfa-freight", "location-freight"):
        rep = _report(a, code, **params)
        assert rep["rows"], code
        assert Decimal(rep["totals"]["freight"]) == total, code
    comp = _report(a, "company-freight", **params)
    assert {"company", "cfa", "invoices", "boxes", "sales", "freight", "freight_per_box", "freight_pct"} <= set(comp["rows"][0])
    cfa = _report(a, "cfa-freight", **params)
    assert {"actual", "variance", "dispatched"} <= set(cfa["rows"][0])


# 21-22 ----------------------------------------------------------------------------------------------
def test_21_22_rate_change_effective_date_and_history(api, db):
    a = api("admin")
    today = local_today()
    old_inv = make_invoice(a, company="COB", cfa="BSP", customer="C-VDP01", boxes=10, inv_date=today - timedelta(days=5))
    assert old_inv["rate_per_box"] == "3.50"
    admin = db.scalar(select(User).where(User.username == "admin"))
    checker = db.scalar(select(User).where(User.username == "checker"))
    cfa_mgr = db.scalar(select(User).where(User.username == "cfa.raipur"))
    rate = db.scalar(select(FreightRate).join(Location, FreightRate.location_id == Location.id)
                     .join(Company, FreightRate.company_id == Company.id).join(CFA, FreightRate.cfa_id == CFA.id)
                     .where(Company.code == "COB", CFA.code == "BSP", Location.code == "BSP",
                            FreightRate.status == RateStatus.APPROVED, FreightRate.effective_to.is_(None)))
    new = rate_service.revise_rate(db, cfa_mgr, rate, rate_per_box=Decimal("4.25"), effective_from=today - timedelta(days=2),
                                   change_reason="Test revision")
    assert new.status == RateStatus.PENDING_APPROVAL
    # maker-checker: creator cannot approve
    import pytest
    from app.services.errors import BusinessError
    with pytest.raises(BusinessError):
        rate_service.approve_rate(db, cfa_mgr, new)
    rate_service.approve_rate(db, checker, new)
    db.commit()
    db.refresh(rate)
    assert rate.effective_to == today - timedelta(days=3)   # old version end-dated, not overwritten
    assert rate.rate_per_box == Decimal("3.50")
    # 22 old invoice keeps the rate it was priced with (even though revision is back-dated to cover it? no:
    #    its date is before the new effective-from), and invoices already priced are never re-priced
    assert a.get(f"/api/v1/invoices/{old_inv['id']}").json()["rate_per_box"] == "3.50"
    fresh = make_invoice(a, company="COB", cfa="BSP", customer="C-VDP01", boxes=10)
    assert fresh["rate_per_box"] == "4.25" and fresh["transportation_expense"] == "42.50"
    # history cannot be overwritten: a version starting before the latest one is refused
    with pytest.raises(BusinessError):
        rate_service.create_rate(db, admin, company_id=rate.company_id, cfa_id=rate.cfa_id, location_id=rate.location_id,
                                 rate_per_box=Decimal("1"), effective_from=today - timedelta(days=10))
    db.rollback()
    # already-priced invoice dated inside a later back-dated version keeps its frozen rate
    newest = db.get(FreightRate, new.id)
    inv_mid = make_invoice(a, company="COB", cfa="BSP", customer="C-VDP01", boxes=4, inv_date=today)
    rev2 = rate_service.revise_rate(db, cfa_mgr, newest, rate_per_box=Decimal("5.00"), effective_from=today - timedelta(days=1),
                                    change_reason="Back-dated test")
    rate_service.approve_rate(db, checker, rev2)
    db.commit()
    assert a.get(f"/api/v1/invoices/{inv_mid['id']}").json()["rate_per_box"] == "4.25"


# 23 -------------------------------------------------------------------------------------------------
def test_23_manual_override_with_approval_and_audit(api, web, db):
    inv = make_invoice(api("admin"), company="COA", cfa="RPR", customer="C-VDP01", boxes=30)
    op = web("operator.raipur")
    op.post(f"/invoices/{inv['id']}/override", {"field": "BOXES", "new_value": "32", "reason": "Recount at dock"})
    i = get_inv(inv["id"])
    assert i.status == "EXCEPTION" and i.exception_code == "OVERRIDE_PENDING" and i.boxes == 30
    # operator cannot approve; checker approves
    from app.models import InvoiceOverride
    ov = db.scalar(select(InvoiceOverride).where(InvoiceOverride.invoice_id == inv["id"]))
    assert op.post(f"/overrides/{ov.id}/decide", {"decision": "approve"}).status_code == 403
    web("checker").post(f"/overrides/{ov.id}/decide", {"decision": "approve", "remarks": "Checked"})
    i = get_inv(inv["id"])
    assert i.status == "PENDING" and i.boxes == 32 and i.freight_amount == Decimal("256.00")  # frozen ₹8 rate
    entries = db.scalars(select(AuditLog).where(AuditLog.record_type == "Invoice", AuditLog.record_id == str(inv["id"]),
                                                AuditLog.module == "OVERRIDE")).all()
    actions = {e.action for e in entries}
    assert {"OVERRIDE_REQUESTED", "OVERRIDE_APPROVED"} <= actions
    ap = next(e for e in entries if e.action == "OVERRIDE_APPROVED")
    assert ap.old_value == "30" and ap.new_value == "32" and ap.username == "checker" and "Recount" in ap.reason
    # rate override applied directly by checker (still audited)
    web("checker").post(f"/invoices/{inv['id']}/override", {"field": "RATE", "new_value": "7.5",
                                                            "reason": "Special negotiated rate"})
    i = get_inv(inv["id"])
    assert i.rate_per_box == Decimal("7.50") and i.freight_amount == Decimal("240.00") and i.rate_level == "MANUAL"
    # reason is mandatory
    r = web("checker").post(f"/invoices/{inv['id']}/override", {"field": "RATE", "new_value": "7", "reason": ""})
    assert "reason" in r.text.lower()


# 24 -------------------------------------------------------------------------------------------------
def test_24_excel_csv_pdf_export(web):
    m = web("management")
    r = m.get("/reports/customer-freight/export?fmt=xlsx&date_from=2000-01-01")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws["A1"].value == "Customer-wise Transportation Expense"
    assert ws.max_row > 5
    r = m.get("/reports/invoice-freight/export?fmt=csv&date_from=2000-01-01")
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][:3] == ["Date", "Company", "Depot"] and len(rows) > 20
    r = m.get("/reports/company-freight/export?fmt=pdf&date_from=2000-01-01")
    assert r.content[:4] == b"%PDF"


# 25 -------------------------------------------------------------------------------------------------
def test_25_role_based_permissions(web, api, db):
    viewer = web("viewer")
    assert viewer.get("/reports/customer-freight").status_code == 200
    assert viewer.get("/reports/customer-freight/export?fmt=csv").status_code == 403   # no export right
    assert viewer.get("/dispatch/pending").status_code == 403
    assert viewer.get("/admin/users").status_code == 403
    op = web("operator.raipur")
    assert op.get("/rates/new").status_code == 403
    assert op.get("/masters/customers").status_code == 403
    assert op.get("/admin/settings").status_code == 403
    # CFA scoping: Bilaspur operator cannot open a Raipur invoice
    rpr_inv = db.scalar(select(Invoice).join(CFA).where(CFA.code == "RPR"))
    assert web("operator.bilaspur").get(f"/dispatch/{rpr_inv.id}").status_code == 403
    assert api("operator.bilaspur").get(f"/api/v1/invoices/{rpr_inv.id}").status_code == 404
    # only admin may cancel an already completed dispatch
    done = db.scalar(select(Invoice).join(CFA).where(Invoice.status == "DISPATCHED", CFA.code == "RPR"))
    r = web("checker").post(f"/dispatch/{done.id}/cancel", {"reason": "Testing permission"})
    assert "Only an administrator" in r.text
    assert get_inv(done.id).status == "DISPATCHED"


# 26-27 ----------------------------------------------------------------------------------------------
def test_26_27_transporter_bill_reconciliation_and_variance(db, web):
    acc = db.scalar(select(User).where(User.username == "accounts"))
    chk = db.scalar(select(User).where(User.username == "checker"))
    tr = db.scalar(select(Transporter).where(Transporter.code == "VDL"))
    today = local_today()
    start = today - timedelta(days=45)
    bill = recon_service.create_bill(db, acc, transporter_id=tr.id, bill_no=f"T-{today:%y%m%d}-VDL", bill_date=today,
                                     bill_amount="0", period_from=start, period_to=today)
    n = recon_service.link_dispatches(db, acc, bill)
    assert n >= 1
    amounts = {ln.id: str(ln.calculated_amount + (Decimal("50") if k == 0 else 0)) for k, ln in enumerate(bill.lines)}
    recon_service.set_line_amounts(db, acc, bill, amounts)
    bill.bill_amount = sum(ln.billed_amount for ln in bill.lines)
    recon_service.rematch(db, bill)
    assert bill.status == "VARIANCE" and bill.variance == Decimal("50.00")
    first = bill.lines[0]
    assert first.match_status == "VARIANCE" and first.variance == Decimal("50.00")
    # one LR line can cover several invoices of the consignment: the billed amount is spread over them
    ids = first.invoice_ids or [first.invoice_id]
    assert sum(db.get(Invoice, i).actual_freight for i in ids) == first.billed_amount
    inv = db.get(Invoice, first.invoice_id)
    import pytest
    from app.services.errors import BusinessError
    with pytest.raises(BusinessError):        # maker cannot approve own bill
        recon_service.approve(db, acc, bill, "self approval attempt")
    with pytest.raises(BusinessError):        # variance approval needs remarks
        recon_service.approve(db, chk, bill, "")
    recon_service.approve(db, chk, bill, "Loading charges accepted")
    recon_service.mark_paid(db, acc, bill, "UTR123456789")
    db.commit()
    assert db.get(TransporterBill, bill.id).status == "PAID"
    # duplicate billing of the same dispatch in another bill is flagged
    b2 = recon_service.create_bill(db, acc, transporter_id=tr.id, bill_no=f"T-{today:%y%m%d}-VDL-2", bill_date=today,
                                   bill_amount="100")
    from app.models import TransporterBillLine
    b2.lines.append(TransporterBillLine(lr_no=inv.lr_no, billed_amount=Decimal("100")))
    db.flush()
    recon_service.rematch(db, b2)
    assert b2.lines[0].match_status == "DUPLICATE_BILLING"
    db.commit()
    # variance report shows the variance line
    rep = web("management").get("/reports/freight-variance?date_from=2000-01-01")
    assert rep.status_code == 200 and bill.bill_no in rep.text
