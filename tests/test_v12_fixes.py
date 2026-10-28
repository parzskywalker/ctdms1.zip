"""Regression tests for the v1.2 fixes (second, independent QA round - IDs QA-03, 07, 09, 10, 11, 13)."""
from __future__ import annotations

import io
import os
import sqlite3
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models import Invoice, InvoiceStatus, Role, User
from app.security.auth import hash_password
from app.services.exporters import XLSX_CELL_LIMIT, to_pdf, to_xlsx
from app.services.pdf_service import table_pdf
from app.services.reports import Col, ReportResult
from tests.helpers import new_invoice_no

ROOT = Path(__file__).resolve().parents[1]


def _make_user(db, username: str, role_code: str, password: str = "Recover#2468x") -> User:
    u = db.scalar(select(User).where(User.username == username))
    role = db.scalar(select(Role).where(Role.code == role_code))
    if not u:
        u = User(username=username, full_name=username.title(), role=role, password_hash=hash_password(password),
                 all_companies=True, all_cfas=True)
        db.add(u)
    else:
        u.role, u.password_hash, u.is_active, u.must_change_password = role, hash_password(password), True, False
    db.commit()
    return u


def _invoice_row() -> dict:
    return {"invoice_no": new_invoice_no("CS"), "invoice_date": "2026-09-01", "company_code": "COA",
            "cfa_code": "RPR", "customer_code": "C-VDP01", "sales_value": "1000", "boxes": 2}


# ----------------------------------------------------------------------------- QA-07 CSRF
@pytest.mark.parametrize("hdr", [{"Authorization": "not-a-bearer-token"}, {"Authorization": "Bearer junk"},
                                 {"X-API-Key": "junk"}])
def test_qa07_fake_auth_header_does_not_skip_csrf(web, hdr):
    """[QA-07] a browser session must still send the CSRF token, whatever other headers are present."""
    w = web("admin")
    row = _invoice_row()
    r = w.c.post("/api/v1/invoices/ingest", json={"invoices": [row]}, headers=hdr)
    # blocked: 403 (CSRF token missing) - or 401 for an invalid Bearer token, which is judged on its own
    assert r.status_code in (401, 403), r.text
    assert w.c.get("/api/v1/invoices", params={"invoice_no": row["invoice_no"]}).json()["total"] == 0


def test_qa07_session_with_csrf_and_real_bearer_still_work(web, api):
    w = web("admin")
    r = w.c.post("/api/v1/invoices/ingest", json={"invoices": [_invoice_row()]}, headers={"X-CSRF-Token": w.token})
    assert r.status_code == 200, r.text
    r = api("admin").post("/api/v1/invoices/ingest", json={"invoices": [_invoice_row()]})
    assert r.status_code == 200, r.text


# ----------------------------------------------------------------------------- QA-03 recovery
def test_qa03_cli_recovery_signs_out_old_sessions_and_tokens(web, api, db):
    """[QA-03] CLI password recovery must end every existing session and API token of that account."""
    from app.cli import main
    _make_user(db, "qa.recover", "ADMIN")
    browser = web("qa.recover", "Recover#2468x")
    token_client = api("qa.recover", "Recover#2468x")
    assert token_client.get("/api/v1/me").status_code == 200
    main(["create-admin", "--username", "qa.recover", "--password", "NewRecover#9753"])
    assert token_client.get("/api/v1/me").status_code == 401
    r = browser.get("/admin/users", follow_redirects=False)
    assert r.status_code in (302, 303) and "/login" in r.headers["location"]
    # the new password works
    assert web("qa.recover", "NewRecover#9753").get("/admin/users").status_code == 200


def test_qa03_cli_will_not_silently_promote(web, db):
    """[QA-03] create-admin on an existing non-admin needs an explicit --promote."""
    from app.cli import main
    _make_user(db, "qa.promote", "MANAGEMENT")
    old = web("qa.promote", "Recover#2468x")
    with pytest.raises(SystemExit):
        main(["create-admin", "--username", "qa.promote", "--password", "Promote#97531"])
    db.expire_all()
    assert db.scalar(select(User).where(User.username == "qa.promote")).role.code == "MANAGEMENT"
    main(["create-admin", "--username", "qa.promote", "--password", "Promote#97531", "--promote"])
    db.expire_all()
    assert db.scalar(select(User).where(User.username == "qa.promote")).role.code == "ADMIN"
    # the session opened before the promotion does not inherit administrator rights
    r = old.get("/admin/users", follow_redirects=False)
    assert r.status_code in (302, 303) and "/login" in r.headers["location"]


def test_qa03_role_change_invalidates_api_token(api, db):
    _make_user(db, "qa.rolechg", "MANAGEMENT")
    c = api("qa.rolechg", "Recover#2468x")
    assert c.get("/api/v1/me").status_code == 200
    u = db.scalar(select(User).where(User.username == "qa.rolechg"))
    u.role = db.scalar(select(Role).where(Role.code == "ADMIN"))
    db.commit()
    assert c.get("/api/v1/me").status_code == 401


def test_qa03_admin_role_change_signs_user_out(web, db):
    _make_user(db, "qa.rolechg2", "VIEWER")
    target = web("qa.rolechg2", "Recover#2468x")
    u = db.scalar(select(User).where(User.username == "qa.rolechg2"))
    mgmt = db.scalar(select(Role).where(Role.code == "MANAGEMENT"))
    admin = web("admin")
    r = admin.post("/admin/users/save", data={"id": u.id, "username": u.username, "full_name": u.full_name,
                                              "role_id": mgmt.id, "is_active": "1", "all_companies": "1",
                                              "all_cfas": "1"})
    assert r.status_code == 200
    r = target.get("/", follow_redirects=False)
    assert r.status_code in (302, 303) and "/login" in r.headers["location"]


# ----------------------------------------------------------------------------- QA-10 big PDF
def _lrs(n):
    return ", ".join(f"LR{100000 + i}" for i in range(n))


def test_qa10_pdf_survives_huge_cell():
    """[QA-10] one row with thousands of LR numbers used to raise LayoutError (HTTP 500)."""
    from pypdf import PdfReader
    big = _lrs(3000)
    pdf = table_pdf("Transporter", "test", [("Transporter", "text"), ("Invoices", "int"), ("LR Numbers", "text")],
                    [["Big Roadways", "3000", big], ["Small Carrier", "2", "LR1, LR2"]], None, "Org")
    assert pdf.startswith(b"%PDF")
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    compact = text.replace("\n", "").replace(" ", "")
    for probe in ("LR100000", "LR101500", "LR102999", "SmallCarrier"):
        assert probe in compact, probe  # nothing silently dropped


def test_qa10_transporter_report_pdf_route(web):
    r = web("admin").get("/reports/transporter/export?fmt=pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_qa10_xlsx_cell_limit_is_respected():
    res = ReportResult([Col("transporter", "Transporter"), Col("lr_numbers", "LR Numbers")],
                       [{"transporter": "Big", "lr_numbers": _lrs(5000)}], {})
    ws = load_workbook(io.BytesIO(to_xlsx(res, "T", "s"))).active
    val = ws.cell(row=5, column=2).value
    assert len(val) <= XLSX_CELL_LIMIT and "more" in val and "CSV" in val
    assert to_pdf(res, "T", "s", "Org").startswith(b"%PDF")


# ----------------------------------------------------------------------------- QA-11 paging
def test_qa11_pending_api_is_paged_and_complete(api, db):
    """[QA-11] the pending queue API reports totals and pages instead of silently cutting at 2,000."""
    c = api("admin")
    first = c.get("/api/v1/dispatch/pending", params={"page_size": 7}).json()
    assert {"items", "total", "page", "page_size", "has_more"} <= set(first)
    total = first["total"]
    assert total > 7 and first["has_more"] is True and len(first["items"]) == 7
    seen, page = [], 1
    while True:
        body = c.get("/api/v1/dispatch/pending", params={"page_size": 250, "page": page}).json()
        seen += [i["id"] for i in body["items"]]
        if not body["has_more"]:
            break
        page += 1
    assert len(seen) == len(set(seen)) == total
    n_open = db.scalar(select(func.count()).select_from(Invoice).where(Invoice.status.in_(InvoiceStatus.OPEN)))
    assert total == n_open
    assert c.get("/api/v1/dispatch/pending", params={"page_size": 5000}).status_code == 422


def test_qa11_pending_api_scope_still_applies(api):
    body = api("operator.raipur").get("/api/v1/dispatch/pending", params={"page_size": 1000}).json()
    assert body["items"] and all(i["cfa"]["code"] == "RPR" for i in body["items"])


# ----------------------------------------------------------------------------- QA-09 startup
def _run_prepare(db_path: Path) -> subprocess.CompletedProcess:
    env = {**os.environ, "CTDMS_DATABASE_URL": f"sqlite:///{db_path}", "CTDMS_TEST_DATABASE_URL": "",
           "CTDMS_ENVIRONMENT": "development"}
    return subprocess.run([sys.executable, "-m", "app.cli", "prepare-demo"], cwd=ROOT, env=env,
                          capture_output=True, text=True, timeout=600)


def _count(db_path: Path, table: str) -> int:
    con = sqlite3.connect(db_path)
    try:
        return con.execute(f"select count(*) from {table}").fetchone()[0]
    finally:
        con.close()


def test_qa09_existing_empty_database_gets_demo_data(tmp_path):
    """[QA-09] an existing but empty ctdms.db used to skip setup, so the advertised demo logins gave 401."""
    db_path = tmp_path / "ctdms.db"
    db_path.write_bytes(b"")  # file exists, nothing inside
    r = _run_prepare(db_path)
    assert r.returncode == 0, r.stdout + r.stderr
    assert _count(db_path, "users") >= 9 and _count(db_path, "invoices") > 0
    # second run: populated database is left exactly as it is
    inv = _count(db_path, "invoices")
    con = sqlite3.connect(db_path)
    con.execute("update users set full_name='KEEP ME' where username='admin'")
    con.commit()
    con.close()
    r = _run_prepare(db_path)
    assert r.returncode == 0 and "ready" in r.stdout.lower(), r.stdout + r.stderr
    assert _count(db_path, "invoices") == inv
    con = sqlite3.connect(db_path)
    assert con.execute("select full_name from users where username='admin'").fetchone()[0] == "KEEP ME"
    con.close()


def test_qa09_data_without_users_is_never_wiped(tmp_path):
    db_path = tmp_path / "ctdms.db"
    assert _run_prepare(db_path).returncode == 0
    con = sqlite3.connect(db_path)
    con.execute("pragma foreign_keys=off")
    con.execute("delete from user_sessions")
    con.execute("update audit_logs set user_id=null")
    con.execute("update users set is_active=0")
    con.commit()
    con.close()
    inv = _count(db_path, "invoices")
    r = _run_prepare(db_path)
    assert r.returncode == 2, r.stdout + r.stderr
    assert "create-admin" in r.stdout
    assert _count(db_path, "invoices") == inv


# ----------------------------------------------------------------------------- QA-13 concurrency
def _sqlite_path():
    url = os.environ["CTDMS_DATABASE_URL"]
    return url.replace("sqlite:///", "") if url.startswith("sqlite") else None


@pytest.mark.skipif(not os.environ["CTDMS_DATABASE_URL"].startswith("sqlite"), reason="SQLite locking only")
def test_qa13_page_views_do_not_wait_for_a_writer(web):
    """[QA-13] a page view must not queue behind someone else's save (v1.1 waited up to 30 s, then error 500)."""
    import time
    w = web("admin")
    assert w.get("/").status_code == 200  # session bookkeeping done
    con = sqlite3.connect(_sqlite_path(), timeout=1, isolation_level=None)
    con.execute("BEGIN IMMEDIATE")  # another user is in the middle of saving
    try:
        t0 = time.perf_counter()
        for url in ("/", "/invoices", "/dispatch/pending", "/reports/company-freight", "/api/v1/dispatch/pending"):
            assert w.get(url).status_code == 200, url
        assert time.perf_counter() - t0 < 10
    finally:
        con.execute("ROLLBACK")
        con.close()


def test_qa13_read_that_collides_is_retried(api, monkeypatch):
    import sqlite3 as _sq
    from sqlalchemy.exc import OperationalError
    import app.api.v1 as v1
    real, calls = v1.dashboard_data, []

    def flaky(*a, **kw):
        calls.append(1)
        if len(calls) == 1:
            raise OperationalError("SELECT ...", {}, _sq.OperationalError("database is locked"))
        return real(*a, **kw)

    monkeypatch.setattr(v1, "dashboard_data", flaky)
    r = api("admin").get("/api/v1/dashboard/summary")
    assert r.status_code == 200 and len(calls) == 2


# ----------------------------------------------------------------------------- follow-ups from the independent review
@pytest.mark.skipif(not os.environ["CTDMS_DATABASE_URL"].startswith("sqlite"), reason="SQLite locking only")
def test_review_idle_session_bookkeeping_never_waits(web, db):
    """A session not seen for > 60 s needs a 'last seen' write; that must not make the page wait for a writer."""
    import time
    from datetime import timedelta
    from app.models import UserSession, utcnow
    w = web("admin")
    for s in db.scalars(select(UserSession).where(UserSession.revoked_at.is_(None))):
        s.last_seen_at = utcnow() - timedelta(minutes=5)
    db.commit()
    con = sqlite3.connect(_sqlite_path(), timeout=1, isolation_level=None)
    con.execute("BEGIN IMMEDIATE")
    try:
        t0 = time.perf_counter()
        assert w.get("/").status_code == 200
        assert time.perf_counter() - t0 < 5
    finally:
        con.execute("ROLLBACK")
        con.close()


def test_review_no_retry_after_something_was_saved(api, db, monkeypatch):
    """The retry safety net must never run a request twice once it has saved something."""
    import sqlite3 as _sq
    from sqlalchemy.exc import OperationalError
    import app.api.v1 as v1
    from app.models import AuditLog
    calls = []

    def saves_then_collides(dbs, user, f):
        calls.append(1)
        dbs.add(AuditLog(user_id=user.id, username=user.username, action="QA_PROBE", module="QA", ref="retry"))
        dbs.commit()
        raise OperationalError("SELECT ...", {}, _sq.OperationalError("database is locked"))

    monkeypatch.setattr(v1, "dashboard_data", saves_then_collides)
    c = api("admin")
    c_noraise = c
    with pytest.raises(OperationalError):
        c_noraise.get("/api/v1/dashboard/summary")
    assert len(calls) == 1
    assert db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action == "QA_PROBE")) == 1


def test_review_cursor_paging_skips_nothing_while_queue_changes(api, db):
    c = api("admin")
    first = c.get("/api/v1/dispatch/pending", params={"page_size": 5}).json()
    ids = [i["id"] for i in first["items"]]
    victim = db.get(Invoice, ids[1])  # leaves the open queue after page 1 was read
    old_status, victim.status = victim.status, InvoiceStatus.CANCELLED
    db.commit()
    try:
        cur = first["next_cursor"]
        while cur:
            body = c.get("/api/v1/dispatch/pending", params={"page_size": 250, "after": cur}).json()
            ids += [i["id"] for i in body["items"]]
            cur = body["next_cursor"]
        n_open = db.scalar(select(func.count()).select_from(Invoice).where(Invoice.status.in_(InvoiceStatus.OPEN)))
        assert len(ids) == len(set(ids)) == n_open + 1  # everything still open, plus the one read before it left
    finally:
        victim.status = old_status
        db.commit()
    assert c.get("/api/v1/dispatch/pending", params={"page": 10 ** 19}).status_code == 422
    assert c.get("/api/v1/invoices", params={"page": 10 ** 19}).status_code == 422
    assert c.get("/api/v1/dispatch/pending", params={"after": "garbage"}).status_code == 422


def test_review_bearer_token_decides_identity(web, api):
    """Cookie + valid Bearer: the token decides (no CSRF needed); cookie + invalid Bearer: 401, no fallback."""
    w = web("admin")
    op_token = api("operator.raipur").headers["Authorization"]
    me = w.c.get("/api/v1/me", headers={"Authorization": op_token}).json()
    assert me["username"] == "operator.raipur"
    assert w.c.get("/api/v1/me", headers={"Authorization": "Bearer not.a.jwt"}).status_code == 401
    r = w.c.post("/api/v1/invoices/ingest", json={"invoices": [_invoice_row()]},
                 headers={"Authorization": api("admin").headers["Authorization"]})
    assert r.status_code == 200, r.text


def test_review_reactivated_user_old_token_stays_dead(web, api, db):
    _make_user(db, "qa.react", "VIEWER")
    c = api("qa.react", "Recover#2468x")
    assert c.get("/api/v1/me").status_code == 200
    u = db.scalar(select(User).where(User.username == "qa.react"))
    admin = web("admin")
    form = {"id": u.id, "username": u.username, "full_name": u.full_name, "role_id": u.role_id,
            "all_companies": "1", "all_cfas": "1"}
    admin.post("/admin/users/save", data={**form, "is_active": "0"})
    assert c.get("/api/v1/me").status_code == 401
    admin.post("/admin/users/save", data={**form, "is_active": "1"})
    assert c.get("/api/v1/me").status_code == 401  # v1.2 first cut: 200 again
    assert api("qa.react", "Recover#2468x").get("/api/v1/me").status_code == 200  # fresh token works


def test_review_pdf_note_labels():
    from pypdf import PdfReader
    big = _lrs(400)
    pdf = table_pdf("T", "s", [("LR Numbers", "text"), ("Count", "int")], [[big, "400"], ["", "1"]], None, "Org")
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert text.replace("\n", "").replace(" ", "").count("LR100399") == 1  # printed once, not twice
    pdf = table_pdf("T", "s", [("Transporter", "text"), ("LR Numbers", "text")], [["", big]], None, "Org")
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "None" not in text and "row 1" in text


def test_review_non_ascii_csrf_token_is_403_not_500(web):
    w = web("admin")
    r = w.c.post("/api/v1/invoices/ingest", json={"invoices": [_invoice_row()]},
                 headers={"X-CSRF-Token": "tökén".encode("utf-8")})
    assert r.status_code == 403
    r = w.c.post("/admin/users/save", data={"csrf_token": "ünïcödé", "username": "x"})
    assert r.status_code == 403


def test_review_prepare_never_wipes_partial_data(tmp_path):
    from sqlalchemy import create_engine, text
    from app.db import Base
    import app.models  # noqa: F401
    db_path = tmp_path / "ctdms.db"
    eng = create_engine(f"sqlite:///{db_path}")
    Base.metadata.create_all(eng)
    with eng.begin() as con:  # masters typed in by hand, no users / companies / invoices yet
        con.execute(text("insert into locations (code, name, is_active, created_at, updated_at) "
                         "values ('KRB', 'Korba', 1, '2026-09-01', '2026-09-01')"))
    eng.dispose()
    r = _run_prepare(db_path)
    assert r.returncode == 2, r.stdout + r.stderr
    assert _count(db_path, "locations") == 1 and _count(db_path, "invoices") == 0


def test_upgrade_from_v11_database_keeps_data(tmp_path):
    """Starting v1.2 on a v1.1 database adds the new column (migration 0002) and keeps every record."""
    db_path = tmp_path / "ctdms.db"
    assert _run_prepare(db_path).returncode == 0
    env = {**os.environ, "CTDMS_DATABASE_URL": f"sqlite:///{db_path}", "CTDMS_TEST_DATABASE_URL": ""}
    down = subprocess.run([sys.executable, "-m", "alembic", "downgrade", "0002"], cwd=ROOT, env=env,
                          capture_output=True, text=True, timeout=600)  # v1.3 -> v1.2 structure first
    assert down.returncode == 0, down.stdout + down.stderr
    con = sqlite3.connect(db_path)  # turn it back into a v1.1 database
    con.execute("alter table users drop column tokens_valid_after")
    con.execute("update alembic_version set version_num='0001'")
    con.execute("update users set full_name='KEEP ME' where username='admin'")
    con.commit()
    inv = con.execute("select count(*) from invoices").fetchone()[0]
    con.close()
    r = _run_prepare(db_path)
    assert r.returncode == 0 and "ready" in r.stdout.lower(), r.stdout + r.stderr
    con = sqlite3.connect(db_path)
    assert con.execute("select version_num from alembic_version").fetchone()[0] == "0014"  # current head
    inv_cols = [c[1] for c in con.execute("pragma table_info(invoices)")]
    assert {"taxable_value", "delivery_status", "eway_bill_no", "qr_signature"} <= set(inv_cols)
    assert "tokens_valid_after" in [c[1] for c in con.execute("pragma table_info(users)")]
    assert "party_type" in [c[1] for c in con.execute("pragma table_info(customers)")]
    assert con.execute("select full_name from users where username='admin'").fetchone()[0] == "KEEP ME"
    assert con.execute("select count(*) from invoices").fetchone()[0] == inv
    con.close()
