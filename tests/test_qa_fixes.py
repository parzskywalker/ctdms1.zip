"""Regression tests for the issues found in the v1.0 QA round (QA report IDs in brackets)."""
from __future__ import annotations

import io
from datetime import timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import (
    CFA, Company, Customer, ERPCodeMapping, FreightRate, ImportBatch, Invoice, InvoiceStatus, Location, Transporter,
    TransporterBill, TransporterBillLine, User,
)
from app.services import recon_service
from app.services.errors import BusinessError
from app.services.utils import csv_safe, local_today, parse_decimal, safe_int
from tests.helpers import make_invoice

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _flash(html: str) -> str:
    import re
    return " | ".join(re.sub(r"<[^>]+>", "", m).strip()
                      for m in re.findall(r'class="alert[^"]*"[^>]*>(.*?)</div>', html, re.S))


# ----------------------------------------------------------------------------- data scope
def _scoped_user(web, username, role_code, cfa_codes):
    """A user of `role_code` restricted to the given depots (created through the real Users screen)."""
    rid = db_role_id(role_code)
    ids = []
    from tests.conftest import _test_session
    s_ = _test_session()
    try:
        ids = [str(s_.scalar(select(CFA.id).where(CFA.code == c))) for c in cfa_codes]
    finally:
        s_.close()
    web("admin").post("/admin/users/save", data={"username": username, "full_name": username, "role_id": str(rid),
                                                 "password": "Scope@2026xy", "is_active": "1", "all_companies": "1",
                                                 "cfas": ids})
    return web(username, "Scope@2026xy")


def db_role_id(code):
    from app.models import Role
    from tests.conftest import _test_session
    s_ = _test_session()
    try:
        return s_.scalar(select(Role.id).where(Role.code == code))
    finally:
        s_.close()


def test_sec01_transporter_bill_scope(web, db):
    """[SEC-01] a depot-restricted user cannot see or approve another depot's transporter bill. (A Depot Manager has no
    bill permission at all since v1.6.1 - 403; the scope rule is shown with a depot-restricted Accounts-type user.)"""
    acc = db.scalar(select(User).where(User.username == "accounts"))
    bsp = db.scalar(select(CFA).where(CFA.code == "BSP"))
    co = db.scalar(select(Company).where(Company.code == "COA"))
    tr = db.scalar(select(Transporter).where(Transporter.code == "CGR"))
    bill = recon_service.create_bill(db, acc, transporter_id=tr.id, bill_no="QA-XCFA-1", bill_date=local_today(),
                                     bill_amount="100", company_id=co.id, cfa_id=bsp.id)
    db.commit()
    assert web("cfa.raipur").get(f"/recon/{bill.id}").status_code == 403
    cfa = _scoped_user(web, "qa.acc.rpr", "ACCOUNTS", ["RPR"])
    assert cfa.get(f"/recon/{bill.id}").status_code == 404
    assert cfa.post(f"/recon/{bill.id}/approve", {"remarks": "cross-cfa approve"},
                    follow_redirects=False).status_code == 404
    assert "QA-XCFA-1" not in cfa.get("/recon").text
    assert "QA-XCFA-1" not in cfa.get("/approvals").text
    assert web("accounts").get(f"/recon/{bill.id}").status_code == 200  # all-scope user still sees it


def test_sec02_import_batch_scope(web, db):
    """[SEC-02] a restricted user cannot open / edit / cancel another CFA's import batch."""
    admin = web("admin")
    bsp = db.scalar(select(CFA).where(CFA.code == "BSP"))
    csv = ("company_code,invoice_no,invoice_date,customer_code,location,sales_value,boxes\n"
           f"COA,QA-BSP-{local_today():%y%m%d}-1,{local_today():%d-%m-%Y},C-ABC01,Bilaspur,1000,3\n")
    r = admin.post("/imports/upload", {"cfa_id": str(bsp.id)}, files={"file": ("bsp.csv", csv.encode(), "text/csv")})
    assert r.status_code == 200
    batch = db.scalar(select(ImportBatch).order_by(ImportBatch.id.desc()))
    cfa = web("cfa.raipur")
    assert cfa.get(f"/imports/{batch.id}").status_code == 404
    assert cfa.post(f"/imports/{batch.id}/cancel", follow_redirects=False).status_code == 404
    db.expire_all()
    assert db.get(ImportBatch, batch.id).status != "CANCELLED"
    assert batch.batch_no not in cfa.get("/imports").text


def test_sec03_importer_cannot_remap_existing_erp_code(web, db):
    """[SEC-03] changing an existing ERP customer mapping needs Masters-edit rights."""
    m = db.scalar(select(ERPCodeMapping).where(ERPCodeMapping.entity_type == "CUSTOMER"))
    assert m is not None
    other = db.scalar(select(Customer).where(Customer.id != m.target_id))
    batch = db.scalar(select(ImportBatch).where(ImportBatch.erp_system_id == m.erp_system_id)
                      .order_by(ImportBatch.id.desc()))
    if batch is None:
        pytest.skip("no batch for the mapped ERP in demo data")
    batch.created_by_id = db.scalar(select(User.id).where(User.username == "cfa.raipur"))  # make it visible
    db.commit()
    r = web("cfa.raipur").post(f"/imports/{batch.id}/map-customer",
                               {"source_code": m.source_code, "customer_id": str(other.id)})
    assert "already mapped" in _flash(r.text)
    db.expire_all()
    assert db.get(ERPCodeMapping, m.id).target_id != other.id


def test_sec07_rate_cfa_scope(web, db):
    """[SEC-07] a depot-restricted user cannot create / revise another depot's rates. (A Depot Manager cannot touch
    rates at all since v1.6.1 - 403; the scope rule is shown with an Administrator-role user limited to Raipur.)"""
    bsp_rate = db.scalar(select(FreightRate).join(CFA, FreightRate.cfa_id == CFA.id)
                         .where(CFA.code == "BSP", FreightRate.status == "APPROVED"))
    loc = db.scalar(select(Location))
    dm = web("cfa.raipur")
    assert dm.post(f"/rates/{bsp_rate.id}/revise", {"rate_per_box": "99", "effective_from":
                   (local_today() + timedelta(days=40)).isoformat(), "change_reason": "x"},
                   follow_redirects=False).status_code == 403
    assert dm.post("/rates", {"company_id": str(bsp_rate.company_id), "cfa_id": str(bsp_rate.cfa_id),
                              "location_id": str(loc.id), "rate_per_box": "9", "effective_from":
                              (local_today() + timedelta(days=41)).isoformat(), "remarks": "x"},
                   follow_redirects=False).status_code == 403
    cfa = _scoped_user(web, "qa.admin.rpr", "ADMIN", ["RPR"])
    r = cfa.post(f"/rates/{bsp_rate.id}/revise", {"rate_per_box": "99", "effective_from":
                 (local_today() + timedelta(days=40)).isoformat(), "change_reason": "cross scope"})
    assert r.status_code == 404 or "outside your" in _flash(r.text)
    r = cfa.post("/rates", {"company_id": str(bsp_rate.company_id), "cfa_id": str(bsp_rate.cfa_id),
                            "location_id": str(loc.id), "rate_per_box": "9", "effective_from":
                            (local_today() + timedelta(days=41)).isoformat(), "remarks": "x"})
    assert "only create rates for your own" in _flash(r.text)


# ----------------------------------------------------------------------------- numbers / crashes
@pytest.mark.parametrize("bad", ["Infinity", "-inf", "NaN", "1e400", "sNaN", "1000000000000"])
def test_sec04_parse_decimal_rejects_non_finite(bad):
    """[SEC-04] NaN / Infinity / absurd magnitudes are rejected instead of poisoning reports."""
    with pytest.raises(ValueError):
        parse_decimal(bad)
    with pytest.raises(ValueError):
        parse_decimal(float("inf"))


def test_sec04_infinity_row_rejected_and_reports_stay_up(web):
    admin = web("admin")
    csv = ("company_code,cfa_code,invoice_no,invoice_date,customer_code,location,sales_value,boxes\n"
           f"COA,RPR,QA-INF-1,{local_today():%d-%m-%Y},C-ABC01,Bilaspur,Infinity,3\n")
    r = admin.post("/imports/upload", files={"file": ("inf.csv", csv.encode(), "text/csv")})
    assert r.status_code == 200 and "Not a number" in r.text
    for url in ("/", "/mis", "/reports/company-freight", "/invoices"):
        assert admin.get(url).status_code == 200


@pytest.mark.parametrize("url", ["/imports/upload", "/rates/import"])
@pytest.mark.parametrize("name", ["bad.xlsx", "bad.xls"])
def test_bug01_corrupt_excel_gives_friendly_error(web, url, name):
    """[BUG-01] a damaged / fake Excel file no longer crashes the importers with HTTP 500."""
    r = web("admin").post(url, files={"file": (name, b"this is not an excel file", XLSX)})
    assert r.status_code == 200
    assert "Could not read" in _flash(r.text)


def test_bug01_corrupt_excel_on_bill_import(web, db):
    bill = db.scalar(select(TransporterBill).where(TransporterBill.status.in_(recon_service.EDITABLE)))
    r = web("accounts").post(f"/recon/{bill.id}/import", files={"file": ("bill.xlsx", b"garbage", XLSX)})
    assert r.status_code == 200 and "Could not read" in _flash(r.text)


def test_sec05_upload_size_limit_on_all_importers(web, monkeypatch):
    """[SEC-05] rate / bill imports now honour the upload limit (a 28 MB CSV used to tie up the server)."""
    from app.config import get_settings
    monkeypatch.setattr(get_settings(), "max_upload_mb", 1)
    big = b"company,cfa,location,rate_per_box,effective_from\n" + b"COA,RPR,Raipur,1,01-01-2027\n" * 60_000
    r = web("admin").post("/rates/import", files={"file": ("big.csv", big, "text/csv")})
    assert "too large" in _flash(r.text).lower()


@pytest.mark.parametrize("url", ["/invoices?page=abc", "/dispatch/pending?page=abc", "/dispatch/register?page=x",
                                 "/masters/customers?page=abc", "/reports/company-freight?company_id=99999999999999999999",
                                 "/masters/erp-mappings?entity=FOO&erp_id=abc", "/mis/data?p_company=99999999999999999999"])
def test_bug02_bad_query_params_do_not_500(web, url):
    """[BUG-02] non-numeric / oversized query parameters are ignored instead of raising HTTP 500."""
    assert web("admin").get(url).status_code in (200, 400)


@pytest.mark.parametrize("url", ["/invoices/99999999999999999999", "/rates/99999999999999999999",
                                 "/api/v1/invoices/99999999999999999999"])
def test_bug02_huge_ids_are_404(web, url):
    assert web("admin").get(url).status_code == 404


def test_bug03_api_validation_error_with_form_body_is_422(web):
    """[BUG-03] validation errors that contain raw bytes used to crash the JSON error handler."""
    from fastapi.testclient import TestClient

    from app.main import app
    r = TestClient(app).post("/api/v1/auth/token", data={"username": "admin", "password": "x"})
    assert r.status_code == 422


def test_bug04_admin_user_save_bad_role(web):
    r = web("admin").post("/admin/users/save", {"username": "qa.badrole", "full_name": "QA", "role_id": "abc",
                                                "password": "Strong#Pass2026"})
    assert r.status_code == 200 and "Select a role" in _flash(r.text)


# ----------------------------------------------------------------------------- exports
def test_sec06_formula_injection_neutralised(web, db, api):
    """[SEC-06] text starting with = + - @ is exported as text, not as a live Excel formula."""
    assert csv_safe('=HYPERLINK("x")') == "'=HYPERLINK(\"x\")" and csv_safe(-5) == -5 and csv_safe("ABC") == "ABC"
    admin = web("admin")
    loc = db.scalar(select(Location).where(Location.name == "Raipur"))
    admin.post("/masters/customers/save", {"code": "C-QAFX", "name": '=HYPERLINK("http://evil.example","x")',
                                           "location_id": str(loc.id), "is_active": "on"})
    make_invoice(api("admin"), customer="C-QAFX", location="Raipur", boxes=2, sales="1000")
    d0 = (local_today() - timedelta(days=1)).isoformat()
    x = admin.get(f"/reports/invoice-freight/export?fmt=xlsx&date_from={d0}&date_to={local_today().isoformat()}")
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    cells = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert cells and all(v.startswith("'") for v in cells)
    c = admin.get(f"/reports/invoice-freight/export?fmt=csv&date_from={d0}&date_to={local_today().isoformat()}")
    assert ',=HYPERLINK' not in c.text and ',"=HYPERLINK' not in c.text
    m = admin.get("/masters/customers?export=csv")
    assert ',=HYPERLINK' not in m.text and ',"=HYPERLINK' not in m.text


# ----------------------------------------------------------------------------- auth
def test_sec08_forced_password_change_is_enforced(web, db):
    """[SEC-08] a user flagged 'must change password' cannot use anything else (web or API) until they do."""
    u = db.scalar(select(User).where(User.username == "management"))
    u.must_change_password = True
    db.commit()
    try:
        from fastapi.testclient import TestClient

        from app.main import app
        c = web("management")
        r = c.get("/reports", follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/account/password"
        assert c.get("/account/password").status_code == 200
        t = TestClient(app).post("/api/v1/auth/token", json={"username": "management", "password": "Demo@123"})
        assert t.status_code == 403
    finally:
        u.must_change_password = False
        db.commit()


def test_sec10_login_messages_and_password_blocklist(web):
    from fastapi.testclient import TestClient

    from app.main import app
    from app.security.auth import password_problems
    c = TestClient(app)
    import re
    tok = re.search(r'name="csrf_token" value="([^"]+)"', c.get("/login").text).group(1)
    a = c.post("/login", data={"username": "nosuchuser", "password": "x", "csrf_token": tok})
    b = c.post("/login", data={"username": "checker", "password": "wrong", "csrf_token": tok})
    msg = lambda r: re.search(r'class="alert[^"]*"[^>]*>([^<]+)', r.text).group(1)  # noqa: E731
    assert msg(a) == msg(b) and "lock the account" in msg(a)  # same text + tells the rule up front
    assert password_problems("password1") and password_problems("Checker2026x", "checker")
    assert not password_problems("Truck#Route42")


def test_sec09_audit_ip_not_spoofable(web, db):
    from app.models import AuditLog
    c = web("viewer")
    c.c.post("/logout", data={"csrf_token": c.token}, headers={"X-Forwarded-For": "203.0.113.66"})
    last = db.scalar(select(AuditLog).where(AuditLog.username == "viewer").order_by(AuditLog.id.desc()))
    assert last is None or last.ip != "203.0.113.66"


def test_sec11_security_headers(web):
    r = web("admin").get("/")
    assert "frame-ancestors 'self'" in r.headers.get("content-security-policy", "")


# ----------------------------------------------------------------------------- recon integrity
def test_rec01_bill_total_excludes_problem_lines_and_blocks_approval(db):
    """[REC-01] not-dispatched / other-transporter lines no longer inflate the system total or pass approval."""
    acc = db.scalar(select(User).where(User.username == "accounts"))
    chk = db.scalar(select(User).where(User.username == "checker"))
    tr = db.scalar(select(Transporter).where(Transporter.code == "CGR"))
    disp = db.scalar(select(Invoice).where(Invoice.status == InvoiceStatus.DISPATCHED, Invoice.transporter_id == tr.id,
                                           Invoice.transporter_bill_id.is_(None)))
    pend = db.scalar(select(Invoice).where(Invoice.status == InvoiceStatus.PENDING, Invoice.freight_amount.isnot(None)))
    bill = recon_service.create_bill(db, acc, transporter_id=tr.id, bill_no="QA-REC-01", bill_date=local_today(),
                                     bill_amount=str(disp.freight_amount + pend.freight_amount))
    bill.lines.append(TransporterBillLine(invoice_no=disp.invoice_no, billed_amount=disp.freight_amount))
    bill.lines.append(TransporterBillLine(invoice_no=pend.invoice_no, billed_amount=pend.freight_amount))
    db.flush()
    recon_service.rematch(db, bill)
    assert {ln.match_status for ln in bill.lines} == {"MATCHED", "NOT_DISPATCHED"}
    assert bill.calculated_amount == disp.freight_amount          # pending invoice not counted
    assert bill.variance == pend.freight_amount                   # over-billing is now visible
    with pytest.raises(BusinessError, match="cannot be paid"):
        recon_service.approve(db, chk, bill, "please approve anyway")
    db.rollback()


def test_rec02_approver_who_edited_amounts_cannot_approve(db):
    acc = db.scalar(select(User).where(User.username == "accounts"))
    # a transporter that still has unbilled dispatches (other tests bill Vidarbha's - this test used to depend on order)
    tid = db.scalar(select(Invoice.transporter_id).where(
        Invoice.status == InvoiceStatus.DISPATCHED, Invoice.transporter_bill_id.is_(None),
        Invoice.transporter_id.isnot(None), Invoice.dispatch_date >= local_today() - timedelta(days=60)).limit(1))
    tr = db.get(Transporter, tid)
    admin = db.scalar(select(User).where(User.username == "admin"))
    bill = recon_service.create_bill(db, admin, transporter_id=tr.id, bill_no="QA-REC-02", bill_date=local_today(),
                                     bill_amount="0", period_from=local_today() - timedelta(days=60),
                                     period_to=local_today())
    recon_service.link_dispatches(db, admin, bill)
    recon_service.set_line_amounts(db, acc, bill, {ln.id: str(ln.calculated_amount) for ln in bill.lines})
    bill.bill_amount = sum(ln.billed_amount for ln in bill.lines)
    recon_service.rematch(db, bill)
    with pytest.raises(BusinessError, match="another approver"):
        recon_service.approve(db, acc, bill, "edited then approved")
    db.rollback()


# ----------------------------------------------------------------------------- UX regressions
def test_ux_prefilled_date_time_not_reported_missing(web, api):
    inv = make_invoice(api("admin"), boxes=5)
    html = web("operator.raipur").get(f"/dispatch/{inv['id']}").text
    assert "Dispatch date missing" not in html and "Dispatch time missing" not in html


def test_ux_blocked_invoices_sorted_last_in_queue(web):
    import re
    html = web("operator.raipur").get("/dispatch/pending").text
    statuses = re.findall(r'class="badge b-([A-Z_]+)"', html.split("<tbody>", 1)[1])
    blocked = [i for i, s in enumerate(statuses) if s in ("RATE_MISSING", "EXCEPTION")]
    ready = [i for i, s in enumerate(statuses) if s not in ("RATE_MISSING", "EXCEPTION")]
    assert not blocked or not ready or min(blocked) > max(ready)


def test_ux_roles_without_invoice_view_see_no_dead_invoice_links(web, db):
    """A role made by the Master with reports but no invoice register must not get links into invoices."""
    from app.models import Role
    a = web("admin")
    from tests.helpers import roles_matrix
    a.post("/admin/roles/save", data=roles_matrix(db, {"new_code": "RPT_ONLY", "new_name": "Reports only", "new_desc": "test"}))
    rid = db.scalar(select(Role.id).where(Role.code == "RPT_ONLY"))
    a.post("/admin/roles/save", data=roles_matrix(db, {f"p_{rid}_report.view": "1", f"p_{rid}_master.use": "1"}))
    a.post("/admin/users/save", data={"username": "qa.rptonly", "full_name": "Rpt", "role_id": str(rid),
                                      "password": "Rptonly@2026", "is_active": "1", "all_companies": "1",
                                      "all_cfas": "1"})
    v = web("qa.rptonly", "Rptonly@2026")
    assert 'href="/invoices/' not in v.get("/reports/dispatch-register").text
    assert v.get("/mis").status_code == 403  # the drill-down MIS is a final report: not for such a role
    assert "CAN_OPEN_INVOICE = true" in web("management").get("/mis").text


def test_ux_irreversible_actions_ask_for_confirmation(web, api, db):
    inv = make_invoice(api("admin"), boxes=5)
    assert "data-confirm=" in web("admin").get(f"/dispatch/{inv['id']}").text            # cancel dispatch asks first
    open_slip = db.scalar(select(Invoice.dispatch_slip_id).where(Invoice.status == InvoiceStatus.SLIP_GENERATED))
    assert "data-confirm=" in web("admin").get(f"/gatepass/{open_slip}").text            # so does voiding a gatepass


def test_ux_read_only_role_sees_slip_number(web, db):
    inv = db.scalar(select(Invoice).where(Invoice.status == InvoiceStatus.SLIP_GENERATED))
    html = web("checker").get(f"/dispatch/{inv.id}").text
    assert "No active slip" not in html and inv.dispatch_no in html and f"/gatepass/{inv.dispatch_slip_id}" in html


def test_ux_import_flags_implausible_box_count(web):
    csv = ("company_code,cfa_code,invoice_no,invoice_date,customer_code,location,sales_value,boxes\n"
           f"COA,RPR,QA-BIG-{local_today():%y%m%d},{local_today():%d-%m-%Y},C-ABC01,Bilaspur,50000,99999\n")
    r = web("admin").post("/imports/upload", files={"file": ("big.csv", csv.encode(), "text/csv")})
    assert "unusually high" in r.text or "% of the sales value" in r.text


def test_safe_int_helper():
    assert safe_int("12") == 12 and safe_int("abc") is None and safe_int("99999999999999999999") is None
    assert safe_int("-3", 1, minimum=1) == 1 and safe_int(None, 7) == 7
    assert parse_decimal("1,234.50") == Decimal("1234.50")
